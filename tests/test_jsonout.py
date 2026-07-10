"""리뷰 JSON(check --json)과 AI 소견 시트(annotate-ai) — /dsd-review 계약."""
import json

from openpyxl import load_workbook

from dsdxlsx import ties
from dsdxlsx.check import run_checks
from dsdxlsx.jsonout import build_review_json
from dsdxlsx.parse import parse_contents

from fixtures import factory


def _payload(**kw):
    doc = parse_contents(factory.standard_contents(**kw))
    report = run_checks(doc)
    sugg = ties.suggest_ties(doc, existing=[])
    pts = ties.resolve_points(doc, sugg)
    return doc, build_review_json(doc, report, [], pts, source="합성.dsd")


def test_review_json_structure():
    _, data = _payload()
    assert data["schema"] == 1 and data["source"] == "합성.dsd"

    d = data["document"]
    labels = {s["label"] for s in d["sections"]}
    assert "재무상태표" in labels and "주석4" in labels
    assert {n["no"] for n in d["notes"]} == {1, 4, 5, 7}

    # FS 행: 라벨·참조·기간별 금액
    ar = next(r for r in d["fs_rows"] if r["label"] == "매출채권")
    assert ar["refs"] == [5] and ar["amounts"] == {"당기": 5000, "전기": 4000}

    # 문단 텍스트 (오타 검토 재료) — 주석 헤딩 포함
    texts = [p["text"] for p in d["paragraphs"]]
    assert any("5. 매출채권" in t for t in texts)

    # 표: 행 라벨·헤더·높이
    t = next(t for t in d["tables"] if t["section"] == "주석5")
    assert {"매출채권", "대손충당금", "차감후 장부금액"} <= {r["label"] for r in t["rows"]}
    assert t["height"] > 0

    # 대사 제안이 JSON에 실려 나감 (스킬의 '대사 제안 정제' 재료)
    assert data["ties"]["suggested"]
    s = data["ties"]["suggested"][0]
    assert "tie" in s and "a_value" in s and "b_value" in s

    # 검증 결과·커버리지 동봉 + JSON 직렬화 가능
    assert data["coverage"] and "counts" in data["checks"]
    json.dumps(data, ensure_ascii=False, default=str)


def test_review_json_pgbrk_positions():
    contents = factory.document(
        factory.p("표지"),
        factory.title("재무제표", 2),
        factory.statement_title_table("재 무 상 태 표"),
        factory.pgbrk(),
        factory.table([["과목", "당기", "전기"]], [["현금", "1", "1"]], unit="원"),
    )
    doc = parse_contents(contents)
    data = build_review_json(doc, run_checks(doc), [], [])
    brks = data["document"]["pgbrks"]
    assert len(brks) == 1
    assert brks[0]["after_block"] == "t001"  # 제목표 뒤·본문 표 앞


def test_annotate_ai_sheet_and_import_safety(tmp_path):
    from dsdxlsx.annotate_xlsx import AI_SHEET, annotate_ai
    from dsdxlsx.export_xlsx import export_xlsx
    from dsdxlsx.import_xlsx import import_xlsx

    contents = factory.standard_contents()
    dsd_path = tmp_path / "s.dsd"
    dsd_path.write_bytes(factory.dsd_bytes(contents))
    xlsx = str(tmp_path / "s.xlsx")
    doc = parse_contents(contents)
    export_xlsx(doc, dsd_path.read_bytes(), str(dsd_path), xlsx)
    _, contents2, imp = import_xlsx(xlsx)
    doc2 = parse_contents(contents2)

    row_id = doc2.tables["t002"].rows[2].id  # 아무 표 행
    findings = [
        {"kind": "오타", "severity": "warn", "row_id": row_id,
         "section": "재무상태표", "message": "계정명 오타 의심", "suggestion": "…"},
        {"kind": "주기", "severity": "info", "section": "주석5",
         "message": "참조 적절성 참고"},
    ]
    stats = annotate_ai(xlsx, findings, doc2, imp.row_positions)
    assert stats["findings"] == 2

    wb = load_workbook(xlsx)
    assert AI_SHEET in wb.sheetnames
    ws = wb[AI_SHEET]
    cells = {c.value for row in ws.iter_rows() for c in row if c.value}
    assert "계정명 오타 의심" in cells
    # 하이퍼링크 생성 (표 행 ID 지정 시)
    linked = [c for row in ws.iter_rows(min_col=4, max_col=4) for c in row
              if c.hyperlink is not None]
    assert linked

    # AI검토 시트가 있어도 import 무변경
    _, _, imp2 = import_xlsx(xlsx)
    assert not imp2.changed

    # 재실행 시 교체 (누적 아님)
    annotate_ai(xlsx, findings[:1], doc2, imp.row_positions)
    wb2 = load_workbook(xlsx)
    vals = [c.value for row in wb2[AI_SHEET].iter_rows() for c in row if c.value]
    assert "참조 적절성 참고" not in vals
