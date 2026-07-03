"""샘플 DSD 기반 E2E: 무수정 왕복 바이트 동일, 셀 수정, 행 추가/삭제, 가드."""
import hashlib
import re

import pytest
from openpyxl import load_workbook

from dsdxlsx import grid as grid_mod
from dsdxlsx.cli import main
from dsdxlsx.dsd_io import read_dsd
from dsdxlsx.import_xlsx import ImportError_, import_xlsx
from dsdxlsx.parse import parse_contents


@pytest.fixture(scope="module")
def workspace(sample_dsd, tmp_path_factory):
    """export 1회 실행 후 (원본 contents, doc, xlsx 경로) 제공."""
    tmp = tmp_path_factory.mktemp("dsdx")
    xlsx = str(tmp / "sample.xlsx")
    assert main(["export", sample_dsd, "-o", xlsx]) == 0
    package = read_dsd(sample_dsd)
    doc = parse_contents(package.contents)
    return {"contents": package.contents, "doc": doc, "xlsx": xlsx, "tmp": tmp}


def sha(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def exported_table_ids(doc):
    from dsdxlsx.export_xlsx import exported_sections

    return {b.item.id for s in exported_sections(doc) for b in s.blocks if b.kind == "table"}


def find_editable_table(doc):
    """행 조작 테스트용: 내보낸 표 + 고정 아님 + TBODY 3행 이상 + 병합 없음 + 숫자 셀."""
    exported = exported_table_ids(doc)
    for table in doc.tables.values():
        if table.id not in exported or table.fixtable or table.irregular:
            continue
        g = grid_mod.expand(table)
        body = [i for i, r in enumerate(table.rows) if not r.in_thead]
        if len(body) < 3:
            continue
        if any(gc and gc.kind != "cell" for i in body for gc in g[i] if gc):
            continue
        if not all(table.rows[i].acopy and table.rows[i].adelete for i in body):
            continue
        has_number = any(
            re.match(r"^\d{1,3}(,\d{3})+$", c.text.strip())
            for i in body
            for c in table.rows[i].cells
        )
        if has_number:
            return table
    pytest.fail("테스트 조건에 맞는 표를 찾지 못함")


def locate_row(xlsx_path, row_id):
    wb = load_workbook(xlsx_path)
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_col=1, max_col=1):
            if row and str(row[0].value) == row_id:
                return wb, ws, row[0].row
    pytest.fail(f"xlsx에서 행 {row_id}를 찾지 못함")


# ---------------------------------------------------------------- M1/M2

def test_parse_counts(workspace):
    doc = workspace["doc"]
    assert len(doc.tables) > 0
    assert sum(1 for b in doc.blocks if b.kind == "p") > 0
    # 모든 노드의 바이트 범위가 원본과 정합: 슬라이스가 태그로 열리고 닫힘
    data = workspace["contents"]
    for table in doc.tables.values():
        assert data[table.start : table.start + 6] == b"<TABLE"
        assert data[table.end - 1 : table.end] == b">"


def test_noedit_roundtrip_is_byte_identical(workspace):
    package, new_contents, report = import_xlsx(workspace["xlsx"])
    assert not report.changed
    assert sha(new_contents) == sha(workspace["contents"])


def test_noedit_full_cli_roundtrip(workspace, sample_dsd):
    out = str(workspace["tmp"] / "roundtrip.dsd")
    assert main(["import", workspace["xlsx"], "-o", out]) == 0
    pkg_out = read_dsd(out)
    pkg_in = read_dsd(sample_dsd)
    assert sha(pkg_out.contents) == sha(pkg_in.contents)
    assert sha(pkg_out.meta) == sha(pkg_in.meta)


# ---------------------------------------------------------------- M3 셀 수정

def test_cell_edit_is_localized(workspace):
    doc = workspace["doc"]
    table = find_editable_table(doc)
    target = None
    for row in table.rows:
        if row.in_thead:
            continue
        for cell in row.cells:
            if re.match(r"^\d{1,3}(,\d{3})+$", cell.text.strip()):
                target = cell
                break
        if target:
            break
    row_id = ".".join(target.id.split(".")[:2])

    wb, ws, r = locate_row(workspace["xlsx"], row_id)
    # 해당 행에서 숫자 값을 가진 열을 찾아 수정
    col = next(
        c.column for c in ws[r]
        if c.column > 1 and isinstance(c.value, (int, float)) and c.value == int(target.text.strip().replace(",", ""))
    )
    ws.cell(row=r, column=col, value=999888777)
    edited = str(workspace["tmp"] / "edit1.xlsx")
    wb.save(edited)

    _, new_contents, report = import_xlsx(edited)
    assert report.cell_edits == [target.id]
    assert report.row_inserts == [] and report.row_deletes == []
    assert b"999,888,777" in new_contents
    # 변경은 해당 셀 하나뿐: 원본과 새 바이트의 차이가 그 노드 범위 안에만 존재
    old = workspace["contents"]
    prefix = target.inner_start
    assert new_contents[:prefix] == old[:prefix]
    tail_old = old[target.inner_end:]
    assert new_contents.endswith(tail_old)


def test_text_edit_with_newline_and_quote(workspace):
    from dsdxlsx.export_xlsx import exported_sections

    doc = workspace["doc"]
    exported_paras = {
        b.item.id for s in exported_sections(doc) for b in s.blocks if b.kind == "p"
    }
    pid = next(
        i for i, n in doc.nodes.items()
        if i in exported_paras and not n.complex
    )
    wb, ws, r = locate_row(workspace["xlsx"], pid)
    ws.cell(row=r, column=2, value='첫 줄 "인용"\n둘째 줄 <태그> & 앰퍼샌드')
    edited = str(workspace["tmp"] / "edit2.xlsx")
    wb.save(edited)

    _, new_contents, report = import_xlsx(edited)
    assert f"{pid} (문단)" in report.cell_edits
    assert b"&quot;" in new_contents
    assert "첫 줄 &quot;인용&quot;&amp;cr;둘째 줄 &lt;태그&gt; &amp; 앰퍼샌드".encode() in new_contents
    parse_contents(new_contents)  # well-formed


# ---------------------------------------------------------------- M4 행 추가/삭제

def test_row_insert(workspace):
    doc = workspace["doc"]
    table = find_editable_table(doc)
    body_rows = [r for r in table.rows if not r.in_thead]
    anchor = body_rows[1]

    wb, ws, r = locate_row(workspace["xlsx"], anchor.id)
    ws.insert_rows(r + 1)
    ws.cell(row=r + 1, column=2, value="신규계정과목")
    ws.cell(row=r + 1, column=4, value=12345678)
    edited = str(workspace["tmp"] / "edit3.xlsx")
    wb.save(edited)

    _, new_contents, report = import_xlsx(edited)
    assert len(report.row_inserts) == 1
    assert report.cell_edits == []  # 행 삽입이 아래 표들의 값을 오염시키지 않음
    assert report.row_deletes == []
    assert "신규계정과목".encode() in new_contents
    assert b"12,345,678" in new_contents
    doc2 = parse_contents(new_contents)
    assert len(doc2.tables[table.id].rows) == len(table.rows) + 1


def test_row_delete(workspace):
    doc = workspace["doc"]
    table = find_editable_table(doc)
    victim = [r for r in table.rows if not r.in_thead][1]

    wb, ws, r = locate_row(workspace["xlsx"], victim.id)
    ws.delete_rows(r)
    edited = str(workspace["tmp"] / "edit4.xlsx")
    wb.save(edited)

    _, new_contents, report = import_xlsx(edited)
    assert len(report.row_deletes) == 1
    assert report.cell_edits == []
    assert report.row_inserts == []
    doc2 = parse_contents(new_contents)
    assert len(doc2.tables[table.id].rows) == len(table.rows) - 1


# ---------------------------------------------------------------- 가드

def test_guard_readonly_tu_cell(workspace):
    doc = workspace["doc"]
    exported = exported_table_ids(doc)
    tu = next(
        (n for n in doc.nodes.values()
         if n.readonly and n.tag == "TU" and n.id.split(".")[0] in exported),
        None,
    )
    if tu is None:
        pytest.skip("내보낸 구획에 TU 셀 없음 (표지 TU는 시트 미생성)")
    row_id = ".".join(tu.id.split(".")[:2])
    table = doc.tables[row_id.split(".")[0]]
    i = int(row_id.split(".r")[1]) - 1
    col = grid_mod.cell_grid_col(table, i, tu)

    wb, ws, r = locate_row(workspace["xlsx"], row_id)
    ws.cell(row=r, column=2 + col, value="변조 시도")
    edited = str(workspace["tmp"] / "edit5.xlsx")
    wb.save(edited)

    with pytest.raises(ImportError_, match="자동 필드"):
        import_xlsx(edited)


def test_guard_thead_row_delete(workspace):
    doc = workspace["doc"]
    table = find_editable_table(doc)
    head = next((r for r in table.rows if r.in_thead), None)
    if head is None:
        pytest.skip("이 표에 THEAD 없음")

    wb, ws, r = locate_row(workspace["xlsx"], head.id)
    ws.delete_rows(r)
    edited = str(workspace["tmp"] / "edit6.xlsx")
    wb.save(edited)

    with pytest.raises(ImportError_, match="THEAD"):
        import_xlsx(edited)


def test_guard_fixtable_insert(workspace):
    doc = workspace["doc"]
    exported = exported_table_ids(doc)
    fixed = next(
        (t for t in doc.tables.values() if t.fixtable and t.rows and t.id in exported),
        None,
    )
    if fixed is None:
        pytest.skip("내보낸 구획에 고정 표 없음 (표지 고정 표는 시트 미생성)")

    wb, ws, r = locate_row(workspace["xlsx"], fixed.rows[-1].id)
    ws.insert_rows(r + 1)
    ws.cell(row=r + 1, column=2, value="불법 삽입")
    edited = str(workspace["tmp"] / "edit7.xlsx")
    wb.save(edited)

    with pytest.raises(ImportError_, match="AFIXTABLE"):
        import_xlsx(edited)
