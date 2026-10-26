"""DSD ↔ 조서(최종분석적검토 등) 대사 — recon.

DSD 재무제표·주석의 숫자 칸마다 조서 엑셀에서 같은 값을 찾는다. 조서 양식은 회사마다
달라(공시형 BS·IS 시트 / 계정형 BS + 공시형 FS·FN 시트 …) 시트·행 위치에 기대지 않는다.

찾는 순서 (칸 하나):
  1) 같은 이름 행(정규화 라벨 키 겹침) × 같은 기간 열(머리글 날짜·3개월) → 일치
  2) 같은 이름 행 여러 개의 합(하위 계정: '지분법적용투자주식(갑)' + '(을)')  → 일치(하위 합)
  3) 같은 이름 행은 있는데 값이 다름 (가장 가까운 조서 값을 보여 줌)          → 불일치
  4) 같은 이름 행이 없을 때만: 같은 범위의 다른 행(재무제표는 같은 기간 열,
     주석은 같은 주석 구획)                                                → 일치(다른 행)
  5) 아무것도 없음 / 같은 이름 행의 그 칸이 비어 있음                       → 상대 없음
비교는 단위(원/천원 반올림 ±1)·비율(27.27 ↔ 0.2727)·손실 표기 부호('분기순손실'을
양수로 적는 관례)를 맞춘다. 조서에 그 기간 열이 없으면(주석 시트에는 보통 당기만 있음)
전기 칸은 '대사 대상 아님'으로 따로 센다 (--prior 로 값 찾기 강제).
"""
from __future__ import annotations

import datetime as _dt
import difflib
import re
from dataclasses import dataclass, field
from decimal import Decimal

from openpyxl import load_workbook

from . import dart_client as D
from . import period as P
from .check import norm
from .model import Document
from .rollinterim import _CUR_MARK_RE, _PREV_MARK_RE, TableRec, _leaf, scan

# 조서 시트 자동 분류 — 시트 이름(정규화 · 소문자, '연결' 머리 · '(공시)/(정산)' 꼬리를 뗀 것)
# 또는 시트 위 제목 칸이 이 이름 그대로면 재무제표 시트. 'BS합산' · 'CF조정' 같은 보조 시트는 아님.
FS_NAME_KIND = {"bs": "BS", "재무상태표": "BS",
                "is": "PL", "pl": "PL", "손익계산서": "PL", "포괄손익계산서": "PL",
                "cf": "CF", "현금흐름표": "CF",
                "sce": "SCE", "ce": "SCE", "자본변동표": "SCE",
                "fs": "FS", "재무제표": "FS"}
_GRADE_RE = re.compile(r"\((공시|정산)\)$")
NOTE_SHEET_RE = re.compile(r"주석|^fn$|^notes?$", re.I)
# 주석마다 따로 둔 시트 ('6. 현금및현금성자산', '18. 차입금 (3)')
NOTE_EACH_RE = re.compile(r"^\s*(\d{1,2})\s*\.\s*(\S.*)$")
_INDEX_RE = re.compile(r"index|목차|색인", re.I)

_DATE_RE = re.compile(r"(20\d{2})\s*[.\-/년]\s*(\d{1,2})\s*[.\-/월]\s*(\d{1,2})")
# 기간 길이 '3개월' — 만기 구간 '3개월 미만 · 3개월에서 1년'은 아님
_3M_RE = re.compile(r"3\s*개\s*월(?!\s*(미만|이내|이하|이상|초과|에서|~|-|～))")
_NOTE_TITLE_RE = re.compile(r"^\s*(\d{1,2})\s*\.\s*(\S.*)$")
_PAREN_RE = re.compile(r"\(([^()]*)\)")
_LOSS_RE = re.compile(r"손실|결손|차감|손$|감소|누계액|\(-\)")
_RESULT_RE = re.compile(r"(영업|순|법인세비용차감전순|법인세차감전순)(이익|손실)$")
# 기간이 아니라 계정인 머리글 ('당기순손익' 열이 당기 열로 잡히지 않게)
_ACCT_HDR_RE = re.compile(r"당기순|당기손익|전기오류")
# 금액이 아닌 열 (공정가치 수준 등) — 대사하지 않음
_SKIP_HDR_RE = re.compile(r"수준|레벨|level|주석", re.I)
_CF_DIR_RE = re.compile(r"취득|처분|지급|상환|증가|감소|배당|차입|납부|수취|유입|유출|회수|예치|인출")
# 행 라벨의 기간 표시 — '당기순이익' 같은 계정 이름은 빼고 '당기초·전분기말'만
_ROW_CUR_RE = re.compile(r"당\s*\d?\s*(분|반)?\s*기\s*(초|말)|\(당\)")
_ROW_PREV_RE = re.compile(r"전\s*\d?\s*(분|반)?\s*기\s*(초|말)|\(전\)")


# ── 라벨 키 ─────────────────────────────────────────────────────────

def _canon(text: str) -> str:
    t = re.sub(r"\(\s*주석[^)]*\)", "", text)          # '(주석3)', '(주석12,13)'
    t = _DATE_RE.sub("", t)
    t = re.sub(r"당\s*(분|반)\s*기", "당기", t)
    t = re.sub(r"전\s*(분|반)\s*기", "전기", t)
    k = D.canon(t).lower()
    k = re.sub(r"의?(증가|감소)\((증가|감소)\)$", "", k)      # 현금흐름표 '미수금의 감소(증가)'
    k = re.sub(r"(.{2,})의(취득|처분|증가|감소|상환|지급|발행|차입|납부|수취|회수)$", r"\1", k)
    k = k.replace("으로인한", "").replace("에의한", "")
    m = _RESULT_RE.search(k)
    if m:                                               # 영업손실 = 영업이익(손실) 행
        k = k[: m.start(2)] + "이익"
    return k


# 같은 줄의 다른 이름 - DSD 표준 문구(K-GAAP 현금흐름표 소제목 등) ↔ 조서 · 정산표의 짧은 이름. 키(정규화) → 덧붙일 키
LABEL_ALIAS = {
    "현금의유출이없는비용등의가산": "비용가산", "현금유출이없는비용등의가산": "비용가산",
    "현금의유입이없는수익등의차감": "수익차감", "현금유입이없는수익등의차감": "수익차감",
    "영업활동자산부채의변동": "자산부채변동", "영업활동자산부채변동": "자산부채변동",
    "기초의현금": "기초현금현금성자산", "기말의현금": "기말현금현금성자산",
    "기초현금": "기초현금현금성자산", "기말현금": "기말현금현금성자산",
    "미처리결손금": "미처분이익잉여금", "이익잉여금(결손금)": "이익잉여금",
}


def label_keys(texts: list[str]) -> set[str]:
    """행의 글자 칸들 → 비교 키 집합: 전체 · 괄호 밖 · 괄호 안 각각의 정규화(+ 같은 줄 다른 이름)."""
    keys: set[str] = set()
    for raw in texts:
        s = raw.replace("\xa0", " ").replace("　", " ").strip()
        if not s:
            continue
        for p in [s, _PAREN_RE.sub("", s)]:
            k = _canon(p)
            if k and not k.isdigit():
                keys.add(k)
        for p in _PAREN_RE.findall(s):          # 괄호 안: '(주)'·'(유)' 회사 형태 표기는 버림
            k = _canon(p)
            if k and k not in ("주", "유", "재", "사", "합") and not k.isdigit():
                keys.add(k)
    keys |= {LABEL_ALIAS[k] for k in keys if k in LABEL_ALIAS}
    return keys


# ── 조서 읽기 ───────────────────────────────────────────────────────

@dataclass
class WpRow:
    sheet: str
    row: int
    label: str
    keys: set[str]
    nums: dict[int, float]          # 열 번호 → 값
    # 그 행에 걸린 머리글: 열 → (날짜 | None, 당기/전기 | None, 3개월 여부)
    cols: dict[int, tuple] = field(default_factory=dict)
    sect: str | None = None          # 표 위 소제목 '① 당기말'의 기간 (열 머리글이 기간을 안 줄 때)
    unit: int = 1                    # 시트 위 '(단위: 천원)' 배수
    names: dict[int, set] = field(default_factory=dict)   # 열 → 머리글 이름 키 ('연이자율' · '감가상각')


@dataclass
class WpSheet:
    name: str
    kind: str                        # fs | note
    rows: list[WpRow] = field(default_factory=list)
    blocks: list[tuple[int | None, str, int, int]] = field(default_factory=list)  # (번호, 제목키, 시작, 끝)
    fs_kind: str = ""                # BS | PL | CF | SCE | FS(여러 재무제표가 한 시트)
    grade: str = ""                  # 공시 | 정산 | ''
    consolidated: bool = False       # 시트 이름 · 제목에 '연결'


def _as_date(v) -> tuple[int, int, int] | None:
    if isinstance(v, (_dt.datetime, _dt.date)):
        return (v.year, v.month, v.day)
    if isinstance(v, str):
        m = _DATE_RE.search(v)
        if m:
            return (int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return None


def _is_num(v) -> bool:
    return isinstance(v, (int, float, Decimal)) and not isinstance(v, bool)


# 머리글 칸의 기간 표시: '당분기말' · '제 40(당) 기' · '(전)' … ('당기순이익' 같은 계정은 _ACCT_HDR_RE 로 뺌)
_HDR_CUR_RE = re.compile(r"\(당\)|당\d?(분|반)?기")
_HDR_PREV_RE = re.compile(r"\(전\)|전\d?(분|반)?기")
_CUM_RE = re.compile(r"누적|누계")
# 표 위 소제목이 기간 하나만 말할 때 ('① 당기말', '(2) 전분기', '<당기>')
_SECT_RE = re.compile(r"^[\[<(]?[①-⑳\d]*[).]?[\[<(]?(당|전)\d?(분|반)?기(말|초)?(현재)?[\]>)]?:?$")
_UNIT_RE = re.compile(r"단위\s*[:：]\s*([^)\s]+)")
_UNIT_MULT = {"원": 1, "천원": 1000, "백만원": 1_000_000, "십억원": 1_000_000_000}


def _hdr_token(v) -> tuple | None:
    """머리글 칸 → (날짜 | None, 당기/전기 | None, 3개월 여부, 누적 여부). 기간 말이 없으면 None."""
    if isinstance(v, (_dt.datetime, _dt.date)):
        return ((v.year, v.month, v.day), None, False, False)
    if not isinstance(v, str):
        return None
    s = norm(v)
    if not s or len(s) > 30 or _ACCT_HDR_RE.search(s):
        return None
    d = _as_date(s)
    if d is None and (m := re.fullmatch(r"(?:fy)?(20\d{2})년?(?:도)?", s, re.I)):
        d = (int(m.group(1)), 0, 0)              # 연도만 - 그 해 날짜 아무것과 맞음
    p = "당기" if _HDR_CUR_RE.search(s) else "전기" if _HDR_PREV_RE.search(s) else None
    t3, cum = bool(_3M_RE.search(s)), bool(_CUM_RE.search(s))
    if d or p or t3 or cum:
        return (d, p, t3, cum)
    return None


_TOTAL_NAMES = {"총계", "계", "합", "합계", "총합계"}


_PERIOD_WORD_RE = re.compile(r"제\s*\d+\s*\(\s*(당|전)\s*\)\s*\d?\s*(분|반)?\s*기\s*(말|초)?"
                             r"|(당|전)\s*\d?\s*(분|반)?\s*기\s*(말|초)?|누적|누계")


def _name_keys(text: str, marks: bool = True) -> set[str]:
    """열 머리글 → 이름 키 ('총장부금액' · '연이자율'). 기간 말 · 날짜 · 숫자는 뺀다."""
    t = _3M_RE.sub("", _DATE_RE.sub("", _PERIOD_WORD_RE.sub("", text)))
    keys = {k for k in label_keys([t]) if len(k) >= 2 and not k.isdigit()}
    out = {"합계" if k in _TOTAL_NAMES else k for k in keys}
    n = norm(text)
    if not marks:
        return out
    if re.search(r"기초|기\)?초$", n):                # '당기초' · '전기초' · '기초' - 변동표 첫 열
        out.add("기초")
    if re.search(r"기말|기\)?말$|분기말$|반기말$", n):     # '당분기말' · '전기말' · '기말'
        out.add("기말")
    if any(k.endswith(("총계", "합계")) for k in out):   # '자본총계' 열 = '총 계' 열
        out.add("합계")
    return out


def _plain(keys: set[str]) -> set[str]:
    return {k.lstrip("^") for k in keys}


def _leafk(keys: set[str]) -> set[str]:
    return {k for k in keys if not k.startswith("^")}


def _names_hit(names: set[str], keys: set[str], exact: bool = False) -> bool:
    """열 이름이 같은가 — exact 가 아니면 한쪽이 다른 쪽을 품어도 같다고 본다('기초' ↔ '기초장부금액').
    keys 의 '^' 머리 키(위 층 머리글 '매출 등')는 정확히 같을 때만."""
    if names & _plain(keys):
        return True
    return not exact and any(len(a) >= 2 and len(b) >= 2 and (a in b or b in a)
                             for a in names for b in _leafk(keys))


def _named(row: "WpRow", cols, keys: set[str]) -> list[int]:
    """열 중 이름이 같은 것 — 정확히 같은 이름 키가 가장 많이 겹치는 열('매출 등/기타' ↔ '매입 등/기타'),
    없으면 품는 이름."""
    leaf, up = _leafk(keys), _plain(keys) - _leafk(keys)
    scored = [(2 * len(row.names.get(c, set()) & leaf) + len(row.names.get(c, set()) & up), c) for c in cols]
    best = max((sc for sc, _c in scored), default=0)
    if best >= 2:                              # 잎 이름이 하나는 맞아야(위 층 '채무'만 맞는 열은 아님)
        return [c for sc, c in scored if sc == best]
    return [c for c in cols if _names_hit(row.names.get(c, set()), _leafk(keys))]


def _unit_of(texts: list[str]) -> int | None:
    for t in texts:
        m = _UNIT_RE.search(t)
        if m and (u := _UNIT_MULT.get(norm(m.group(1)).split(",")[-1])):
            return u
    return None


def classify_sheet(name: str, top: list[list]) -> tuple[str, str, str, bool] | None:
    """시트 → (fs | note, 재무제표 종류, 공시/정산, 연결 여부). 대상이 아니면 None.
    재무제표 = 시트 이름 또는 위 5행 제목 칸이 'BS' · '재무상태표(공시)' · '연결 BS' 그대로.
    주석 = 이름에 '주석' · 'FN' 이 있거나 이름이 'N. 제목'(주석마다 시트)."""
    n = norm(name).lower()
    titles = [norm(str(v)).lower() for r in top[:5] for v in r if isinstance(v, str) and v.strip()]
    cons = n.startswith("연결") or any(t.startswith("연결") for t in titles[:3])

    def fs_name(t: str):
        for part in t.split("|"):             # 제목 칸 'BS | 재무상태표(공시)'
            base = _GRADE_RE.sub("", part)
            base = base[2:] if base.startswith("연결") else base
            if base in FS_NAME_KIND:
                return FS_NAME_KIND[base]
        return None

    fk = fs_name(n) or next((k for t in titles if (k := fs_name(t))), None)
    if fk:
        grade = next((m.group(1) for t in [n] + titles for part in t.split("|")
                      if (m := _GRADE_RE.search(part))), "")
        return ("fs", fk, grade, cons)
    if _INDEX_RE.search(n):
        return None
    if NOTE_SHEET_RE.search(n) or NOTE_EACH_RE.match(name):
        return ("note", "", "", cons)
    return None


def _with_spans(names: dict[int, set], name_spans: list[tuple[int, int | None, set]]) -> dict[int, set]:
    """열 이름 + 그 열을 덮는 병합 머리글 이름(같은 줄 다음 글 칸 앞까지 - 다른 줄에서 이름이 붙은 열만)."""
    out = {c: set(v) for c, v in names.items()}
    cap = max(names) + 1 if names else 0
    for s, e, k in name_spans:
        for c in range(s, cap if e is None else e):
            if c in out:                          # 아래 층에 이름이 있는 열만(빈 열로 번지지 않게)
                out[c] |= k
    return out


def _read_sheet(ws, kind: str) -> WpSheet:
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    sh = WpSheet(ws.title, kind)
    each = NOTE_EACH_RE.match(ws.title) if kind == "note" else None
    if each:                                  # 주석마다 시트 — 시트 전체가 한 구획
        sh.blocks.append((int(each.group(1)), _canon(each.group(2)), 1, 0))
    hdr: dict[int, tuple] = {}                # 지금 표의 머리글: 열 → (날짜, 기간, 3개월)
    spans: list[tuple[int, int, tuple]] = []  # 병합 머리글이 덮는 열 범위 [시작, 끝)
    names: dict[int, set] = {}                # 열 → 머리글 이름 키
    name_spans: list[tuple[int, int | None, set]] = []   # 병합 머리글 이름이 덮는 열 [시작, 끝) - 끝 None = 표 끝
    ref_cols: set[int] = set()
    sect: str | None = None
    unit = 1
    prev_hdr = False
    for ri, r in enumerate(grid, start=1):
        texts = [str(v) for v in r if isinstance(v, str) and v.strip()]
        if (u := _unit_of(texts)) is not None:
            unit = u
        nums = {c: float(v) for c, v in enumerate(r) if _is_num(v) and c not in ref_cols}
        if len(nums) == 1:                     # 머리글 줄 끝의 환율 · 비율 숫자 하나 ('… | 당기말 | 9.4673')
            tc = [c for c, v in enumerate(r) if isinstance(v, str) and v.strip()]
            if len(tc) >= 4 and any((t := _hdr_token(r[c])) and not (t[0] and t[0][1] == 0) for c in tc[1:]):
                nums = {}
        if not nums:
            text_cols = [c for c, v in enumerate(r)
                         if ((isinstance(v, str) and v.strip() and not _UNIT_RE.search(v))
                             or isinstance(v, (_dt.datetime, _dt.date)))]
            texts = [str(r[c]) for c in text_cols if isinstance(r[c], str)]
            if not each:
                hit = False
                for t in texts:
                    m = _NOTE_TITLE_RE.match(t.replace("\xa0", " "))
                    if m and not _DATE_RE.search(t):
                        sh.blocks.append((int(m.group(1)), _canon(m.group(2)), ri, 0))
                        sect, hdr, spans, names, prev_hdr, hit = None, {}, [], {}, False, True
                        name_spans = []
                        break
                if hit:
                    continue
            if len(text_cols) == 1 and text_cols[0] <= 1:      # 왼쪽 한 칸 글 = 제목 · 소제목 · 설명
                if _SECT_RE.match(norm(texts[0]) if texts else ""):
                    sect = "당기" if "당" in norm(texts[0])[:6] else "전기"
                prev_hdr = False
                continue
            if not text_cols:
                prev_hdr = False
                continue
            # 머리글 줄 — 앞 줄도 머리글이면 이어 붙이고, 아니면 새 표
            if not prev_hdr:
                hdr, spans, names, ref_cols = {}, [], {}, set()
                name_spans = []
            two = len({t[1] for t in (_hdr_token(r[c]) for c in text_cols) if t and t[1]}) >= 2
            for k, c in enumerate(text_cols):
                if isinstance(r[c], str):
                    nk = _name_keys(r[c], marks=not two)   # '당기말 | 전기말' 기간 머리글엔 기초 · 기말 이름을 안 붙임
                    names[c] = names.get(c, set()) | nk
                    end = text_cols[k + 1] if k + 1 < len(text_cols) else None   # 마지막 = 표 끝 열까지
                    name_spans.append((c + 1, end, nk))     # 병합 머리글 '매출 등' 이 덮는 오른쪽 열
            toks = {c: t for c, v in enumerate(r) if (t := _hdr_token(v))}
            years = [c for c, t in toks.items() if t[0] and t[0][1] == 0]
            if not 2 <= len({toks[c][0][0] for c in years}) <= 3:   # 연도 머리글 = '2026년 | 2025년'(2 ~ 3개) - '만기 2028년' · 만기별 표는 아님
                for c in years:
                    del toks[c]
            ref_cols |= {c for c, v in enumerate(r)
                         if isinstance(v, str) and re.fullmatch(r"주석|note|notes?", norm(v), re.I)}
            row_spans = []
            for c, (d, p, t3, cum) in toks.items():
                if d or p:
                    hdr[c] = (d, p, t3)
                    if p and not d:            # 병합 머리글 '당분기말' — 같은 줄 다음 글 칸 앞까지 덮음
                        end = next((x for x in text_cols if x > c), 10 ** 6)
                        row_spans.append((c, end, (None, p, t3)))
            for c, (d, p, t3, cum) in toks.items():
                if d or p:
                    continue
                if c in hdr:                   # 위 줄 날짜 칸 아래 '누적' · '최종3개월'
                    hdr[c] = (hdr[c][0], hdr[c][1], t3)
                    continue
                # 기간 말 없는 '3개월 · 누적' 칸: 위 줄 병합 머리글('제 40(당) 기')이 덮으면 그 기간,
                # 아니면 날짜 칸 중 오른쪽(없으면 왼쪽) 가장 가까운 것 ('6월 누적 | 최종3개월 | 9월 누적')
                sp = next((v for s, e, v in reversed(spans) if s <= c < e), None)
                if sp:
                    hdr[c] = (None, sp[1], t3)
                    continue
                dated = [x for x in hdr if hdr[x][0] and x != c]
                left_all = sorted((x for x in dated if x < c), reverse=True)
                if cum and left_all and toks.get(left_all[0], (None, None, False))[2]:
                    parent = left_all[0]           # 날짜 하나가 '3개월 | 누적' 두 칸 위에 병합 - 누적 칸은 왼쪽 날짜
                else:
                    dated = [x for x in dated if not hdr[x][2]]
                    right = sorted(x for x in dated if x > c)
                    left = sorted((x for x in dated if x < c), reverse=True)
                    parent = right[0] if right else (left[0] if left else None)
                if parent is not None:
                    hdr[c] = (hdr[parent][0], hdr[parent][1], t3)
            spans += row_spans
            if any(hdr[c][1] or hdr[c][0] for c in toks if c in hdr):
                sect = None                    # 열 머리글이 기간을 주면 소제목 기간은 끝
            prev_hdr = True
            continue
        prev_hdr = False
        first = min(nums)
        label_texts = [str(v) for c, v in enumerate(r) if c < first and isinstance(v, str)]
        cols = dict(hdr)
        cap = max(names) + 1 if names else 10 ** 6   # 머리글 글이 있는 열까지 (표 오른쪽 점검 칸 'True' · 0 은 아님)
        for s, e, v in spans:
            for c in nums:
                if s <= c < min(e, cap) and c not in cols:
                    cols[c] = v
        words = {v[1] for v in hdr.values() if v[1] and not v[0]}
        if len(words) == 1 and not any(v[0] for v in hdr.values()):
            only = words.pop()                 # '당기초 | 취득 | 지분법손익 | 당기말' — 표 전체가 한 기간
            for c in nums:
                if c < cap:
                    cols.setdefault(c, (None, only, False))
        sh.rows.append(WpRow(ws.title, ri, " ".join(t.strip() for t in label_texts if t.strip()),
                             label_keys(label_texts), nums, cols, sect, unit, _with_spans(names, name_spans)))
    sh.blocks = [(n, k, s, (sh.blocks[i + 1][2] - 1 if i + 1 < len(sh.blocks) else len(grid)))
                 for i, (n, k, s, _) in enumerate(sh.blocks)]
    if kind == "note" and not sh.blocks:     # '투자부동산주석' — 시트 이름이 곧 주석 제목
        title = re.sub(r"주석|\(.*?\)", "", ws.title).strip()
        sh.blocks.append((None, _canon(title), 1, len(grid)))
    return sh


def load_workpaper(path: str, fs_sheets: list[str] | None = None,
                   note_sheets: list[str] | None = None, skip: list[str] | None = None,
                   fs_map: dict[str, list[str]] | None = None) -> list[WpSheet]:
    """조서 엑셀(계산값)을 읽어 재무제표·주석 시트만 행 단위로 색인.
    fs_map = {'BS': ['BS(공시)'], …} 이나 fs_sheets/note_sheets 로 시트를 지정하면 그 시트만.
    지정이 없으면 이름 · 제목으로 자동 판정하고, 같은 종류에 공시 · 정산 시트가 다 있으면 공시만."""
    import warnings
    with warnings.catch_warnings():          # 인쇄 영역·데이터 유효성 확장 경고는 대사와 무관
        warnings.simplefilter("ignore")
        wb = load_workbook(path, data_only=True, read_only=True)
    given = bool(fs_sheets or note_sheets or fs_map)
    fs_kind_of = {s: k for k, names in (fs_map or {}).items() for s in names}
    fs_kind_of.update({s: "" for s in (fs_sheets or []) if s not in fs_kind_of})
    notes = set(note_sheets or [])
    skip_set = set(skip or [])
    out: list[WpSheet] = []
    for ws in wb.worksheets:
        if ws.title in skip_set:
            continue
        top = [list(r) for r in ws.iter_rows(values_only=True, max_row=5)]
        auto = classify_sheet(ws.title, top)
        if given:
            if ws.title in fs_kind_of:
                kind, fk = "fs", fs_kind_of[ws.title] or (auto[1] if auto and auto[0] == "fs" else "FS")
            elif ws.title in notes:
                kind, fk = "note", ""
            else:
                continue
            grade, cons = (auto[2], auto[3]) if auto else ("", False)
        else:
            if auto is None or getattr(ws, "sheet_state", "visible") != "visible":
                continue
            kind, fk, grade, cons = auto
        sh = _read_sheet(ws, kind)
        sh.fs_kind, sh.grade, sh.consolidated = fk, grade, cons
        out.append(sh)
    wb.close()
    if not given:                             # 같은 종류에 공시 · 정산이 다 있으면 공시만
        public = {s.fs_kind for s in out if s.kind == "fs" and s.grade == "공시"}
        out = [s for s in out if not (s.kind == "fs" and s.grade == "정산" and s.fs_kind in public)]
    return out


def sheet_config(sheets: list[WpSheet]) -> dict:
    """자동 판정 결과 → 건별 설정 모양 (사람이 고쳐 쓰라고 파일로 남김)."""
    fs: dict[str, list[str]] = {}
    for s in sheets:
        if s.kind == "fs":
            fs.setdefault(s.fs_kind or "FS", []).append(s.name)
    return {"재무제표": fs, "주석": [s.name for s in sheets if s.kind == "note"],
            "빼는 시트": [], "빼는 칸": [], "이름 다른 행": {}}


# ── 값 비교 ─────────────────────────────────────────────────────────

def compare(v, w: float | None, unit: int, label: str, wp_unit: int = 1) -> str | None:
    """DSD 값 v 와 조서 값 w → '같음' | '반올림' | '부호' | None.
    unit = DSD 표 단위, wp_unit = 조서 시트 '(단위: 천원)' 배수."""
    if w is None:
        return None
    v = float(v)
    cands = [w]
    if unit > 1:
        cands.append(w / unit)                    # 조서 원 단위, DSD 천원
    if wp_unit != unit:
        cands.append(w * wp_unit / unit)          # 조서 천원, DSD 원 (또는 백만원)
    scaled = unit > 1 or wp_unit > 1
    if unit == 1 and wp_unit == 1 and abs(v) >= 100 and w == int(w) and abs(w) >= 100_000:
        cands.append(w / 1000)                    # DSD 표에 단위 표시가 없고 행 이름에만 '(단위: 천원)'
        scaled = True
    for x in cands:
        if x == v:
            return "같음"
        if v and (abs(x - v) <= (1 if scaled else 0) or (v == int(v) and round(x) == v)):
            return "반올림"
    if v != int(v) or abs(v) <= 100:              # 비율: 27.27 ↔ 0.2727, 4.6 ↔ 0.046
        if w and abs(w * 100 - v) <= 0.006:
            return "같음"
    for x in cands:
        if abs(x) == abs(v) or (scaled and abs(abs(x) - abs(v)) <= 1):
            return "부호"
    return None


# ── DSD 칸 순회 ─────────────────────────────────────────────────────

@dataclass
class Item:
    sheet: str           # DSD 구획 (재무상태표 / 주석4 …)
    table: str
    row_id: str
    grid_col: int
    label: str
    col: str
    period: str          # 당기 | 전기
    dsd: float
    wp: float | None
    verdict: str         # 일치 | 불일치 | 부호 다름 | 상대 없음 | 대사 대상 아님
    how: str = ""        # 같은 행 · 하위 합 · 다른 행 · 반올림 · 부호 표기
    where: str = ""      # 조서 위치 '시트!B12'
    wp_label: str = ""
    unit: int = 1        # DSD 표 단위 (원=1, 천원=1000)

    @property
    def wp_in_unit(self) -> float | None:
        """조서 값을 DSD 단위로 (조서가 원 단위면 나눔)."""
        if self.wp is None:
            return None
        if self.unit > 1 and abs(self.wp / self.unit - self.dsd) < abs(self.wp - self.dsd):
            return self.wp / self.unit
        return self.wp


@dataclass
class ReconReport:
    period: str = ""
    items: list[Item] = field(default_factory=list)
    zero_skipped: int = 0
    warnings: list[str] = field(default_factory=list)

    def count(self, verdict: str) -> int:
        return sum(1 for x in self.items if x.verdict == verdict)


def _table_marker(rec: TableRec) -> str | None:
    tv = rec.tv
    rows = (tv.head or []) + tv.body[:1]
    for i in rows:
        t = norm(tv.text(i, 0))
        if not t or t.startswith("(단위"):
            continue
        if _CUR_MARK_RE.search(t) or t.startswith("당"):
            return "당기"
        if _PREV_MARK_RE.search(t) or t.startswith("전"):
            return "전기"
    return None


def _cell_periods(rec: TableRec, year: int | None = None) -> dict[tuple[int, int], str]:
    """(행, 열) → 당기/전기. 열 머리글 > 표 위 캡션 > 표 왼쪽 위 칸 > 행 라벨 순.
    행 라벨은 '당기초'·'전분기말' 같은 말 또는 '2025.01.01(기초자본)'의 연도로 본다."""
    tv = rec.tv
    out: dict[tuple[int, int], str] = {}
    numeric = {c for i in tv.body for c in range(tv.ncols) if tv.value(i, c) is not None}
    acct = {c for c in range(tv.ncols) if _ACCT_HDR_RE.search(norm(tv.col_header(c)))}
    cur = [c for c in rec.cur if c in numeric and c not in acct]
    prev = [c for c in rec.prev if c in numeric and c not in acct]
    table_mark = None
    if cur and not prev:                # '전기초 | 지분법손익 | 전기말' — 머리글 없는 열도 같은 기간
        table_mark = "당기"
    elif prev and not cur:
        table_mark = "전기"
    elif not cur and not prev:
        if rec.caption:
            table_mark = "당기" if rec.caption_is_cur else "전기"
        elif (lc := [c for c in rec.cur if c not in acct]) or [c for c in rec.prev if c not in acct]:
            table_mark = "당기" if lc else "전기"   # 왼쪽 위 칸이 '전기말'인 세로형 표
        else:
            table_mark = _table_marker(rec)
    row_mark = None
    for i in tv.body:
        lab = norm(tv.row_label(i))
        dm = _DATE_RE.search(lab)
        if dm and year:
            row_mark = "당기" if int(dm.group(1)) >= year else "전기"
        elif _ROW_PREV_RE.search(lab):
            row_mark = "전기"
        elif _ROW_CUR_RE.search(lab):
            row_mark = "당기"
        for c in range(tv.ncols):
            if c in cur:
                out[(i, c)] = "당기"
            elif c in prev:
                out[(i, c)] = "전기"
            else:
                out[(i, c)] = row_mark or table_mark or "당기"
    return out


def _target_date(per: P.Period | None, period: str, rec: TableRec):
    if per is None:
        return None
    y, m, d = per.end
    if period == "당기":
        return (y, m, d)
    if per.is_annual or rec.is_balance:
        return (y - 1, 12, 31)
    return (y - 1, m, d)


def _match_block(sheets: list[WpSheet], note_no: int | None, title: str) -> list[tuple[WpSheet, int, int]]:
    """DSD 주석 → 조서 주석 구획(들). 제목 유사도 우선, 없으면 번호.
    주석마다 시트를 두고 한 주석을 여러 시트로 나눈 경우('18. 차입금(1),(2)' · '18. 차입금 (3)')
    가장 비슷한 시트와 번호가 같은 시트를 함께 본다."""
    key = _canon(title)
    best, best_r = [], 0.0
    by_no = []
    for sh in sheets:
        if sh.kind != "note":
            continue
        for n, k, s, e in sh.blocks:
            r = difflib.SequenceMatcher(None, key, k).ratio() if key and k else 0.0
            if r > best_r + 1e-9:
                best, best_r = [(sh, n, s, e)], r
            elif abs(r - best_r) < 1e-9 and r > 0:
                best.append((sh, n, s, e))
            if note_no is not None and n == note_no:
                by_no.append((sh, s, e))
    if best_r >= 0.6:
        nos = {n for sh, n, _s, _e in best if n is not None and NOTE_EACH_RE.match(sh.name)}
        out = [(sh, s, e) for sh, _n, s, e in best]
        for sh in sheets:
            if sh.kind == "note" and NOTE_EACH_RE.match(sh.name):
                out += [(sh, s, e) for n, _k, s, e in sh.blocks if n in nos and (sh, s, e) not in out]
        return out
    return by_no


def rec_scope(rec: TableRec, sheets: list[WpSheet], fs_rows: list[WpRow]) -> tuple[list[WpRow], str]:
    """DSD 표 → 조서에서 찾을 범위(행들)와 구획 이름(재무상태표 · 주석N)."""
    if rec.kind == "NOTE":
        scope = []
        for sh, s, e in _match_block(sheets, rec.sec.note_no, rec.sec.note_title or ""):
            scope += [r for r in sh.rows if s <= r.row <= e]
        return scope, (f"주석{rec.sec.note_no}" if rec.sec.note_no else "주석")
    mine = [sh for sh in sheets if sh.kind == "fs" and sh.fs_kind in (rec.kind, "FS")]
    return ([r for sh in mine for r in sh.rows] if mine else fs_rows), rec.sec.title


def common_keys(scope: list[WpRow]) -> set[str]:
    """범위 안 4행 이상에 나오는 이름 키 = 흔한 키('종속기업' · '기타')."""
    freq: dict[str, int] = {}
    for r in scope:
        for k in r.keys:
            freq[k] = freq.get(k, 0) + 1
    return {k for k, n in freq.items() if n >= 4}


def table_strict(rec: TableRec) -> bool:
    """기간이 없는 이름 있는 열이 여럿인 표(범주별 · 변동표 · 자본 구성요소)."""
    tv = rec.tv
    named = {frozenset(k) for c in range(tv.ncols)
             if c not in rec.cur and c not in rec.prev and not _ACCT_HDR_RE.search(norm(tv.col_header(c)))
             and (k := _name_keys(_leaf(tv, c)))}
    return len(named) >= 2


def col_keys_of(rec: TableRec, c: int) -> tuple[set[str] | None, bool]:
    """DSD 열 이름 키(머리글 모든 층 - '매출 등/기타')와 '기간이 없는 열' 여부."""
    parts = rec.tv.col_header(c).split("/")
    marks = not (rec.cur and rec.prev)             # 기초 · 기말 이름은 한 기간 표('당기초 | … | 당분기말')에서만
    keys: set[str] = set(_name_keys(parts[-1], marks))
    for part in parts[:-1]:                        # 위 층은 '^' 표시
        keys |= {"^" + k for k in _name_keys(part, marks) if k not in keys}
    neutral = c not in rec.cur and c not in rec.prev and not _ACCT_HDR_RE.search(norm(rec.tv.col_header(c)))
    if not _leafk(keys):
        return None, neutral
    return keys, neutral


def same_name_rows(keys: set[str], own: set[str] | None, scope: list[WpRow], common: set[str],
                   col_keys: set[str] | None = None) -> list[WpRow]:
    """같은 이름 행 — 시트마다 키 겹침 최대인 행들. 병합 칸(세로 병합된 회사명 등)을 뺀
    그 행 자신의 이름으로도 한 번 더 고른다 ('(주)갑 | 기초' 아래 '증감' 행)."""
    def best_rows(ks: set[str]) -> list[WpRow]:
        best: dict[str, int] = {}
        by_sheet: dict[str, list[WpRow]] = {}
        spec = ks - common
        for row in scope:
            inter = ks & row.keys
            if spec and not (inter & spec):    # 범위에 흔한 키('종속기업')만 겹침
                continue
            n = len(inter)
            if n and n >= best.get(row.sheet, 0):
                if n > best.get(row.sheet, 0):
                    best[row.sheet], by_sheet[row.sheet] = n, []
                by_sheet[row.sheet].append(row)
        return [r for rows in by_sheet.values() for r in rows]

    cands = best_rows(keys)
    if own and own != keys:
        cands += [r for r in best_rows(own) if r not in cands]
    if col_keys and len(cands) > 1:            # 같은 회사 · 계정 이름이 여러 표에 — DSD 열 이름이 있는 표
        pref = [r for r in cands if any(_names_hit(r.names.get(c, set()), col_keys, True) for c in r.nums)]
        pref = pref or [r for r in cands if any(_names_hit(r.names.get(c, set()), col_keys) for c in r.nums)]
        if pref:
            cands = pref
    return cands


def row_keys(tv, i: int) -> tuple[str, set[str], set[str]]:
    """DSD 행 → (이름 글, 이름 키, 병합 칸을 뺀 자기 이름 키)."""
    texts = [tv.text(i, c) for c in range(tv.ncols) if tv.text(i, c).strip() and tv.value(i, c) is None]
    own = label_keys([tv.cell(i, c).text for c in range(tv.ncols)
                      if tv.cell(i, c) is not None and tv.value(i, c) is None])
    return " ".join(t.strip() for t in texts), label_keys(texts), own


@dataclass
class ReconConfig:
    """건별 설정(_engagement/dsd_recon.json 의 '별도' · '연결' 한쪽)."""
    exclude: list[dict] = field(default_factory=list)      # [{'구획': '주석9', '행': '합계', '열': '당분기'}]
    aliases: dict[str, str] = field(default_factory=dict)  # DSD 행 이름 → 조서 행 이름

    @classmethod
    def from_dict(cls, d: dict | None) -> "ReconConfig":
        d = d or {}
        return cls(list(d.get("빼는 칸") or []), dict(d.get("이름 다른 행") or {}))

    def excluded(self, item: "Item") -> bool:
        for x in self.exclude:
            if x.get("구획") and x["구획"] != item.sheet:
                continue
            if x.get("행") and norm(x["행"]) not in norm(item.label):
                continue
            if x.get("열") and norm(x["열"]) not in norm(item.col):
                continue
            if x.get("기간") and x["기간"] != item.period:
                continue
            return True
        return False

    def alias_keys(self, keys: set[str]) -> set[str]:
        extra: set[str] = set()
        for src, dst in self.aliases.items():
            if label_keys([src]) & keys:
                extra |= label_keys([dst])
        return extra


def reconcile(doc: Document, sheets: list[WpSheet], include_prior: bool = False,
              config: ReconConfig | None = None) -> ReconReport:
    per = P.detect_period(doc)
    cfg = config or ReconConfig()
    rep = ReconReport(period=per.label if per else "(기간 미상)")
    if per is None:
        rep.warnings.append("DSD 보고기간을 못 읽어 조서 머리글 날짜를 쓰지 않고 행 안에서만 찾음")
    elif sheets and per.consolidated != any(s.consolidated for s in sheets):
        rep.warnings.append("DSD 는 %s 재무제표인데 조서 시트 제목은 %s - 조서 파일을 확인"
                            % (("연결", "별도") if per.consolidated else ("별도", "연결")))
    one_q = bool(per and per.quarter == 1)    # 1분기는 3개월 = 누적
    fs_rows = [r for sh in sheets if sh.kind == "fs" for r in sh.rows]
    if sheets and not any(r.nums for sh in sheets for r in sh.rows):
        rep.warnings.append("조서에서 숫자를 못 읽음 — 계산값이 없는 파일이면 엑셀에서 열어 저장한 뒤 다시 실행")

    recs, _ = scan(doc)
    for rec in recs:
        if rec.kind in ("OTHER", "RE"):
            continue
        tv = rec.tv
        scope, ctx = rec_scope(rec, sheets, fs_rows)
        periods = _cell_periods(rec, per.year if per else None)
        unit = tv.unit
        common = common_keys(scope)
        strict = table_strict(rec)
        has_cols: dict[tuple, bool] = {}          # (기간, 3개월) → 범위에 그 기간 열이 있나
        for i in tv.body:
            label, keys, own = row_keys(tv, i)
            keys |= cfg.alias_keys(keys)
            for c in range(tv.ncols):
                v = tv.value(i, c)
                if v is None or _SKIP_HDR_RE.search(norm(tv.col_header(c))):
                    continue
                period = periods[(i, c)]
                target = _target_date(per, period, rec)
                is3m = bool(_3M_RE.search(_leaf(tv, c)))
                item = Item(ctx, tv.table.id, tv.table.rows[i].id, c, label[:40],
                            tv.col_header(c)[:30], period, float(v), None, "", unit=unit)
                if cfg.exclude and cfg.excluded(item):
                    item.verdict = "제외(설정)"
                    rep.items.append(item)
                    continue
                if not scope:
                    if v != 0:
                        item.verdict = "조서에 없음"
                        rep.items.append(item)
                    continue
                # DSD 열 이름 — 기간이 없는 열('연이자율(%)', 변동표 '취득')은 이름으로 조서 열을 찾고,
                # 기간 열('전기말/총장부금액')은 기간이 맞는 열 중 같은 이름 열을 먼저 본다
                col_keys, neutral = col_keys_of(rec, c)
                hk = (period, is3m, target, frozenset(col_keys or ()), neutral)
                if hk not in has_cols:
                    has_cols[hk] = any(_row_cols(r, target, period, is3m, one_q, col_keys, False,
                                                 neutral, strict)[0] for r in scope)
                _judge(item, v, keys, scope, target, period, is3m, one_q, has_cols[hk], unit,
                       include_prior or period == "당기", rec.kind, own, col_keys, neutral, strict,
                       common)
                if item.verdict == "0 생략":
                    rep.zero_skipped += 1
                    continue
                rep.items.append(item)
    return rep


def _row_period(row: WpRow) -> str | None:
    """행 전체의 기간 — 표 위 소제목('① 당기말') 또는 행 이름('전기초' · '당분기말')."""
    if row.sect:
        return row.sect
    lab = norm(row.label)
    return "전기" if _ROW_PREV_RE.search(lab) else "당기" if _ROW_CUR_RE.search(lab) else None


def _row_cols(row: WpRow, target, period: str, is3m: bool, one_q: bool,
              col_keys: set[str] | None = None, loose: bool = True,
              neutral: bool = True, strict: bool = False) -> tuple[list[int], bool]:
    """그 행에서 DSD 칸(기간 · 3개월)에 맞는 조서 열과, 행에 기간 정보가 있는지.
    열 머리글의 날짜가 있으면 날짜로, 없으면 '당기 · 전기' 말로. 열이 기간을 안 주면 표 위 소제목.
    col_keys = DSD 열 이름('연이자율(%)' · '감가상각' · '총장부금액') — 기간이 맞는 열 중 같은 이름 열을 먼저,
    기간이 없는 DSD 열(neutral)은 조서에서 같은 이름 열을 쓴다.
    loose = 행 기간을 모를 때도 이름 대조를 쓸지 (당기 칸 · --prior 만).
    strict = DSD 표가 이름 있는 열 여럿(범주 · 변동 항목)이라 같은 이름 열이 없으면 다른 표로 본다."""

    def narrow(cols: list[int]) -> list[int]:
        if col_keys:
            named = _named(row, cols, col_keys)
            if named:
                return named
        return cols

    # 값이 있는 열에 걸린 기간만 ('당기말 | 예금종류 | 금액' 처럼 이름 칸 위 기간 말은 표 전체 표시라 무시)
    info = [c for c, (d, p, _t3) in row.cols.items() if (d or p) and c in row.nums]

    def period_ok(c: int) -> bool:
        d, p, t3 = row.cols[c]
        if not one_q and t3 != is3m:
            return False
        return ((d is not None and target is not None and (d == target or (d[1] == 0 and d[0] == target[0])))
                or (d is None and p == period))

    rp = _row_period(row)
    if col_keys and neutral:
        named = [c for c in _named(row, sorted(row.names), col_keys) if c in row.nums]
        if not named and (any(_names_hit(row.names.get(c, set()), col_keys) for c in row.names)
                          or any(row.names.get(c) for c in row.nums)):
            return [], True                     # 이름이 맞는 열이 이 행에서 빈 칸 / 값 열마다 이름이 있는데 안 맞음(다른 표)
        ok = [c for c in named if (period_ok(c) if c in info
                                   else rp == period or (rp is None and loose))]
        if ok:
            return sorted(ok), True
        if named or strict:                    # 이름 열은 있는데 기간이 다름 / 이름 열이 없는 다른 표
            return [], True
    if info:
        return sorted(narrow([c for c in info if period_ok(c)])), True
    if rp:
        return (sorted(narrow(list(row.nums))) if rp == period else []), True
    return [], False


def _ref(row: WpRow, col: int) -> str:
    from openpyxl.utils import get_column_letter
    return f"{row.sheet}!{get_column_letter(col + 1)}{row.row}"


def _judge(item: Item, v, keys: set[str], scope: list[WpRow], target, period: str, is3m: bool,
           one_q: bool, any_dated: bool, unit: int, search_undated: bool, kind: str = "",
           own: set[str] | None = None, col_keys: set[str] | None = None,
           neutral: bool = True, strict: bool = False, common: set[str] = frozenset()) -> None:
    def val_cols(row: WpRow) -> tuple[list[int], bool]:
        cols, info = _row_cols(row, target, period, is3m, one_q, col_keys, search_undated, neutral,
                               strict)
        if info:
            return cols, True
        return (sorted(row.nums), False) if search_undated else ([], False)

    def val(row: WpRow, col: int) -> float | None:
        return row.nums.get(col)

    def finish(verdict, how, row=None, col=None, w=None):
        item.verdict, item.how = verdict, how
        if row is not None:
            item.where = _ref(row, col) if col is not None else f"{row.sheet}!{row.row}행"
            item.wp_label = row.label[:40]
        item.wp = w
        if verdict == "일치" and how.endswith("부호"):
            lab = norm(item.label)
            ok = _LOSS_RE.search(lab) or (kind in ("CF", "SCE") and _CF_DIR_RE.search(lab))
            item.verdict = "일치" if ok else "부호 다름"
            item.how = how.replace("부호", "부호 표기")

    # 1) 같은 이름 행
    cands = same_name_rows(keys, own, scope, common, col_keys)
    for row in cands:
        cols, _ = val_cols(row)
        for col in cols:
            k = compare(v, val(row, col), unit, item.label, row.unit)
            if k:
                return finish("일치", "같은 행" + ("" if k == "같음" else f" · {k}"), row, col,
                              val(row, col))

    if not search_undated:
        # 같은 이름 행이 있으면 그 행들에, 없으면 범위 전체에 그 기간 열이 있는지 (연결 주석 시트는 당기만)
        dated = any(val_cols(r)[0] for r in cands) if cands else any_dated
        if not dated:
            item.verdict = "대사 대상 아님"
            return

    if v == 0:
        item.verdict = "0 생략"
        for row in cands:
            cols, dated = val_cols(row)
            w = next((val(row, c) for c in cols if c in row.nums), None) if dated else None
            if w:
                return finish("불일치", "같은 행", row, cols[0], w)
        return

    # 2) 하위 합 — 같은 이름 행이 여럿이면 기간 열 합
    if len(cands) > 1:
        tot, used = 0.0, []
        for row in cands:
            cols, dated = val_cols(row)
            if dated and cols and cols[0] in row.nums:
                tot += val(row, cols[0])
                used.append(row)
        if len(used) > 1 and (k := compare(v, tot, unit, item.label, used[0].unit)):
            return finish("일치", f"하위 합 {len(used)}행" + ("" if k == "같음" else f" · {k}"),
                          used[0], None, tot)

    # 3) 같은 이름 행은 있는데 값이 다름 — 가장 가까운 값을 보여 줌
    near = []
    for row in cands:
        cols, _ = val_cols(row)
        for c in cols:
            if c in row.nums:
                w = val(row, c)
                gap = min(abs(w - float(v)), abs(w / unit - float(v)) if unit > 1 else abs(w - float(v)))
                near.append((gap, row, c, w))
    if near:
        _, row, c, w = min(near, key=lambda x: x[0])
        return finish("불일치", "같은 행", row, c, w)
    # 4) 같은 이름 행이 없거나 비어 있을 때만 — 같은 범위 다른 행
    for row in scope:
        if row in cands:
            continue
        cols, _ = val_cols(row)
        for col in cols:
            k = compare(v, val(row, col), unit, item.label, row.unit)
            if k in ("같음", "반올림") or (k == "부호" and kind in ("CF", "SCE")
                                         and _CF_DIR_RE.search(norm(item.label))):
                return finish("일치", "다른 행" + ("" if k == "같음" else f" · {k}"), row, col,
                              val(row, col))

    if cands:
        return finish("상대 없음", "조서 칸 빈칸", cands[0])
    finish("상대 없음", "")


# ── 결과 엑셀 ───────────────────────────────────────────────────────

VERDICTS = ["일치", "불일치", "부호 다름", "상대 없음", "조서에 없음", "대사 대상 아님", "제외(설정)"]
ISSUE_VERDICTS = ("불일치", "부호 다름", "상대 없음")


def summary_rows(rep: ReconReport) -> list[tuple[str, dict[str, int]]]:
    order: list[str] = []
    counts: dict[str, dict[str, int]] = {}
    for x in rep.items:
        if x.sheet not in counts:
            order.append(x.sheet)
            counts[x.sheet] = {v: 0 for v in VERDICTS}
        counts[x.sheet][x.verdict] += 1
    return [(k, counts[k]) for k in order]


def write_recon_report(rep: ReconReport, out_path: str, dsd_path: str, wp_path: str) -> dict:
    import os

    from openpyxl import Workbook
    from openpyxl.styles import Alignment

    from .report_xlsx import (
        _BORDER, _FILL_BAD, _FONT_BAD, _NUMFMT, FONT, FONT_LINK, FONT_TITLE, _write_header,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    ws.cell(row=2, column=2, value="DSD ↔ 조서 대사 결과").font = FONT_TITLE
    info = [("DSD", os.path.basename(dsd_path)), ("조서", os.path.basename(wp_path)),
            ("보고기간", rep.period), ("0 칸 생략", rep.zero_skipped)]
    r = 4
    for k, v in info:
        ws.cell(row=r, column=2, value=k).font = FONT
        ws.cell(row=r, column=3, value=v).font = FONT
        r += 1
    r += 1
    _write_header(ws, r, ["", "구획"] + VERDICTS)
    ws.cell(row=r, column=1).value = None
    r += 1
    for sheet, cnt in summary_rows(rep):
        vals = [sheet] + [cnt[v] for v in VERDICTS]
        for c, v in enumerate(vals, start=2):
            cell = ws.cell(row=r, column=c, value=v)
            bad = c - 2 in (2, 3) and isinstance(v, int) and v > 0
            cell.font = _FONT_BAD if bad else FONT
            cell.border = _BORDER
            if c > 2:
                cell.alignment = Alignment(horizontal="center")
        r += 1
    r += 1
    for line in ["일치: 조서에서 같은 값을 찾음(찾은 방법 열 참고) / 불일치: 같은 이름 행은 있는데 값이 다름",
                 "부호 다름: 크기만 같음 / 상대 없음: 조서에 같은 이름 행이 없거나 칸이 비어 있음",
                 "조서에 없음: 조서에 그 주석 구획이 없음 / 대사 대상 아님: 조서에 그 기간 열이 없음(주로 전기)",
                 *rep.warnings]:
        ws.cell(row=r, column=2, value=line).font = FONT
        r += 1
    ws.column_dimensions["B"].width = 18
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 12

    heads = ["구획", "표", "행", "열", "기간", "DSD 값", "조서 값", "차이", "결과", "찾은 방법",
             "조서 위치", "조서 행"]
    widths = [12, 7, 30, 22, 6, 18, 18, 14, 10, 18, 16, 30]
    wp_abs = os.path.abspath(wp_path)
    n_issue = 0
    for title, items in (("확인 필요", [x for x in rep.items if x.verdict in ISSUE_VERDICTS]),
                         ("대사결과", rep.items)):
        ws = wb.create_sheet(title)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[chr(64 + i)].width = w
        _write_header(ws, 1, heads)
        ws.freeze_panes = "A2"
        rr = 2
        for x in items:
            wpv = x.wp_in_unit
            diff = (x.dsd - wpv) if wpv is not None and x.verdict == "불일치" else None
            vals = [x.sheet, x.table, x.label, x.col.replace("\n", " "), x.period, x.dsd, wpv,
                    diff, x.verdict, x.how, x.where, x.wp_label]
            issue = x.verdict in ISSUE_VERDICTS
            for c, v in enumerate(vals, start=1):
                cell = ws.cell(row=rr, column=c, value=v)
                cell.font = _FONT_BAD if issue and c == 9 else FONT
                cell.border = _BORDER
                if c in (6, 7, 8) and isinstance(v, (int, float)):
                    cell.number_format = _NUMFMT if float(v) == int(v) else "#,##0.00"
                if issue:
                    cell.fill = _FILL_BAD
            if "!" in x.where and not x.where.endswith("행"):
                sh, addr = x.where.split("!", 1)
                link = ws.cell(row=rr, column=11)
                link.hyperlink = f"{wp_abs}#'{sh}'!{addr}"
                link.font = FONT_LINK
            rr += 1
        ws.auto_filter.ref = f"A1:L{max(rr - 1, 1)}"
        if title == "확인 필요":
            n_issue = len(items)
    from .export_xlsx import _T
    if _T.get("modernize"):
        _T["modernize"](wb)
    wb.properties.creator = "이재현"
    wb.save(out_path)
    return {"items": len(rep.items), "issues": n_issue}
