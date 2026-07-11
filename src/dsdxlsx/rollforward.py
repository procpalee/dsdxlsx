"""기간 롤포워드 — 전기(직전 보고서) DSD에서 당기 표준 정산표를 만든다.

'결산 시즌 시작 버튼': 전기 DSD 하나로 ① 기수 라벨 +1 ② TU 자동필드 연도 +1
(스파이크 V4 검증 방식) ③ 각 표의 당기 열 값 → 전기 열 이동 + 당기 열 클리어
④ THEAD의 기간 날짜 +1년 을 수행하고, 결과를 3등급 리포트로 표면화한다.

등급 원칙 (플랜 확정):
- 자동확정: 결정론 치환 — 기수 라벨, TU, 버킷 페어가 명확한 표의 값 이동.
- 확인필요: 건드리지 않았지만 사람이 봐야 하는 것 — 문단·본문 셀의 날짜,
  meta 문서정보(편집기 소관), 클리어된 당기 열(결산 수치 입력 대상).
- 미처리: 구조를 확신할 수 없어 **절대 무변경**한 표 — 블록 경계 인식 실패
  자본변동표, 당/전 열 대응이 없는 표, 당기 열만 있는 표.

보고서 종류 전환(1분기→반기 등)은 범위 밖 — 날짜는 일괄 +1년(연차·동일분기 롤).
"""
from __future__ import annotations

import bisect
import re
from dataclasses import dataclass, field

from . import splice
from .check import TableView, build_views, norm, use_profile
from .model import Document

_GISU_RE = re.compile(r"(제\s*)(\d{1,3})(\s*\(?[당전]?\)?\s*기)")
_DATE_RE = re.compile(r"(19|20)\d{2}(?=\s*[년.\-/])")
_YEAR_RE = re.compile(rb"(19|20)\d{2}")
# 세로형 구분 문단의 마지막 줄: "1) 당기말" / "2) 전기말" 류
_CUR_MARK_RE = re.compile(r"당\s*(분|반)?\s*기\s*말?\s*$")
_PREV_MARK_RE = re.compile(r"전\s*(분|반)?\s*기\s*말?\s*$")
# '② 전기 중 …변동은 없습니다.' 류 — 앞부분(마커로 남길 '② 전기')을 캡처
_PREV_PROSE_HEAD_RE = re.compile(r"^([\s0-9①-⑳()\.\-]*전(?:분|반)?기(?:\s*말)?)")
# 자본변동표 블록 라벨 플립: 구 당기 블록이 신 전기 블록이 되므로
# (긴 패턴 먼저 — '당분기말'이 '당기말'보다 앞서 치환되도록)
_SOCIE_FLIP = [("당분기말", "전분기말"), ("당반기말", "전반기말"),
               ("당기초", "전기초"), ("당기말", "전기말")]
# 주석 안 행블록 변동표(<전기>/<당기>) 인식 — 명시적 전/당 접두 라벨만
_PREV_OPEN_RE = re.compile(r"전(분|반)?기초")
_PREV_CLOSE_RE = re.compile(r"전(분|반)?기말")
_CUR_OPEN_RE = re.compile(r"당(분|반)?기초")
_CUR_CLOSE_RE = re.compile(r"당(분|반)?기말")
# 당기 전용 열 헤더 → 기대되는 전기 상대 헤더 ("당기말" → "전기말")
_CUR_IN_HEADER_RE = re.compile(r"당(분|반)?기")

_CR_SEP = b"&amp;cr;"  # P inner 원시 바이트의 줄 구분자 (escape_text 관례)


@dataclass
class RollItem:
    grade: str      # 자동확정 | 확인필요 | 미처리
    where: str
    item: str
    before: str = ""
    after: str = ""
    note: str = ""


@dataclass
class RollReport:
    items: list[RollItem] = field(default_factory=list)
    # 자동 이동 표시(색칠) 대상 — roll_ops 종료 시 롤 적용 후 바이트 위치로 확정:
    # ("cells",[(row_start,gc)..]) | ("table",table_start) | ("socie",table_start)
    painted: list[tuple] = field(default_factory=list)

    def add(self, grade, where, item, before="", after="", note=""):
        self.items.append(RollItem(grade, where, item, str(before)[:40],
                                   str(after)[:40], note))

    def count(self, grade: str) -> int:
        return sum(1 for i in self.items if i.grade == grade)


def _gisu_mapping(doc: Document) -> dict[int, int]:
    """문서에 등장하는 기수 중 상위 3개만 +1 — '제 1 기'류 서술(설립 연혁 등)
    오치환 방지."""
    nums = set()
    for node in doc.nodes.values():
        for m in _GISU_RE.finditer(node.text):
            nums.add(int(m.group(2)))
    top = sorted(nums, reverse=True)[:3]
    return {n: n + 1 for n in top}


def _period_years(doc: Document) -> set[int]:
    """보고기간 연도 집합 — TU 자동필드(기계값 YYYYMMDD)에서 추출한 최신 2개
    연도만. 이 연도의 날짜만 +1년 대상 — 설립·상장 등 역사적 연도 오염 방지."""
    years: set[int] = set()
    for node in doc.nodes.values():
        if node.tag == "TU":
            for m in _YEAR_RE.finditer(doc.data[node.start:node.end]):
                years.add(int(m.group()))
    if not years:  # TU 없는 문서: 본문 날짜 중 최신 연도로 추정
        for node in doc.nodes.values():
            for m in _DATE_RE.finditer(node.text):
                years.add(int(m.group()))
    if not years:
        return set()
    top = max(years)
    return {top, top - 1}  # 당기말·전기말 연도만 — 그 이전(역사적) 연도는 제외


def _bump_year(text: str, allowed: set[int]) -> str:
    return _DATE_RE.sub(
        lambda m: str(int(m.group()) + 1) if int(m.group()) in allowed else m.group(),
        text)


def _roll_text(text: str, mapping: dict[int, int], allowed: set[int]) -> str:
    """기수 라벨 +1, 기수가 있던 텍스트에 한해 보고기간 연도 +1 (② 규칙)."""
    new = _GISU_RE.sub(
        lambda m: (m.group(1)
                   + str(mapping.get(int(m.group(2)), int(m.group(2))))
                   + m.group(3)),
        text)
    if new != text:
        new = _bump_year(new, allowed)
    return new


def _numeric_cells(tv: TableView, i: int):
    """행의 값 성격 셀(숫자·'-' 플레이스홀더)만 — 라벨·주석 텍스트는 제외."""
    from .check import cell_value
    for c in range(tv.ncols):
        cell = tv.cell(i, c)
        if cell is None or cell.tag == "TU":
            continue
        t = cell.text.strip()
        if t and (t == "-" or cell_value(cell.text) is not None):
            yield c, cell


# ── P 원시 바이트 ↔ 표시 줄 대응 (문단 분할용) ──────────────────────


def _p_segments(doc: Document, node) -> tuple[bytes, list[bytes]] | None:
    """P inner 원시 바이트를 표시 텍스트 줄과 1:1 세그먼트로 분해.

    재구성 검증(세그먼트 조합 == node.text)에 실패하면 None —
    자식 요소·수치참조가 섞인 복합 문단은 자동 분할하지 않는다."""
    raw = doc.data[node.inner_start:node.inner_end]
    segs = raw.split(_CR_SEP)

    def display(seg: bytes) -> str:
        s = seg.decode("utf-8", "replace")
        s = s.replace("\r\n", "").replace("\n", "").replace("\r", "")
        return (s.replace("&lt;", "<").replace("&gt;", ">")
                .replace("&quot;", '"').replace("&amp;", "&"))

    if "\n".join(display(s) for s in segs) != node.text:
        return None
    return raw, segs


def _seg_offset(segs: list[bytes], k: int) -> int:
    return sum(len(s) for s in segs[:k]) + k * len(_CR_SEP)


def _split_mixed_headings(doc: Document, report: RollReport, ops: list,
                          touched: set, mapping: dict, allowed_years: set):
    """혼합 문단 분리: 앞 주석 꼬리에 '20. 제목' 헤딩이 병합된 P를 헤딩 경계에서
    두 P로 나눈다 — 여는 태그·원시 바이트를 그대로 재사용하는 클론 분할이라
    편집기 수용 방식(clone-only)을 유지한다. 롤 산출물의 주석 구획·시트가
    온전해지는 물리적 교정."""
    from .export_xlsx import _note_heading, refined_sections

    for sec in refined_sections(doc):
        if sec.note_no is None and not sec.notes_intro:
            continue
        base_n = sec.note_no or 0
        for block in sec.blocks:
            if block.kind != "p":
                continue
            node = block.item
            if node.readonly or node.id in touched:
                continue
            lines = node.text.split("\n")
            cuts, n = [], base_n
            for k, line in enumerate(lines):
                if k == 0:
                    continue
                h = _note_heading(line, n)
                if h:
                    cuts.append(k)
                    n = h[0]
            if not cuts:
                continue
            seg = _p_segments(doc, node)
            if seg is None:
                report.add("확인필요", node.id,
                           "주석 헤딩이 복합 문단에 병합 — 자동 분리 불가",
                           before=lines[cuts[0]].strip()[:36],
                           note="편집기에서 헤딩을 별도 문단으로 분리하세요")
                continue
            raw, segs = seg
            bounds = [0] + [_seg_offset(segs, k) for k in cuts] + [len(raw)]
            open_tag = doc.data[node.start:node.inner_start]
            close_tag = doc.data[node.inner_end:node.end]
            parts = []
            for a, b in zip(bounds, bounds[1:]):
                text = raw[a:b].decode("utf-8")
                parts.append(open_tag
                             + _roll_text(text, mapping, allowed_years).encode()
                             + close_tag)
            ops.append(splice.Op(pos=node.start, start=node.start,
                                 end=node.end, payload=b"\r\n".join(parts)))
            touched.add(node.id)
            report.add("자동확정", node.id,
                       f"문단에 병합된 주석 헤딩 {len(cuts)}곳을 별도 문단으로 분리",
                       note="; ".join(lines[k].strip()[:24] for k in cuts))
            for m in _DATE_RE.finditer(node.text):
                ctx = node.text[max(0, m.start() - 12):m.end() + 10]
                report.add("확인필요", node.id, "날짜 발견 (자동 치환 안 함)",
                           before=ctx.replace("\n", " ").strip(),
                           note="보고기간 표기면 +1년으로 직접 수정")
                break


# ── 세로형 표 쌍 / '전기 변동 없음' 문단 ────────────────────────────


def _clear_cur_table(doc: Document, sec_title: str, cur_t, ops: list,
                     touched: set) -> int:
    cur_tv = TableView(cur_t, sec_title)
    n_clear = 0
    for i in cur_tv.body:
        for _, cell in _numeric_cells(cur_tv, i):
            ops.append(splice.replace_inner(cell, ""))
            touched.add(cell.id)
            n_clear += 1
    return n_clear


def _swap_prose_to_table(doc: Document, sec, cur_t, node, first_k: int,
                         report: RollReport, ops: list, touched: set,
                         painted: list, handled: set) -> bool:
    """'② 전기 중 …변동은 없습니다.' 문단 → '② 전기' 문단 + 구 당기 표 클론.

    구 당기에 변동이 있었으므로 신 전기에는 표가 와야 한다. 문단을 마커('② 전기')
    로 줄이고 그 아래 구 당기 표를 통째 복제, 당기 표는 값을 비운 스켈레톤으로."""
    lines = node.text.split("\n")
    line = lines[first_k].strip()
    m = _PREV_PROSE_HEAD_RE.match(line)
    if m is None:
        return False
    marker_text = m.group(1).strip()
    seg = _p_segments(doc, node)
    if seg is None:
        report.add("미처리", _sec_where(sec),
                   f"{cur_t.id} 당기 표의 전기 상대가 '변동 없음' 복합 문단",
                   before=line[:36], note="자동 교체 불가 — 전기 표를 수기 작성")
        return True
    raw, segs = seg
    off_k = _seg_offset(segs, first_k)
    open_tag = doc.data[node.start:node.inner_start]
    close_tag = doc.data[node.inner_end:node.end]
    p1 = (open_tag + raw[:off_k] + splice.escape_text(marker_text) + close_tag)
    clone = doc.data[cur_t.start:cur_t.end]
    payload = p1 + b"\r\n\r\n" + clone
    rest = raw[off_k + len(segs[first_k]):]
    if rest:
        payload += b"\r\n" + open_tag + rest + close_tag
    ops.append(splice.Op(pos=node.start, start=node.start, end=node.end,
                         payload=payload))
    touched.add(node.id)
    n_clear = _clear_cur_table(doc, sec.title, cur_t, ops, touched)
    handled.add(cur_t.id)
    painted.append(("ptable", node.start, len(p1) + 4))
    report.add("자동확정", _sec_where(sec),
               f"'{line[:28]}' 문단 → '{marker_text}' + 표 신설(구 당기 클론), "
               f"당기 표 값 {n_clear}칸 비움",
               note="새 당기 변동이 없으면 당기 표를 지우고 문구로 되돌리세요")
    return True


def _sec_where(sec) -> str:
    return f"주석{sec.note_no}" if sec.note_no else sec.title


def _roll_paired_tables(doc: Document, report: RollReport, ops: list,
                        touched: set, painted: list) -> set[str]:
    """세로형 표 쌍: 'N) 당기말' 문단 아래 표 ↔ 'N) 전기말' 문단 아래 표.

    기수 간 행 구성이 다를 수 있으므로(파생상품 신규 등) 셀 이동이 아니라
    **전기 표를 구 당기 표의 원시 바이트 클론으로 통째 교체**하고, 당기 표의
    값 셀만 비운다 — clone-only 원칙(편집기 수용 검증 방식).
    전기 쪽이 '변동 없음' 문단이면 문단을 마커로 줄이고 표를 신설한다."""
    from .export_xlsx import refined_sections

    handled: set[str] = set()
    for sec in refined_sections(doc):
        pending = None   # (table, 당기마커 문단 요약)
        fresh = False    # pending 직후 블록인가 — 문단 교체는 바로 다음 P만
        last_mark = ""
        for block in sec.blocks:
            if block.kind == "p":
                node = block.item
                stripped = [(k, l.strip()) for k, l in
                            enumerate(node.text.split("\n")) if l.strip()]
                if stripped:
                    first_k, first = stripped[0]
                    if (pending is not None and fresh
                            and node.id not in touched and not node.readonly
                            and "없" in first
                            and _PREV_PROSE_HEAD_RE.match(first)):
                        if _swap_prose_to_table(doc, sec, pending[0], node,
                                                first_k, report, ops, touched,
                                                painted, handled):
                            pending = None
                    last_mark = stripped[-1][1]
                    if not first.startswith("(*"):
                        fresh = False  # 각주(*)·빈 문단만 pending 직후 상태 유지
                continue
            if block.kind != "table":
                fresh = False
                continue
            t = block.item
            tags = {c.tag for r in t.rows for c in r.cells}
            eligible = ("TU" not in tags and not t.fixtable
                        and t.group_aclass is None and not t.in_insertion)
            if _CUR_MARK_RE.search(last_mark) and eligible:
                pending = (t, last_mark)
                fresh = True
            elif (_PREV_MARK_RE.search(last_mark) and eligible
                  and pending is not None and pending[0].ncols == t.ncols):
                cur_t, cur_mark = pending
                clone = doc.data[cur_t.start:cur_t.end]
                ops.append(splice.Op(pos=t.start, start=t.start, end=t.end,
                                     payload=clone))
                for row in t.rows:
                    for cell in row.cells:
                        touched.add(cell.id)
                n_clear = _clear_cur_table(doc, sec.title, cur_t, ops, touched)
                handled.update({cur_t.id, t.id})
                painted.append(("table", t.start))
                report.add("자동확정", _sec_where(sec),
                           f"{t.id}(전기) ← {cur_t.id}(구 당기) 표 교체, "
                           f"당기 표 값 {n_clear}칸 비움",
                           note=f"'{cur_mark}'/'{last_mark}' 쌍 — 각주(*) 문단은 확인")
                pending = None
                fresh = False
            else:
                pending = None
                fresh = False
    return handled


# ── 행 블록 변동표 (자본변동표 + 주석 <전기>/<당기> 블록) ───────────


def _roll_socie(doc: Document, views, report: RollReport, ops: list,
                touched: set, painted: list, allowed_years: set,
                handled_in: set) -> set[str]:
    """행 블록 변동표: [전기초..전기말] 블록 ← [당기초..당기말] 블록 통째 클론
    (라벨의 당기↔전기 플립), 신 당기초 값 ← 구 당기말 값(이월), 당기 변동·기말
    클리어. 자본변동표 본표 외에 주석의 <전기>/<당기> 행 블록 표(자본금 변동
    내용 등)도 명시적 전기초/당기초 라벨이 있으면 같은 방식으로 처리.
    블록 경계가 불명확하면 절대 무변경(미처리)."""
    from .check import norm as _norm

    handled: set[str] = set()
    for tv in views:
        t = tv.table
        if t.id in handled_in:
            continue
        is_socie_sec = "자본변동표" in _norm(tv.sec_title)
        if not is_socie_sec:
            if _skip_table(tv):
                continue
            if (t.fixtable or t.group_aclass is not None or t.in_insertion):
                continue
            # 당/전 열이 있는 표는 열 페어(①) 소관 — 행 블록 오인 방지
            if any(tv.period_bucket(c) is not None for c in range(tv.ncols)):
                continue
        body = tv.body
        opens = [i for i in body if "기초" in _norm(tv.row_label(i))
                 and any(tv.value(i, c) is not None for c in range(tv.ncols))]
        closes = [i for i in body if "기말" in _norm(tv.row_label(i))
                  and any(tv.value(i, c) is not None for c in range(tv.ncols))]
        if len(opens) != 2 or len(closes) != 2:
            continue
        o1, o2 = opens
        c1, c2 = closes
        if not (o1 < c1 < o2 < c2):
            continue
        if not is_socie_sec:
            # 주석 표는 명시적 전/당 접두 라벨(전기초·당기초…)일 때만
            if not (_PREV_OPEN_RE.search(_norm(tv.row_label(o1)))
                    and _PREV_CLOSE_RE.search(_norm(tv.row_label(c1)))
                    and _CUR_OPEN_RE.search(_norm(tv.row_label(o2)))
                    and _CUR_CLOSE_RE.search(_norm(tv.row_label(c2)))):
                continue
        if any(cell.tag == "TU" for i in body for cell in t.rows[i].cells):
            continue
        rows1 = t.rows[o1:c1 + 1]   # 전기 블록
        rows2 = t.rows[o2:c2 + 1]   # 당기 블록
        clone = doc.data[rows2[0].start:rows2[-1].end]
        for a, b in _SOCIE_FLIP:
            clone = clone.replace(a.encode(), b.encode())
        ops.append(splice.Op(pos=rows1[0].start, start=rows1[0].start,
                             end=rows1[-1].end, payload=clone))
        for row in rows1:
            for cell in row.cells:
                touched.add(cell.id)
        # 신 당기초 ← 구 당기말 값 (이월), 나머지 당기 행 값 클리어
        close_vals = {c: tv.cell(c2, c).text for c, _ in _numeric_cells(tv, c2)}
        for c, cell in _numeric_cells(tv, o2):
            new = close_vals.get(c, "")
            if cell.text != new:
                ops.append(splice.replace_inner(cell, new))
            touched.add(cell.id)
        n_clear = 0
        for i in body:
            if o2 < i <= c2:
                for c, cell in _numeric_cells(tv, i):
                    ops.append(splice.replace_inner(cell, ""))
                    touched.add(cell.id)
                    n_clear += 1
        # 당기 블록의 라벨·일자 셀 보고기간 연도 +1
        # ('2025.1.1(당기초)' 라벨, 일자 열 '2025.01.01' 모두 대응)
        for i in body:
            if o2 <= i <= c2:
                for cell in t.rows[i].cells:
                    if cell.tag == "TU" or cell.id in touched:
                        continue
                    new_text = _bump_year(cell.text, allowed_years)
                    if new_text != cell.text:
                        ops.append(splice.replace_inner(cell, new_text))
                        touched.add(cell.id)
        handled.add(t.id)
        painted.append(("socie", t.start))
        report.add("자동확정", tv.context,
                   f"{t.id} 변동표 행 블록 교체 (전기 ← 구 당기, "
                   f"당기초 ← 구 당기말, 변동 {n_clear}칸 비움)",
                   note="당기 실적 행 구성은 결산 내용에 맞게 확인·수정")
    return handled


# ── 열 페어 결정 (당기 N열 ↔ 전기 N열) ──────────────────────────────


def _column_pairs(tv: TableView, cur: list[int], prev: list[int]
                  ) -> tuple[list[tuple[int, int]], list[int]] | None:
    """당기↔전기 열 페어. 열 수가 같으면 순서 대응, 다르면 헤더의 당→전 플립
    일치로 대응('당기말'↔'전기말'). 전기 열이 하나라도 남으면 페어 실패(None).
    상대 없는 당기 전용 열('당기말 연이자율' 등)은 leftover로 반환."""
    if len(cur) == len(prev):
        return list(zip(cur, prev)), []
    by_header: dict[str, list[int]] = {}
    for pc in prev:
        by_header.setdefault(norm(tv.col_header(pc)), []).append(pc)
    pairs: list[tuple[int, int]] = []
    leftover: list[int] = []
    for cc in cur:
        want = _CUR_IN_HEADER_RE.sub(
            lambda m: "전" + (m.group(1) or "") + "기",
            norm(tv.col_header(cc)))
        cands = by_header.get(want)
        if cands:
            pairs.append((cc, cands.pop(0)))
        else:
            leftover.append(cc)
    if not pairs or any(by_header.values()):
        return None
    return pairs, leftover


def roll_ops(doc: Document, profile=None) -> tuple[list[splice.Op], RollReport]:
    ops: list[splice.Op] = []
    report = RollReport()
    mapping = _gisu_mapping(doc)
    allowed_years = _period_years(doc)
    touched: set[str] = set()  # op가 걸린 노드 — 중복 치환·중복 보고 방지
    painted: list[tuple] = []  # 색칠 대상 (원본 바이트 앵커 — 종료 시 롤 위치로 확정)

    with use_profile(profile):
        views, _ = build_views(doc)

        # ── ⓪ 구조 단위 교정·이동 ────────────────────────────────────
        _split_mixed_headings(doc, report, ops, touched, mapping, allowed_years)
        handled = _roll_paired_tables(doc, report, ops, touched, painted)
        handled |= _roll_socie(doc, views, report, ops, touched, painted,
                               allowed_years, handled)

        # ── ① 표별 당기 → 전기 값 이동 (열 페어) ────────────────────
        for tv in views:
            if _skip_table(tv) or tv.table.id in handled:
                continue
            if "자본변동표" in norm(tv.sec_title):
                # ⓪에서 블록 교체에 실패한 자본변동표만 — 항상 표면화
                if any(tv.value(i, c) is not None for i in tv.body
                       for c in range(tv.ncols)):
                    report.add("미처리", tv.context, f"{tv.table.id} 자본변동표",
                               note="블록 경계 인식 실패 — 당기 실적을 수기 작성")
                continue
            cur = [c for c in range(tv.ncols) if tv.period_bucket(c) == "당기"]
            prev = [c for c in range(tv.ncols) if tv.period_bucket(c) == "전기"]
            if not cur:
                continue  # 기간 열 없는 표 — 롤 무관
            paired = _column_pairs(tv, cur, prev)
            if paired is None:
                report.add("미처리", tv.context,
                           f"{tv.table.id} 당기 {len(cur)}열/전기 {len(prev)}열",
                           note="열 구성이 대칭이 아니고 헤더 대응도 없어 자동 이동 안 함")
                continue
            pairs, leftover = paired
            if any(tv.cell(i, c) is not None and tv.cell(i, c).tag == "TU"
                   for i in tv.body for c in cur + prev):
                report.add("미처리", tv.context, f"{tv.table.id} TU 열 포함",
                           note="자동 필드 열은 이동 대상 아님")
                continue
            for cc in leftover:
                report.add("확인필요", tv.context,
                           f"{tv.table.id} 당기 전용 열 "
                           f"'{tv.col_header(cc)[:20]}' — 값 유지",
                           note="상대 전기 열이 없어 이동 안 함 — 새 기수 기준으로 확인")
            n_moved = 0
            moved_cells: list[tuple[int, int]] = []  # (행 시작 오프셋, 전기 열)
            for i in tv.body:
                for cc, pc in pairs:
                    c_cell, p_cell = tv.cell(i, cc), tv.cell(i, pc)
                    if c_cell is None or p_cell is None:
                        continue
                    if c_cell.text != p_cell.text:
                        ops.append(splice.replace_inner(p_cell, c_cell.text))
                        touched.add(p_cell.id)
                    if c_cell.text.strip():
                        # 값이 같아 op가 없어도 전기 칸은 '자동 이월된 값' — 색칠
                        moved_cells.append((tv.table.rows[i].start, pc))
                    if c_cell.text:
                        ops.append(splice.replace_inner(c_cell, ""))
                        touched.add(c_cell.id)
                    n_moved += 1
            if n_moved:
                if moved_cells:
                    painted.append(("cells", moved_cells))
                report.add("자동확정", tv.context,
                           f"{tv.table.id} 당기값 {n_moved}칸 → 전기 열",
                           note="당기 열은 빈 값 — 결산 수치 입력 대상")

        # ── ② 기수 라벨 +1 (TU·값이동 셀 제외) ──────────────────────
        n_label = 0
        for node in doc.nodes.values():
            if node.tag == "TU" or node.id in touched:
                continue
            new_text = _roll_text(node.text, mapping, allowed_years)
            if new_text == node.text:
                continue
            ops.append(splice.replace_inner(node, new_text))
            touched.add(node.id)
            n_label += 1
        if mapping:
            desc = ", ".join(f"제{a}기→제{b}기" for a, b in sorted(mapping.items()))
            report.add("자동확정", "문서 전체", f"기수 라벨 {n_label}곳", note=desc)

        # ── ③ THEAD 날짜 +1년 (기수 없는 헤더 — '2026.3.31 현재'류) ──
        n_head_date = 0
        for tv in views:
            if _skip_table(tv):
                continue
            for i in [i for i in range(tv.n) if tv.table.rows[i].in_thead]:
                for cell in tv.table.rows[i].cells:
                    if cell.tag == "TU" or cell.id in touched:
                        continue  # 기수 셀은 ②에서 이미 처리
                    new_text = _bump_year(cell.text, allowed_years)
                    if new_text != cell.text:
                        ops.append(splice.replace_inner(cell, new_text))
                        touched.add(cell.id)
                        n_head_date += 1
        if n_head_date:
            report.add("자동확정", "표 헤더", f"기간 날짜 +1년 {n_head_date}곳")

        # ── ④ TU 자동필드: AUNITVALUE·표시 연도 +1 (스파이크 V4 방식) ─
        n_tu = 0
        for node in doc.nodes.values():
            if node.tag != "TU":
                continue
            frag = doc.data[node.start:node.end]
            rolled = _YEAR_RE.sub(lambda m: str(int(m.group()) + 1).encode(), frag)
            if rolled != frag:
                ops.append(splice.Op(pos=node.start, start=node.start,
                                     end=node.end, payload=rolled))
                n_tu += 1
        if n_tu:
            report.add("자동확정", "자동 필드", f"TU 기간 +1년 {n_tu}곳")

        # ── ⑤ 손대지 않은 날짜 — 확인필요 목록 (역사적 날짜 오염 방지) ─
        for node in doc.nodes.values():
            if node.tag == "TU" or node.id in touched:
                continue
            for m in _DATE_RE.finditer(node.text):
                ctx = node.text[max(0, m.start() - 12):m.end() + 10]
                report.add("확인필요", node.id, "날짜 발견 (자동 치환 안 함)",
                           before=ctx.replace("\n", " ").strip(),
                           note="보고기간 표기면 +1년으로 직접 수정")
                break  # 노드당 1건만

    report.add("확인필요", "meta/문서정보", "보고서 기간·문서명",
               note="DART 편집기 [문서정보]에서 기수·기간 확인 (도구는 미변경)")
    report.painted = _final_painted(ops, painted)
    return ops, report


def _final_painted(ops: list[splice.Op], painted: list[tuple]) -> list[tuple]:
    """색칠 앵커(원본 바이트 오프셋) → 롤 적용 후 오프셋.

    표 신설 등으로 노드 ID가 밀려도 위치는 결정론적으로 추적된다:
    오프셋 앞에서 끝나는 모든 op의 길이 변화 합이 이동량."""
    deltas = sorted((op.end, len(op.payload) - (op.end - op.start)) for op in ops)
    keys = [end for end, _ in deltas]
    csum = [0]
    for _, d in deltas:
        csum.append(csum[-1] + d)

    def shift(pos: int) -> int:
        return pos + csum[bisect.bisect_right(keys, pos)]

    out: list[tuple] = []
    for spec in painted:
        kind = spec[0]
        if kind == "cells":
            out.append(("cells", [(shift(p), gc) for p, gc in spec[1]]))
        elif kind == "ptable":   # 문단 교체 payload 안에 신설된 표
            out.append(("table", shift(spec[1]) + spec[2]))
        else:                    # "table" | "socie"
            out.append((kind, shift(spec[1])))
    return out


def _skip_table(tv: TableView) -> bool:
    from .check import _P
    return bool(_P.skip_section_re.search(norm(tv.sec_title)))


# ── 자동 이동 셀 색칠 ────────────────────────────────────────────────

MOVED_FILL = "DDEBF7"  # 연한 파랑 — 검증 색(빨강/노랑)과 구분


def paint_moved(xlsx_path: str, rolled_doc: Document, painted: list,
                row_positions: dict) -> int:
    """롤포워드가 자동으로 채운 값 셀에 배경색 — 이동분을 한눈에 식별.

    painted의 앵커는 롤 적용 후 바이트 오프셋 — 재파싱한 문서의 행/표 시작
    오프셋으로 해석한다 (노드 ID 밀림과 무관)."""
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill

    from .annotate_xlsx import _cell_positions
    from .check import use_profile as _up

    pos_of = _cell_positions(rolled_doc, row_positions)
    row_at: dict[int, object] = {}
    tbl_at: dict[int, object] = {}
    for t in rolled_doc.tables.values():
        tbl_at[t.start] = t
        for row in t.rows:
            row_at[row.start] = row
    targets: set[tuple[str, int, int]] = set()

    def add(row, gc: int):
        loc = pos_of.get(row.id)
        if loc is not None:
            targets.add((loc[0], loc[1], 2 + gc))

    with _up(None):
        for spec in painted:
            kind = spec[0]
            if kind == "cells":
                for pos, gc in spec[1]:
                    row = row_at.get(pos)
                    if row is not None:
                        add(row, gc)
                continue
            t = tbl_at.get(spec[1])
            if t is None:
                continue
            tv = TableView(t, "")
            if kind == "table":
                rows = tv.body
            else:  # socie: 전기 블록(첫 기초~첫 기말) + 당기초 행
                opens = [i for i in tv.body if "기초" in norm(tv.row_label(i))
                         and any(tv.value(i, c) is not None
                                 for c in range(tv.ncols))]
                closes = [i for i in tv.body if "기말" in norm(tv.row_label(i))
                          and any(tv.value(i, c) is not None
                                  for c in range(tv.ncols))]
                if len(opens) < 2 or not closes:
                    continue
                rows = [i for i in tv.body if opens[0] <= i <= closes[0]]
                rows.append(opens[1])
            for i in rows:
                for gc, _cell in _numeric_cells(tv, i):
                    add(t.rows[i], gc)

    if not targets:
        return 0
    wb = load_workbook(xlsx_path)
    fill = PatternFill("solid", fgColor=MOVED_FILL)
    n = 0
    for sheet, r, c in targets:
        if sheet in wb.sheetnames:
            wb[sheet].cell(row=r, column=c).fill = fill
            n += 1
    wb.save(xlsx_path)
    return n


# ── 리포트 시트 ──────────────────────────────────────────────────────

ROLL_SHEET = "롤포워드"
_GRADE_ORDER = {"미처리": 0, "확인필요": 1, "자동확정": 2}


def write_roll_sheet(xlsx_path: str, report: RollReport):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, PatternFill

    from .export_xlsx import FILL_HEAD, FONT, FONT_HEAD, FONT_TITLE

    wb = load_workbook(xlsx_path)
    if ROLL_SHEET in wb.sheetnames:
        del wb[ROLL_SHEET]
    ws = wb.create_sheet(ROLL_SHEET, 0)
    ws.column_dimensions["A"].hidden = True
    for col, w in (("B", 9), ("C", 26), ("D", 34), ("E", 32), ("F", 32), ("G", 44)):
        ws.column_dimensions[col].width = w
    rid = iter(range(1, 100000))

    def put(r, c, v, font=FONT, fill=None):
        ws.cell(row=r, column=1, value=f"!r{next(rid):05d}").font = FONT
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = font
        if fill:
            cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=(c in (6, 7)))
        return cell

    put(2, 2, "롤포워드 결과 — 미처리·확인필요부터 위에서 아래로 처리하세요", FONT_TITLE)
    put(3, 2, "미처리 = 도구가 손대지 않은 표(수기 작성) · 확인필요 = 사람 확인 대상 · "
              "자동확정 = 결정론 치환 완료. 파란 배경 셀 = 자동으로 이동·채워진 값. "
              "당기 칸은 비어 있음 — 결산 수치를 입력한 뒤 check로 검산하세요.")
    r = 5
    for c, h in enumerate(["등급", "위치", "항목", "변경 전", "변경 후", "비고"], start=2):
        cell = put(r, c, h, FONT_HEAD)
        cell.fill = FILL_HEAD
    r += 1
    fill_of = {"미처리": PatternFill("solid", fgColor="FDE9E9"),
               "확인필요": PatternFill("solid", fgColor="FFF4DE")}
    for item in sorted(report.items, key=lambda x: _GRADE_ORDER.get(x.grade, 9)):
        put(r, 2, item.grade, fill=fill_of.get(item.grade))
        put(r, 3, item.where)
        put(r, 4, item.item)
        put(r, 5, item.before)
        put(r, 6, item.after)
        put(r, 7, item.note)
        r += 1
    wb.save(xlsx_path)
