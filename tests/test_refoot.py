"""A4 refoot — 전기 구조 고정 재계산: 채움/수정 제안, 고정점, 보류, 0건 게이트."""
from openpyxl import load_workbook

from dsdxlsx import refoot
from dsdxlsx.check import run_checks
from dsdxlsx.cli import main
from dsdxlsx.import_xlsx import import_xlsx
from dsdxlsx.parse import parse_contents

from fixtures import factory


def _exported(tmp_path):
    from dsdxlsx.export_xlsx import export_xlsx

    contents = factory.standard_contents()
    dsd_path = tmp_path / "s.dsd"
    dsd_path.write_bytes(factory.dsd_bytes(contents))
    xlsx = str(tmp_path / "s.xlsx")
    export_xlsx(parse_contents(contents), dsd_path.read_bytes(), str(dsd_path), xlsx)
    return xlsx


def _plan(xlsx):
    _, contents, imp = import_xlsx(xlsx)
    doc = parse_contents(contents)
    report = run_checks(doc)
    return refoot.refoot_plan(doc, report, imp.row_positions)


def _set(xlsx, sheet, label, col, value):
    wb = load_workbook(xlsx)
    ws = wb[sheet]
    for row in ws.iter_rows():
        if row[1].value and label in str(row[1].value):
            # ws.cell(value=None)은 대입을 생략하므로 셀 객체에 직접 대입 (클리어 지원)
            ws.cell(row=row[0].row, column=col).value = value
            break
    wb.save(xlsx)


def test_pristine_zero_proposals(tmp_path):
    """심리 완료 상태에서는 제안 0건 — 오기입 방지 회귀 게이트."""
    items, _ = _plan(_exported(tmp_path))
    assert items == []


def test_fix_and_fill_with_fixpoint(tmp_path):
    xlsx = _exported(tmp_path)
    # 결산 시나리오: 당기 leaf 수정(현금 1,000→1,200), 주석4 leaf 수정(예금 600→800),
    # 자산총계는 비움 — 상위 합계 3곳이 전부 어긋난 상태
    _set(xlsx, "재무상태표", "현금및현금성자산", 4, 1200)
    _set(xlsx, "재무상태표", "자산총계", 4, None)
    _set(xlsx, "4", "보통예금", 3, 800)

    items, skips = _plan(xlsx)
    by_label = {it.label.strip(): it for it in items}
    assert by_label["Ⅰ.유동자산"].new == 6200          # 1,200+5,000 (C2 전기 구조)
    assert by_label["자산총계"].old is None            # 빈 셀 채움
    assert by_label["자산총계"].new == 6200            # 고정점: 유동자산 갱신 후
    assert by_label["합  계"].new == 1200              # 주석4: 400+800

    # 적용 → 재검산 클린
    rc = main(["refoot", xlsx, "--apply"])
    assert rc == 0
    _, contents, _ = import_xlsx(xlsx)
    report = run_checks(parse_contents(contents))
    assert not [f for f in report.findings if f.level == "warn"]


def test_missing_component_is_held_not_guessed(tmp_path):
    xlsx = _exported(tmp_path)
    _set(xlsx, "재무상태표", "현금및현금성자산", 4, None)  # leaf 미입력
    items, skips = _plan(xlsx)
    assert not any("유동자산" in it.label for it in items)
    assert any("미입력" in s.reason for s in skips)


def test_cli_preview_does_not_write(tmp_path, capsys):
    xlsx = _exported(tmp_path)
    _set(xlsx, "재무상태표", "현금및현금성자산", 4, 1200)
    rc = main(["refoot", xlsx])
    out = capsys.readouterr().out
    assert rc == 0 and "미리보기" in out and "6,200" in out
    # 미리보기는 파일을 바꾸지 않는다 — 합계는 여전히 원래 값
    wb = load_workbook(xlsx)
    ws = wb["재무상태표"]
    total = next(r for r in ws.iter_rows() if r[1].value == "Ⅰ.유동자산")
    assert total[3].value == 6000
