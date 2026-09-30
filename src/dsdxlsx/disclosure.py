"""공시사항점검표(한공회 8400) 자동 작성 — DSD 의 재무제표·주석으로 목차 '해당여부'·'주석번호'를 채운다.

입력
  - 감사보고서 DSD(.dsd) 또는 `check --json` 으로 만든 review.json
  - 한공회 2026 표준조서 8400 원본(K-IFRS용·일반기준용) — 환경변수 KICPA_FORMS(xlaudit 과 같은 폴더) 기준

판정 근거(결정론)
  ① 주석 제목에 항목 핵심어가 있음        → 해당(Y), 주석번호 = 그 주석
  ② 재무제표 과목에 핵심어 + 잔액 ≠ 0     → 해당(Y), 주석번호 = 과목의 주석 참조
  ③ 주석 본문·표 행 이름에 핵심어         → 2회 이상 해당(Y), 1회는 '확인 필요'
  ④ 근거 없음                             → 해당 없음(NA)
  항상 공시하는 항목(재무제표 표시·회계정책·특수관계자 등)은 Y. 상장법인 추가 공시는 상장 여부로,
  중간재무보고는 보고서 종류로 가른다. ②는 있는데 ①·③이 없으면 '주석 누락 후보'로 표시한다.

산출(원본 양식 사본 — 원본은 건드리지 않음)
  목차 시트  G 해당여부(Y/NA)·H 주석번호  ← 자동. '확인 필요'는 비워 둔다(사람이 고름)
  8400 시트  F 검토: 해당 없음 항목의 요구사항 줄만 N/A · G 비고: 해당 항목 줄에 주석번호
  _자동판정  항목별 판정·근거·주석 누락 후보·점검표에 대응하지 않는 주석 (audit 테마)
  줄 단위 Y/N 검토(주석이 요구사항을 실제로 충족하는지)는 사람 또는 AI 검토(dsd-review) 몫이다.
"""
from __future__ import annotations

import os
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path

DEFAULT_KICPA = (r"C:\Users\PC\Desktop\Business\FY2026\01. Audit\00. 참고자료\01. 한공회 표준조서"
                 r"\2026 감사조서서식(예시)-엑셀 260827")
FORM_FILES = {
    "kifrs": r"1. K-IFRS_2026\감사조서서식_8400_공시사항점검표_KIFRS용_2026.xlsx",
    "kgaap": r"2. 일반기업회계기준_2026\감사조서서식_8400_공시사항점검표_일반기준용_2026.xlsx",
}
TOC_SHEET = "공시사항점검표 목차(2026)"
BODY_SHEET = "8400"

# ── 판정 규칙 ────────────────────────────────────────────────────────────────
# kind: always(항상 Y) · cond(근거로 판정) · strong(주석 제목·잔액만 Y, 본문만이면 확인 필요)
#       · listed(상장법인만) · interim(중간보고서만) · manual(항상 확인 필요)
# notes: 주석 제목 핵심어 · fs: 재무제표 과목 핵심어(잔액 ≠ 0) · text: 본문 핵심어


def R(kind="cond", notes=(), fs=(), text=()):
    return {"kind": kind, "notes": list(notes), "fs": list(fs), "text": list(text)}


RULES_KIFRS = {
    "DN1.1.1": R("always"), "DN1.1.2": R("always", notes=["회계정책", "작성기준"]),
    "DN1.2": R("always"), "DN1.3": R("always"), "DN1.4": R("always"), "DN1.5": R("always"),
    "DN2.1": R("always", notes=["일반사항", "일반 사항", "회사의 개요"]),
    "DN2.2.1": R("always", notes=["회계정책"]), "DN2.2.2": R("always", notes=["회계정책"]),
    "DN2.2.3": R(text=["회계정책의 변경", "회계추정의 변경", "회계추정치의 변경", "새로 적용한 기준서",
                       "새로 도입된 기준서", "제정·개정 기준서"]),
    "DN2.3": R("always", notes=["회계추정", "추정과 판단", "추정 및 판단"]),
    "DN2.4.1": R(notes=["재고자산"], fs=["재고자산", "원재료", "재공품"]),
    "DN2.4.2": R(notes=["유형자산"], fs=["유형자산"]),
    "DN2.4.3": R(notes=["투자부동산"], fs=["투자부동산"]),
    "DN2.4.4": R(notes=["무형자산"], fs=["무형자산"]),
    "DN2.4.5": R(fs=["영업권"], text=["영업권"]),
    "DN2.4.6": R(text=["자본화된 차입원가", "자본화한 차입원가", "차입원가를 자본화", "자본화이자율", "자본화 이자율"]),
    "DN2.4.7": R(text=["회수가능액", "손상검사", "사용가치"]),
    "DN2.4.8.1": R(notes=["종속기업", "관계기업", "공동기업"], fs=["종속기업", "관계기업", "공동기업"]),
    "DN2.4.8.2": R(notes=["종속기업"], fs=["종속기업"]),
    "DN2.4.8.3": R("strong", notes=["관계기업", "공동기업", "공동약정"], fs=["관계기업", "공동기업"]),
    "DN2.4.8.4": R(text=["비연결구조화기업", "구조화기업"]),
    "DN2.4.8.5": R(text=["별도재무제표"]),
    "DN2.4.9.1": R("always", notes=["금융상품", "현금및현금성자산", "매출채권", "매입채무", "금융자산", "금융부채"]),
    "DN2.4.9.2": R(notes=["범주"], text=["범주별 금융", "금융상품의 범주", "금융자산의 범주"]),
    "DN2.4.9.3": R(fs=["당기손익-공정가치", "당기손익공정가치"], text=["당기손익-공정가치"]),
    "DN2.4.9.4": R(fs=["기타포괄손익-공정가치", "기타포괄손익공정가치"], text=["기타포괄손익-공정가치"]),
    "DN2.4.9.5": R(text=["금융자산의 재분류", "재분류한 금융자산", "사업모형을 변경"]),
    "DN2.4.9.6": R(text=["상계되는 금융", "상계한 금융", "일괄상계약정"]),
    "DN2.4.9.7": R(text=["금융자산의 양도", "양도된 금융자산", "매출채권 할인", "매출채권의 양도", "팩토링"]),
    "DN2.4.9.8": R(notes=["담보"], text=["담보로 제공", "담보제공"]),
    "DN2.4.9.9": R(text=["손실충당금", "대손충당금"]),
    "DN2.4.9.10": R(notes=["전환사채", "신주인수권부사채", "교환사채", "복합금융상품"],
                    fs=["전환사채", "신주인수권부사채", "교환사채", "전환상환우선주", "상환전환우선주"],
                    text=["복합금융상품"]),
    "DN2.4.9.11": R("strong", text=["채무불이행", "약정 위반", "약정위반", "계약위반", "재무약정"]),
    "DN2.4.9.12.1": R(text=["위험회피회계", "위험회피수단"]),
    "DN2.4.9.12.2": R(notes=["공정가치"], text=["공정가치 서열체계", "공정가치서열체계", "수준 3", "수준3"]),
    "DN2.4.9.13.1": R(notes=["위험관리", "재무위험"], text=["재무위험관리", "위험관리 정책"]),
    "DN2.4.9.13.2": R(notes=["위험관리", "재무위험"], text=["신용위험", "유동성위험", "시장위험"]),
    "DN2.4.9.13.3": R(text=["신용위험"]),
    "DN2.4.9.13.4": R(text=["유동성위험"]),
    "DN2.4.9.13.5": R(text=["시장위험", "환위험", "이자율위험", "가격위험"]),
    "DN2.4.9.14": R("manual"),
    "DN2.4.9.15": R(text=["이자율지표", "이자율지표 개혁", "IBOR"]),
    "DN2.4.9.16": R(text=["지속가능성 연계", "ESG 연계", "탄소배출 감축", "ESG연계"]),
    "DN2.4.9.17": R(text=["전력구매계약", "재생에너지 전력", "자연에 의존하는 전력"]),
    "DN2.4.10": R(fs=["계약자산", "계약부채"], text=["계약자산", "계약부채"]),
    "DN2.4.11": R(notes=["충당부채"], fs=["충당부채"]),
    "DN2.4.12": R(notes=["확정급여"], fs=["확정급여", "사외적립자산"], text=["확정급여"]),
    "DN2.4.13": R(notes=["리스"], fs=["사용권자산", "리스부채"]),
    "DN2.4.14": R(text=["정부보조금", "국고보조금"]),
    "DN2.4.15": R("always", notes=["자본금", "자본잉여금", "기타자본", "이익잉여금", "결손금", "자본"]),
    "DN2.4.15.1": R("always", text=["자본관리", "자본위험관리"]),
    "DN2.5.1.1": R("always", notes=["매출", "수익", "영업수익"]),
    "DN2.5.1.2": R(text=["투입법", "진행률", "진행기준"]),
    "DN2.5.2": R("always", notes=["비용의 성격", "성격별", "판매비와관리비", "기타수익", "금융수익", "금융원가"]),
    "DN2.5.3": R("always", notes=["영업손익", "영업이익", "영업부문"]),
    "DN2.5.4": R(notes=["법인세"], fs=["법인세"]),
    "DN2.6": R("always", notes=["현금흐름", "현금의 유입", "현금 유", "유출입"]),
    "DN2.7": R(notes=["사업결합"], text=["사업결합", "이전대가"]),
    "DN2.8.1": R("always", notes=["특수관계"]),
    "DN2.8.2": R(notes=["약정"], text=["약정사항", "주요 약정", "약정한도"]),
    "DN2.8.3": R(notes=["우발"], text=["우발부채", "우발자산", "소송사건", "계류중인 소송"]),
    "DN2.8.4": R("always", notes=["보고기간후", "보고기간 후", "후속사건", "재무제표의 승인"]),
    "DN2.8.5": R(notes=["매각예정", "중단영업"], fs=["매각예정"], text=["중단영업"]),
    "DN2.8.6": R(notes=["가상자산"], fs=["가상자산"], text=["가상자산"]),
    "DN3.1": R(text=["전기오류수정", "전기 오류", "오류수정"]),
    "DN3.2": R(text=["초인플레이션"]),
    "DN3.3": R(notes=["계속기업"], text=["계속기업 가정", "계속기업으로서의 존속능력", "중요한 불확실성"]),
    "DN3.4": R(text=["일탈"]),
    "DN3.5": R(text=["회계연도의 변경", "보고기간 종료일을 변경", "결산일 변경", "결산일을 변경"]),
    "DN3.6": R(text=["연결재무제표를 작성하지"]),
    "DN3.7": R(notes=["주식기준보상", "주식선택권"], fs=["주식선택권", "주식보상비용"], text=["주식기준보상", "주식선택권"]),
    "DN3.8": R(text=["최초채택", "최초 채택"]),
    "DN3.9": R(text=["교환가능성"]),
    "DN4.1.1": R("listed", notes=["영업부문", "부문정보"], text=["보고부문"]),
    "DN4.1.2": R("listed", text=["단일 영업부문", "단일의 영업부문", "하나의 영업부문", "단일 보고부문"]),
    "DN4.2": R("listed", notes=["주당"], text=["주당이익", "주당손익", "주당순이익"]),
    "DN5.1": R(notes=["생물자산"], fs=["생물자산"]),
    "DN5.2": R(text=["민간투자사업", "서비스양허약정"]),
    "DN5.3": R(text=["리스제공자", "운용리스수익", "금융리스채권", "운용리스료수익"]),
    "DN5.4": R(text=["복구기금", "환경정화"]),
    "DN5.5": R(text=["규제이연계정"]),
    "DN6": R("interim"),
    "DN7.1": R(text=["조기적용", "조기 적용"]),  # 1118호 문맥은 아래 _extra 에서 한 번 더 좁힘
    "FN2013-1": R("always", notes=["특수관계"]),
    "FN2013-2": R(notes=["종속기업", "관계기업", "공동기업"], fs=["종속기업", "관계기업", "공동기업"]),
    "FN2014-1": R(notes=["공정가치"], text=["공정가치 서열체계", "공정가치서열체계"]),
    "FN2015-1": R(text=["금융자산의 양도", "양도된 금융자산", "매출채권 할인", "팩토링"]),
    "FN2016-1": R("always", notes=["금융상품", "수익"]),
    "FN2018-1": R(notes=["리스"], fs=["사용권자산", "리스부채"]),
    "FN2018-2": R(fs=["개발비"], text=["개발비"]),
    "FN2021-1": R(fs=["배출권", "배출부채"], text=["온실가스 배출권", "배출권"]),
    "FN2021-2": R(fs=["보험계약부채", "보험계약자산"], text=["보험계약"]),
    "FN2023-1": R(text=["전기오류수정", "오류수정"]),
    "FN2023-2": R(fs=["가상자산"], text=["가상자산"]),
    "FN2023-3": R(text=["건설계약", "공사계약", "책임준공", "PF 보증"]),
    "FN2026-1": R("always", text=["제1118호"]),
}

RULES_KGAAP = {
    "DN1.1": R("always", notes=["일반사항", "일반 사항", "회사의 개요"]),
    "DN1.2": R("always", notes=["회계정책", "회계처리방침"]),
    "DN2.1.1": R(notes=["현금및현금성자산", "금융상품", "사용제한"], fs=["현금및현금성자산", "금융상품", "예금"]),
    "DN2.1.2": R(notes=["매출채권", "채권"], fs=["매출채권", "미수금"]),
    "DN2.1.3": R(notes=["재고자산"], fs=["재고자산", "원재료", "재공품"]),
    "DN2.1.4": R(notes=["유가증권", "매도가능", "만기보유"], fs=["유가증권", "매도가능증권", "만기보유증권", "단기매매증권"]),
    "DN2.1.5": R(notes=["지분법"], fs=["지분법"]),
    "DN2.1.6": R(text=["조인트벤처"]),
    "DN2.1.7": R(fs=["투자부동산", "장기대여금", "장기금융상품"], text=["투자부동산"]),
    "DN2.1.8": R(notes=["유형자산"], fs=["유형자산"]),
    "DN2.1.9": R(notes=["무형자산"], fs=["무형자산"]),
    "DN2.1.10": R("always", text=["담보", "손상차손", "보험"]),
    "DN2.2.1": R(notes=["차입금", "사채"], fs=["차입금", "사채"]),
    "DN2.2.2": R(fs=["전환사채", "신주인수권부사채", "교환사채", "상환전환우선주"], text=["복합금융상품", "전환권"]),
    "DN2.2.3": R(text=["금융부채"]),
    "DN2.2.4": R(notes=["퇴직급여"], fs=["퇴직급여충당부채", "퇴직연금"]),
    "DN2.3": R("always", notes=["외화", "화폐성", "담보", "충당부채"]),
    "DN2.4": R(text=["채권채무조정", "출자전환", "조건변경"]),
    "DN2.5.1": R("always", notes=["자본금", "자본", "이익잉여금", "결손금"]),
    "DN2.5.2": R(notes=["주식기준보상", "주식선택권"], fs=["주식선택권"], text=["주식기준보상", "주식선택권"]),
    "DN3.1": R("always", notes=["매출", "수익"]),
    "DN3.2": R("always", notes=["판매비와관리비", "판매비와 관리비"]),
    "DN3.3": R("always", notes=["부가가치"]),
    "DN3.4": R(notes=["법인세"], fs=["법인세"]),
    "DN3.5": R(notes=["주당"], text=["주당이익", "주당순이익", "주당손익"]),
    "DN3.6": R(notes=["포괄손익"], text=["포괄손익"]),
    "DN3.7": R(text=["중단사업"]),
    "DN4": R("always", notes=["현금흐름", "현금의 유입", "현금 유", "유출입"]),
    "DN5.1": R(notes=["우발", "약정"], text=["우발부채", "소송", "지급보증", "약정"]),
    "DN5.2": R("always", notes=["특수관계"]),
    "DN5.3": R("always", notes=["보고기간후", "보고기간 후", "재무제표의 확정", "재무제표 확정", "확정된 날",
                                 "재무제표의 승인"]),
    "DN5.4": R("manual"),
    "DN6.1": R(text=["회계변경", "회계정책의 변경", "회계추정의 변경"]),
    "DN6.2": R(text=["전기오류수정", "오류수정"]),
    "DN6.3": R(text=["재분류", "표시방법의 변경", "분류방법의 변경"]),
    "DN6.4": R(text=["중소기업 회계처리 특례", "중소기업회계처리특례", "제31장"]),
    "DN6.5": R(notes=["리스"], fs=["리스"], text=["운용리스", "금융리스"]),
    "DN6.6": R(notes=["사업결합"], text=["사업결합", "합병"]),
    "DN6.6.1": R(fs=["영업권", "염가매수차익"], text=["영업권", "염가매수차익"]),
    "DN6.7": R(text=["분할합병", "물적분할", "인적분할"]),
    "DN6.8": R(text=["최초채택", "최초 채택"]),
    "DN6.9": R(text=["한국채택국제회계기준 도입", "K-IFRS 도입"]),
    "DN6.10": R("manual", notes=["계속기업"]),
    "DN7.1": R(text=["공사계약", "공사수익", "진행기준", "건설형 공사"]),
    "DN7.2": R("manual"),
    "DN7.3": R(text=["특수활동"]),
    "DN7.4": R(fs=["배출권", "배출부채"], text=["온실가스 배출권", "배출권"]),
    "DN8": R(text=["연결재무제표"]),
    "DN9": R(text=["금융업"]),
    "FN2013-1": R("always", notes=["특수관계"]),
    "FN2021-1": R(fs=["배출권", "배출부채"], text=["온실가스 배출권", "배출권"]),
    "FN2023-2": R(fs=["가상자산"], text=["가상자산"]),
}

VERDICT_Y, VERDICT_NA, VERDICT_Q = "Y", "NA", "확인 필요"


@dataclass
class Item:
    code: str
    title: str
    row: int
    level: int
    verdict: str = ""
    notes: list[int] = field(default_factory=list)
    basis: str = ""
    warn: str = ""


# ── 문서 증거 ────────────────────────────────────────────────────────────────

def _norm(s: str) -> str:
    s = re.sub(r"\(주석[^)]*\)|\(주\s*[\d,\s]+\)", "", s or "")
    s = s.replace("－", "-").replace("–", "-")
    return re.sub(r"[\s·ㆍ]", "", s)


class Evidence:
    def __init__(self, review: dict):
        doc = review["document"]
        self.notes = [(n["no"], n["title"] or "") for n in doc["notes"]]
        # 회계정책 주석 — 정책 서술은 '그런 거래가 있으면 이렇게 한다'는 일반 문장이라 해당 근거로 약하다
        self.policy_notes = {no for no, t in self.notes if "회계정책" in _norm(t) or "작성기준" in _norm(t)}
        self.fs = []
        for r in doc["fs_rows"]:
            amts = [v for v in (r.get("amounts") or {}).values() if isinstance(v, (int, float))]
            self.fs.append((_norm(r["label"]), r["label"], any(abs(v) > 0 for v in amts), r.get("refs") or [],
                            _norm(r.get("section", ""))))
        # 본문: (섹션 라벨, 정규화 텍스트) — 문단 + 표 행 이름
        self.texts: list[tuple[str, str]] = []
        for p in doc["paragraphs"]:
            self.texts.append((p["section"], _norm(p["text"])))
        for t in doc["tables"]:
            for row in t["rows"]:
                self.texts.append((t["section"], _norm(row["label"])))
        self.all_text = " ".join(t for _, t in self.texts)
        raw = " ".join(p["text"] for p in doc["paragraphs"][:80])
        self.form = "kifrs" if self.all_text.count("한국채택국제회계기준") >= self.all_text.count("일반기업회계기준") else "kgaap"
        self.listed = bool(re.search(r"유가증권시장|코스닥|코넥스|KOSDAQ|KOSPI", self.all_text))
        self.separate = "별도재무제표" in _norm(raw)[:4000] or "(별도)" in raw[:2000]
        self.interim = bool(re.search(r"분기|반기", _norm(raw)[:3000])) and "감사보고서" not in _norm(raw)[:600]
        self.company = ""
        m = re.search(r"(?:주식회사|㈜)\s*([^\s\n]+)|([^\s\n]+)\s*(?:주식회사|㈜)", raw)
        if m:
            self.company = (m.group(0) or "").strip()
        m = re.search(r"(20\d\d)\s*년\s*(\d{1,2})\s*월\s*(\d{1,2})\s*일\s*현재", raw)
        self.period_end = f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}" if m else ""

    @staticmethod
    def _note_no(section: str) -> int | None:
        m = re.match(r"주석(\d+)", section or "")
        return int(m.group(1)) if m else None

    def note_hits(self, kws: list[str]) -> list[int]:
        out = []
        for no, title in self.notes:
            t = _norm(title)
            if any(_norm(k) in t for k in kws):
                out.append(no)
        return out

    def fs_hits(self, kws: list[str]) -> tuple[list[str], list[int]]:
        labels, refs = [], []
        for norm, label, nonzero, rf, sec in self.fs:
            # 잔액 근거는 재무상태표 과목만 — 손익 과목(제품매출·지분법이익 등)은 오탐이 많다
            if sec and "재무상태표" not in sec and "대차대조표" not in sec:
                continue
            if nonzero and any(_norm(k) in norm for k in kws):
                labels.append(re.sub(r"\s+", " ", label).strip())
                refs += [x for x in rf if isinstance(x, int)]
        return labels, refs

    def text_hits(self, kws: list[str]) -> tuple[int, list[int], int]:
        """(정책 주석 밖 적중 수, 적중 주석 번호, 정책 주석 안 적중 수)."""
        n, notes, pol = 0, [], 0
        for sec, t in self.texts:
            if any(_norm(k) in t for k in kws):
                no = self._note_no(sec)
                if no in self.policy_notes:
                    pol += 1
                    continue
                n += 1
                if no is not None and no not in notes:
                    notes.append(no)
        return n, notes, pol


# ── 판정 ────────────────────────────────────────────────────────────────────

def judge(item: Item, rule: dict | None, ev: Evidence) -> None:
    if rule is None:
        item.verdict, item.basis = VERDICT_Q, "판정 규칙 없음"
        return
    kind = rule["kind"]
    nh = ev.note_hits(rule["notes"]) if rule["notes"] else []
    fl, fr = ev.fs_hits(rule["fs"]) if rule["fs"] else ([], [])
    tkw = rule["text"] + rule["fs"]          # 과목 이름도 본문에서 찾는다(주석 안 설명·표 행)
    tn, tnotes, tpol = ev.text_hits(tkw) if tkw else (0, [], 0)
    basis = []
    if nh:
        basis.append("주석 제목: " + ", ".join(f"{n}. {dict(ev.notes).get(n, '')}" for n in nh[:3]))
    if fl:
        basis.append("재무제표: " + ", ".join(fl[:3]))
    if tn:
        basis.append(f"본문 {tn}곳" + (f"(주석 {', '.join(map(str, tnotes[:4]))})" if tnotes else ""))
    notes = sorted(set(nh or fr or tnotes))[:4] if (nh or fr or tnotes) else []

    if kind == "interim":
        item.verdict = VERDICT_Y if ev.interim else VERDICT_NA
        item.basis = "중간재무제표" if ev.interim else "연차 재무제표"
        return
    if kind == "manual":
        item.verdict, item.basis, item.notes = VERDICT_Q, "; ".join(basis) or "자동 판정 대상 아님 - 직접 확인", notes
        return
    if kind == "listed" and not ev.listed and not (nh or fl or tn):
        item.verdict, item.basis = VERDICT_NA, "비상장(상장 표시 없음)"
        return
    if kind == "always":
        item.verdict, item.notes = VERDICT_Y, notes
        item.basis = "항상 공시" + (" · " + "; ".join(basis) if basis else "")
        return
    if kind == "strong" and not (nh or fl):
        if tn:
            item.verdict, item.notes = VERDICT_Q, notes
            item.basis = "; ".join(basis) + " - 본문 언급만(특수관계자·일반 문구일 수 있음)"
        else:
            item.verdict = VERDICT_NA
            item.basis = f"회계정책 주석에만 일반 서술({tpol}곳)" if tpol else "근거 없음"
        return
    if nh or fl or tn >= 2:
        item.verdict, item.notes, item.basis = VERDICT_Y, notes, "; ".join(basis)
        if fl and not nh and not tnotes:
            item.warn = ("주석 누락 후보(재무제표 잔액은 있으나 별도 주석 없음"
                         + (" - 회계정책 주석에만 언급)" if tpol else ")"))
    elif tn == 1:
        item.verdict, item.notes, item.basis = VERDICT_Q, notes, "; ".join(basis) + " - 한 곳뿐"
    else:
        item.verdict = VERDICT_NA
        item.basis = f"회계정책 주석에만 일반 서술({tpol}곳)" if tpol else "근거 없음"
    if kind == "listed" and item.verdict == VERDICT_Y and not ev.listed:
        item.verdict, item.basis = VERDICT_Q, item.basis + " (상장 표시 없음)"


def _extra(items: dict[str, Item], ev: Evidence, form: str) -> None:
    """규칙 한 줄로 못 담는 조정."""
    if form == "kifrs":
        it = items.get("DN2.4.8.5")
        if it:
            it.verdict = VERDICT_Y if ev.separate else VERDICT_NA
            it.basis = "표지·제목에 별도재무제표" if ev.separate else "별도재무제표 표시 없음"
        it = items.get("DN2.4.12")
        if it and it.verdict == VERDICT_NA and ev.note_hits(["퇴직급여", "종업원급여"]):
            it.verdict, it.notes = VERDICT_Q, ev.note_hits(["퇴직급여", "종업원급여"])
            it.basis = "퇴직급여 주석은 있으나 확정급여 언급 없음 - 확정기여만인지 확인"
        it = items.get("DN7.1")
        if it and it.verdict != VERDICT_NA and "제1118호" not in ev.all_text:
            it.verdict, it.basis = VERDICT_NA, "1118호 조기적용 언급 없음"
        a, b = items.get("DN4.1.1"), items.get("DN4.1.2")
        if a and b and a.verdict == VERDICT_Y and b.verdict == VERDICT_Y:
            a.verdict, a.basis = VERDICT_Q, a.basis + " - 단일 부문 문구와 함께 나옴"
        if ev.listed:
            for code in ("DN4.1.1", "DN4.1.2"):
                it = items.get(code)
                if it and it.verdict == VERDICT_NA and (a.verdict != VERDICT_Y and b.verdict != VERDICT_Y):
                    it.verdict, it.basis = VERDICT_Q, "상장법인 - 부문 공시 확인"


def roll_up(items: list[Item]) -> None:
    """상위 항목 = 하위 판정 합성(하나라도 Y 면 Y, 모두 NA 면 NA, 나머지는 확인 필요)."""
    codes = [i.code for i in items]
    by = {i.code: i for i in items}
    for it in sorted(items, key=lambda x: -x.level):
        kids = [by[c] for c in codes if c != it.code and c.startswith(it.code + ".") and
                c.count(".") == it.code.count(".") + 1]
        if it.code == "FN":
            kids = [by[c] for c in codes if c.startswith("FN") and c != "FN"]
        if not kids or it.verdict:
            continue
        vs = {k.verdict for k in kids}
        it.verdict = VERDICT_Y if VERDICT_Y in vs else VERDICT_NA if vs == {VERDICT_NA} else VERDICT_Q
        it.notes = sorted({n for k in kids if k.verdict == VERDICT_Y for n in k.notes})[:6]
        it.basis = "하위 항목 종합"


# ── 양식 읽기·쓰기 ──────────────────────────────────────────────────────────

def form_path(form: str, kicpa_dir: str | None = None) -> Path:
    base = Path(kicpa_dir or os.getenv("KICPA_FORMS") or DEFAULT_KICPA)
    p = base / FORM_FILES[form]
    if not p.exists():
        raise FileNotFoundError(f"한공회 8400 원본이 없습니다: {p} (KICPA_FORMS 환경변수로 폴더 지정)")
    return p


def read_toc(ws) -> list[Item]:
    items = []
    for r in range(11, ws.max_row + 1):
        code = ws.cell(r, 10).value
        if not code or not str(code).startswith(("DN", "FN")):
            continue
        vals = [ws.cell(r, c).value for c in range(1, 7)]
        texts = [str(v).strip() for v in vals if v not in (None, "")]
        code = str(code).strip()
        level = code.count(".") + (0 if code.startswith("DN") else (1 if code != "FN" else 0))
        items.append(Item(code=code, title=texts[-1] if texts else "", row=r, level=level))
    return items


def _label_value_cell(ws, row: int, label: str):
    for c in range(1, 12):
        v = ws.cell(row, c).value
        if isinstance(v, str) and label in v.replace(" ", ""):
            target = ws.cell(row, c + 1)
            for rng in ws.merged_cells.ranges:
                if target.coordinate in rng:
                    return ws.cell(rng.min_row, rng.min_col)
            return target
    return None


def _free_path(p: Path) -> Path:
    if not p.exists():
        return p
    i = 2
    while True:
        q = p.with_name(f"{p.stem}_{i}{p.suffix}")
        if not q.exists():
            return q
        i += 1


def run(review: dict, out: Path, form: str | None = None, kicpa_dir: str | None = None,
        listed: bool | None = None, company: str | None = None, period_end: str | None = None) -> dict:
    import openpyxl

    ev = Evidence(review)
    form = form or ev.form
    if listed is not None:
        ev.listed = listed
    src = form_path(form, kicpa_dir)
    out = _free_path(out)
    shutil.copyfile(src, out)
    wb = openpyxl.load_workbook(out)
    toc, body = wb[TOC_SHEET], wb[BODY_SHEET]

    items = read_toc(toc)
    rules = RULES_KIFRS if form == "kifrs" else RULES_KGAAP
    codes = {i.code for i in items}
    for it in items:
        is_parent = any(c.startswith(it.code + ".") for c in codes) or it.code == "FN"
        if it.code in rules:
            judge(it, rules[it.code], ev)
        elif not is_parent:
            judge(it, None, ev)
    _extra({i.code: i for i in items}, ev, form)
    roll_up(items)

    # 목차 G·H
    for it in items:
        g = toc.cell(it.row, 7)
        g.value = it.verdict if it.verdict in (VERDICT_Y, VERDICT_NA) else None
        toc.cell(it.row, 8).value = ", ".join(map(str, it.notes)) if it.verdict == VERDICT_Y and it.notes else None
    name = company or ev.company
    pe = period_end or ev.period_end
    for ws, rows in ((toc, (5, 7)), (body, (7, 9, 10, 12))):
        for r in rows:
            for label, val in (("회사명", name), ("결산일", pe)):
                cell = _label_value_cell(ws, r, label)
                if cell is not None and val and cell.value in (None, ""):
                    cell.value = val

    # 8400 본문: 해당 없음 → F N/A, 해당 → G 주석번호
    by = {i.code: i for i in items}
    na_rows = y_rows = 0
    for r in range(1, body.max_row + 1):
        code = body.cell(r, 9).value
        it = by.get(str(code).strip()) if code else None
        if not it:
            continue
        a = str(body.cell(r, 1).value or "").strip()
        num = it.code[2:] if it.code.startswith("DN") else ""
        heading = bool(num) and (a.startswith(num + " ") or a.startswith(num + ".") and len(a.split()) > 0
                                 and a.split()[0].rstrip(".") == num)
        mid = [str(body.cell(r, c).value or "").strip() for c in range(2, 6)]
        header_row = any(x.replace(" ", "") in ("내용", "검토", "관련K-IFRS", "참조기준")
                         for x in [a] + mid + [str(body.cell(r, 6).value or "").strip()])
        if heading or header_row or not (a or any(mid)):
            continue
        f, g = body.cell(r, 6), body.cell(r, 7)
        if type(f).__name__ == "MergedCell" or type(g).__name__ == "MergedCell":
            continue                    # 요구사항 문장이 검토·비고 칸까지 병합된 줄(설명 줄)
        if it.verdict == VERDICT_NA and f.value in (None, ""):
            f.value = "N/A"
            na_rows += 1
        elif it.verdict == VERDICT_Y and it.notes and g.value in (None, ""):
            g.value = "주석 " + ", ".join(map(str, it.notes))
            y_rows += 1

    _write_result_sheet(wb, items, ev, form, name, pe)
    wb.save(out)
    cnt = {v: sum(1 for i in items if i.verdict == v) for v in (VERDICT_Y, VERDICT_NA, VERDICT_Q)}
    unmatched = _unmatched_notes(items, ev)
    return {"out": str(out), "form": form, "listed": ev.listed, "interim": ev.interim,
            "counts": cnt, "warns": [i.code for i in items if i.warn],
            "unmatched_notes": unmatched, "na_rows": na_rows, "y_rows": y_rows}


def _unmatched_notes(items: list[Item], ev: Evidence) -> list[tuple[int, str]]:
    used = {n for i in items if i.verdict in (VERDICT_Y, VERDICT_Q) for n in i.notes}
    return [(no, t) for no, t in ev.notes if no not in used]


def _write_result_sheet(wb, items, ev, form, company, period_end):
    import sys
    sys.path.insert(0, str(Path.home() / ".claude" / "skills" / "excel-theme"))
    from excel_theme import write_audit_header, write_section_bar, write_table, set_tab  # type: ignore

    ws = wb.create_sheet("_자동판정", 0)
    hdr = write_audit_header(ws, {"회사명": company or "", "조서번호": "8400", "결산일": period_end or "",
                                  "조서명": "공시사항점검표", "작성자": "자동 판정"}, theme="audit")
    last = 8                                   # B~H
    sum_hdr = write_section_bar(ws, hdr, "1. 판정 요약", theme="audit", last_col=last)
    det_bar = sum_hdr + 1 + 4 + 1              # 요약표 머리글 + 4줄 + 빈 줄
    det_hdr = det_bar + 2
    n = len(items)
    rng = f"$E${det_hdr + 1}:$E${det_hdr + n}"
    warn_rng = f"$H${det_hdr + 1}:$H${det_hdr + n}"
    write_table(ws, ["구분", "건수", "설명"], [
        ["해당(Y)", f'=COUNTIF({rng},"Y")', "목차 G열 Y, H열 주석번호"],
        ["해당 없음(NA)", f'=COUNTIF({rng},"NA")', "목차 G열 NA, 8400 요구사항 줄 F열 N/A"],
        ["확인 필요", f'=COUNTIF({rng},"확인 필요")', "목차 G열 빈칸 - 직접 선택"],
        ["주석 누락 후보", f'=COUNTIF({warn_rng},"주석 누락*")', "재무제표 잔액은 있으나 관련 주석 없음"],
    ], theme="audit", header_row=sum_hdr, number_cols={1: "accounting"})
    write_section_bar(ws, det_bar, "2. 항목별 판정", theme="audit", last_col=last)
    data = [[it.code, it.title, it.level, it.verdict, ", ".join(map(str, it.notes)), it.basis, it.warn]
            for it in items]
    write_table(ws, ["일련번호", "목차 항목", "수준", "판정", "주석번호", "판정 근거", "주의"], data,
                theme="audit", header_row=det_hdr, number_cols={2: "0"})
    um_bar = det_hdr + n + 2
    um_hdr = write_section_bar(ws, um_bar, "3. 대응 없는 주석", theme="audit", last_col=last)
    um = _unmatched_notes(items, ev)
    write_table(ws, ["주석번호", "주석 제목", "확인 사항"],
                [[no, t, "해당 목차 항목 확인(회사 고유 주석이면 기타)"] for no, t in um] or
                [["-", "없음", ""]], theme="audit", header_row=um_hdr)
    for col, w in (("B", 12), ("C", 30), ("D", 8), ("E", 10), ("F", 10), ("G", 30), ("H", 30)):
        ws.column_dimensions[col].width = w
    for r in range(det_hdr + 1, det_hdr + n + 1):
        for c in (3, 7, 8):
            cell = ws.cell(r, c)
            cell.alignment = cell.alignment.copy(wrap_text=True)
    set_tab(ws, "output", "audit")
