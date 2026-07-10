"""대사맵(Tie) 엔진 — selector 해석, 제안, 검증, 저장, 시트 왕복, import 안전성."""
from openpyxl import load_workbook

from dsdxlsx import ties
from dsdxlsx.check import CheckReport, run_checks
from dsdxlsx.parse import parse_contents

from fixtures import factory


def _doc(**kw):
    return parse_contents(factory.standard_contents(**kw))


# ── selector 해석 ─────────────────────────────────────────────────────

def test_resolve_statement_and_note_rows():
    sections = ties.section_views(_doc())
    hit, why = ties.resolve(sections, ties.Selector("재무상태표", "매출채권"))
    assert hit is not None and why == ""
    assert hit.buckets().keys() == {"당기", "전기"}

    hit5, _ = ties.resolve(sections, ties.Selector("5", "차감후장부금액"))
    assert hit5 is not None
    assert hit5.value(hit5.buckets()["당기"]) == 5_000


def test_resolve_failure_is_explained():
    sections = ties.section_views(_doc())
    hit, why = ties.resolve(sections, ties.Selector("5", "없는라벨"))
    assert hit is None and "없음" in why
    hit, why = ties.resolve(sections, ties.Selector("99", "매출채권"))
    assert hit is None and "구획" in why


def test_resolve_nth_disambiguates():
    # '매출채권'은 주석5·주석7에 각 1회 — 주석 안에서는 유일, nth=2는 실패해야
    sections = ties.section_views(_doc())
    hit, why = ties.resolve(sections, ties.Selector("5", "매출채권", row_nth=2))
    assert hit is None and "2번째" in why


# ── 제안 생성 ─────────────────────────────────────────────────────────

def test_suggest_fs_note_and_note_note():
    doc = _doc()
    sugg = ties.suggest_ties(doc, existing=[])
    by_pair = {(t.a.section, t.a.row, t.b.section, t.b.row): t for t in sugg}

    # FS 현금 ↔ 주석4 합계 (합계행 우선 매치, high)
    t1 = by_pair.get(("재무상태표", "현금및현금성자산", "4", "합계"))
    assert t1 is not None and t1.status == "suggested" and t1.confidence == "high"

    # FS 매출채권 ↔ 주석5 차감후 장부금액 (차감후행 스코어, high)
    t2 = by_pair.get(("재무상태표", "매출채권", "5", "차감후장부금액"))
    assert t2 is not None and t2.confidence == "high"

    # 주석5 매출채권(총액) ↔ 주석7 담보 매출채권 (라벨+금액 일치)
    t3 = by_pair.get(("5", "매출채권", "7", "매출채권"))
    assert t3 is not None and t3.origin == "주석↔주석"


def test_suggest_dedups_against_existing_and_suppressed():
    doc = _doc()
    first = ties.suggest_ties(doc, existing=[])
    # 하나를 '제외'로 저장했다고 치면 다시 제안되지 않는다
    excluded = first[0]
    excluded.status = "suppressed"
    again = ties.suggest_ties(doc, existing=[excluded])
    assert excluded.pair_key() not in {t.pair_key() for t in again}
    assert len(again) == len(first) - 1


# ── 검증 (C8) ────────────────────────────────────────────────────────

def _confirm(tie: ties.Tie) -> ties.Tie:
    tie.status = "confirmed"
    return tie


def test_verify_confirmed_pass_both_buckets():
    doc = _doc()
    tie = _confirm(ties.Tie(a=ties.Selector("재무상태표", "매출채권"),
                            b=ties.Selector("5", "차감후장부금액")))
    report = CheckReport()
    pts = ties.verify_ties(doc, [tie], report=report)
    assert [p.bucket for p in pts] == ["당기", "전기"]
    assert all(p.ok for p in pts)
    assert not report.findings and report.counts.get("C8") == 2


def test_verify_mismatch_is_warned_with_diff():
    doc = _doc(cash_note_total="1,100")  # 주석4 합계를 FS와 어긋나게
    tie = _confirm(ties.Tie(a=ties.Selector("재무상태표", "현금및현금성자산"),
                            b=ties.Selector("4", "합계")))
    report = CheckReport()
    pts = ties.verify_ties(doc, [tie], report=report)
    bad = [p for p in pts if not p.ok]
    assert len(bad) == 1 and bad[0].bucket == "당기" and bad[0].diff == -100
    assert any(f.check == "C8" and f.level == "warn" for f in report.findings)


def test_verify_unresolved_is_surfaced_not_silent():
    doc = _doc()
    tie = _confirm(ties.Tie(a=ties.Selector("재무상태표", "작년에만있던라벨"),
                            b=ties.Selector("4", "합계")))
    report = CheckReport()
    pts = ties.verify_ties(doc, [tie], report=report)
    assert pts[0].ok is None and "없음" in pts[0].reason
    assert any(f.check == "C8" and "해석 불가" in f.message for f in report.findings)


# ── 저장소 왕복 ───────────────────────────────────────────────────────

def test_ties_json_roundtrip(tmp_path):
    path = tmp_path / "profiles" / "합성테스트(주)" / "ties.json"
    t = _confirm(ties.Tie(a=ties.Selector("재무상태표", "매출채권"),
                          b=ties.Selector("5", "차감후장부금액", row_nth=1),
                          relation="equal", note="메모", origin="FS↔주석"))
    ties.save_ties(path, [t], company="합성테스트(주)")
    loaded = ties.load_ties(path)
    assert len(loaded) == 1
    assert loaded[0].id == t.id and loaded[0].a == t.a and loaded[0].b == t.b
    assert loaded[0].note == "메모" and loaded[0].status == "confirmed"


# ── 대사맵 시트 왕복 + import 안전성 (E2E) ───────────────────────────

def _exported(tmp_path):
    from dsdxlsx.export_xlsx import export_xlsx
    from dsdxlsx.import_xlsx import import_xlsx

    contents = factory.standard_contents()
    dsd = factory.dsd_bytes(contents)
    dsd_path = tmp_path / "합성.dsd"
    dsd_path.write_bytes(dsd)
    xlsx_path = tmp_path / "합성.xlsx"
    doc = parse_contents(contents)
    export_xlsx(doc, dsd, str(dsd_path), str(xlsx_path))
    _, contents2, imp = import_xlsx(str(xlsx_path))
    return xlsx_path, parse_contents(contents2), imp


def test_tie_sheet_roundtrip_and_import_safety(tmp_path):
    from dsdxlsx.annotate_xlsx import annotate
    from dsdxlsx.import_xlsx import import_xlsx

    xlsx_path, doc, imp = _exported(tmp_path)
    report = run_checks(doc)
    sugg = ties.suggest_ties(doc, existing=[])
    assert sugg
    pts = ties.resolve_points(doc, sugg)
    annotate(str(xlsx_path), report, doc, imp.row_positions,
             tie_points=[], tie_suggestions=pts)

    # 시트에서 그대로 복원되는가 (_data JSON 계약)
    sheet_ties = ties.read_tie_sheet(str(xlsx_path))
    assert {t.id for t in sheet_ties} == {t.id for t in sugg}
    assert all(t.status == "suggested" for t in sheet_ties)

    # 사용자가 상태 셀만 '확정'으로 수정
    target_id = sugg[0].id
    wb = load_workbook(str(xlsx_path))
    ws = wb[ties.SHEET_NAME]
    for row in ws.iter_rows(min_col=3, max_col=3):
        if row[0].value == target_id:
            ws.cell(row=row[0].row, column=2, value="확정")
    wb.save(str(xlsx_path))

    sheet_ties2 = ties.read_tie_sheet(str(xlsx_path))
    status_of = {t.id: t.status for t in sheet_ties2}
    assert status_of[target_id] == "confirmed"

    # merge → persistable: 확정만 저장 대상
    merged = ties.merge_sheet_state([], sheet_ties2)
    keep = ties.persistable(merged)
    assert [t.id for t in keep] == [target_id]

    # 대사맵 시트가 있어도 import 결과는 무변경 (!t ID 행은 무시)
    _, _, imp2 = import_xlsx(str(xlsx_path))
    assert not imp2.changed


def test_confirmed_tie_writes_gating_formula(tmp_path):
    """확정 대사쌍은 recheck가 인식하는 ',TRUE,' 수식으로 기록된다."""
    from dsdxlsx.annotate_xlsx import annotate

    xlsx_path, doc, imp = _exported(tmp_path)
    tie = _confirm(ties.Tie(a=ties.Selector("재무상태표", "매출채권"),
                            b=ties.Selector("5", "차감후장부금액")))
    report = run_checks(doc)
    pts = ties.verify_ties(doc, [tie], report=report)
    annotate(str(xlsx_path), report, doc, imp.row_positions,
             tie_points=pts, tie_suggestions=[])

    wb = load_workbook(str(xlsx_path))
    ws = wb[ties.SHEET_NAME]
    formulas = [c.value for row in ws.iter_rows(min_col=10, max_col=10)
                for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert len(formulas) == 2  # 당기·전기
    assert all(",TRUE," in f for f in formulas)
    # 금액 셀은 원본 시트를 참조하는 살아있는 수식
    amounts = [c.value for row in ws.iter_rows(min_col=5, max_col=5)
               for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert amounts and all("재무상태표" in f for f in amounts)
