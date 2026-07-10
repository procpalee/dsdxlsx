"""verify_annotated — annotate가 심은 검산 수식의 헤드리스(OfficeCLI) 재평가.

합성 xlsx(annotate 출력 패턴 replicate, 고객데이터 미사용)로 검증한다.
officecli 바이너리가 없으면 skip."""
from __future__ import annotations

import pytest
from openpyxl import Workbook

from dsdxlsx import verify_annotated as VA

pytestmark = pytest.mark.skipif(
    not VA.officecli_available(), reason="officecli 미설치 (npm i -g @officecli/officecli)"
)


def _build_annotated(path: str):
    """annotate 출력 패턴: A열 행ID / B 라벨 / C 값 / 우측 =IF(...,TRUE,...) 검산 셀."""
    wb = Workbook()
    ws = wb.active
    ws.title = "재무상태표"
    for r, rid, label, val in [
        (5, "t014.r001", "현금및현금성자산", 300),
        (6, "t014.r002", "매출채권", 500),
        (7, "t014.r003", "재고자산", 200),
        (8, "t014.r004", "유동자산 소계", 1000),
        (9, "t014.r005", "자산총계", 1900),
        (10, "t014.r006", "부채와자본총계", 1880),  # 대차 20 불일치
    ]:
        ws.cell(r, 1, rid)
        ws.cell(r, 2, label)
        ws.cell(r, 3, val)
    ws["E8"] = "=IF(C8=(C5+C6+C7),TRUE,C8-(C5+C6+C7))"  # 통과
    ws["E10"] = "=IF(C9=C10,TRUE,C9-C10)"                # 불일치 20

    n5 = wb.create_sheet("주석5")
    for i, v in enumerate([1000, 1900, 555], start=1):
        n5.cell(i, 2, v)  # 자산총계 1900 존재, 재고 200 없음
    ws["E9"] = ("=IF(OR(COUNTIF('주석5'!$B$1:$Z$10,C9)>0,"
                "COUNTIF('주석5'!$B$1:$Z$10,-C9)>0),TRUE,\"미확인\")")  # 통과(크로스시트)
    ws["E7"] = ("=IF(OR(COUNTIF('주석5'!$B$1:$Z$10,C7)>0,"
                "COUNTIF('주석5'!$B$1:$Z$10,-C7)>0),TRUE,\"미확인\")")  # 미확인
    wb.save(path)


def test_find_checks(tmp_path):
    p = str(tmp_path / "annot.xlsx")
    _build_annotated(p)
    checks = VA.find_checks(p)
    coords = {c.coord for c in checks}
    assert coords == {"E7", "E8", "E9", "E10"}
    # 라벨은 A열 행ID가 아니라 B열 계정명
    by_coord = {c.coord: c.label for c in checks}
    assert by_coord["E10"] == "부채와자본총계"
    assert not any(c.label.startswith("t0") for c in checks)


def test_verify_headless(tmp_path):
    p = str(tmp_path / "annot.xlsx")
    _build_annotated(p)
    rep = VA.verify_annotated(p)
    assert not rep.ok
    # 대차 불일치 20 검출
    fails = {r.coord: r.diff for r in rep.fails}
    assert fails == {"E10": 20.0}
    # 크로스시트 대사: 자산총계(E9) 통과, 재고(E7) 미확인
    assert {r.coord for r in rep.unconfirmed} == {"E7"}
    assert {r.coord for r in rep.passes} == {"E8", "E9"}


def test_all_pass_exit(tmp_path):
    """불일치를 고치면 종합 PASS (ok=True)."""
    p = str(tmp_path / "annot.xlsx")
    _build_annotated(p)
    # E10을 통과로: C10을 1900으로 맞춤
    from openpyxl import load_workbook
    wb = load_workbook(p)
    wb["재무상태표"]["C10"] = 1900
    # 미확인(E7) 제거 — 재고 200을 주석5에 추가
    wb["주석5"].cell(4, 2, 200)
    wb.save(p)
    rep = VA.verify_annotated(p)
    assert rep.ok, str(rep)
