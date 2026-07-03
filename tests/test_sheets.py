"""시트 세분화(재무제표별/주석번호별) + 주석 목차 레이아웃 테스트."""
import re

import pytest
from openpyxl import load_workbook

from dsdxlsx.cli import main
from dsdxlsx.import_xlsx import import_xlsx


@pytest.fixture(scope="module")
def layout_xlsx(sample_dsd, tmp_path_factory):
    tmp = tmp_path_factory.mktemp("sheets")
    xlsx = str(tmp / "layout.xlsx")
    assert main(["export", sample_dsd, "-o", xlsx]) == 0
    return xlsx


@pytest.fixture(scope="module")
def sheet_names(layout_xlsx):
    wb = load_workbook(layout_xlsx, read_only=True)
    names = list(wb.sheetnames)
    wb.close()
    return names


def test_statement_sheets_exist(sheet_names):
    joined = " / ".join(sheet_names)
    assert any(n == "재무상태표" for n in sheet_names), joined
    assert any("손익계산서" in n for n in sheet_names), joined


def test_cover_and_report_sheets_removed(sheet_names):
    """표지·검토(감사)보고서·(첨부) 표지부는 시트를 만들지 않는다 — 원본 보존."""
    joined = " / ".join(sheet_names)
    assert "표지" not in sheet_names, joined
    assert not any("보고서" in n for n in sheet_names), joined
    assert not any("(첨부)" in n for n in sheet_names), joined


def test_note_sheets_are_number_only(sheet_names):
    notes = [n for n in sheet_names if re.match(r"^\d{1,2}$", n)]
    assert len(notes) >= 10, f"주석 시트가 {len(notes)}개뿐: {sheet_names}"
    nums = [int(n) for n in notes]
    assert nums == sorted(nums)
    assert notes[0] == "1"


def test_no_notes_leak_into_statement_sheets(sheet_names):
    # "30. 현금흐름표" 주석이 재무제표 시트로 오분류되지 않았는지: 현금흐름표 시트는 1개만
    assert sum(1 for n in sheet_names if n == "현금흐름표") == 1


def test_note_index_with_hyperlinks(layout_xlsx, sheet_names):
    wb = load_workbook(layout_xlsx)  # 하이퍼링크는 read_only에서 안 읽힘
    ws = wb["주석"]
    links = {}
    found_header = False
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == "주석 목차":
                found_header = True
            if cell.hyperlink is not None:
                links[cell.value] = cell.hyperlink.target
    wb.close()
    assert found_header
    note_sheets = [n for n in sheet_names if re.match(r"^\d{1,2}$", n)]
    assert len(links) == len(note_sheets)
    assert any("회사의 개요" in str(k) for k in links)
    for target in links.values():
        m = re.match(r"^#'(\d{1,2})'!A1$", target)
        assert m and m.group(1) in sheet_names, target


def test_index_rows_do_not_break_import(layout_xlsx):
    # 목차 행(!x..)이 신규 행으로 오검출되지 않고 무수정 왕복이 유지되는지
    _, new_contents, report = import_xlsx(layout_xlsx)
    assert not report.changed
