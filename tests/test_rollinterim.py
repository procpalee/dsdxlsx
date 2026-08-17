"""분·반기 당전기 밀기 — 기간 판정, 어휘·날짜 치환, 당기 비움, 전기 ← B/DART, 리포트."""
from __future__ import annotations

import pytest

from dsdxlsx import dart_client as D
from dsdxlsx import period as P
from dsdxlsx import rollinterim as RI
from dsdxlsx import splice
from dsdxlsx.parse import parse_contents

from fixtures import factory as F


# ── 합성 중간보고서 ─────────────────────────────────────────────────

def _header(doc_name: str, code: str, cik: str = "00000001") -> str:
    return (f'<DOCUMENT-HEADER><DOCUMENT-NAME ACODE="{code}">{doc_name}</DOCUMENT-NAME>'
            f'<COMPANY-NAME AREGCIK="{cik}" AACCOUNTTYPE="A">합성테스트(주)</COMPANY-NAME>'
            "</DOCUMENT-HEADER>")


def interim_doc(q: int, year: int, *, cash_cur: str, cash_prev: str, sales_3m: str,
                sales_cum: str, sales_3m_prev: str, sales_cum_prev: str,
                note_flow_cur: str, note_flow_prev: str) -> bytes:
    """분기(q=1)·반기(q=2) 검토보고서 축소판: 표지(TU) + BS + PL + CF + SCE + 주석 2개."""
    word = {1: "1분기", 2: "반기", 3: "3분기"}[q]
    dang = {1: "당분기", 2: "당반기", 3: "당분기"}[q]
    jeon = {1: "전분기", 2: "전반기", 3: "전분기"}[q]
    m, d = P.QUARTER_END[q]
    end = f"{year}{m:02d}{d:02d}"
    end_txt = f"{year}년 {m:02d}월 {d:02d}일"
    pend_txt = f"{year - 1}년 {m:02d}월 {d:02d}일"
    doc_name = "반기검토보고서" if q == 2 else "분기검토보고서"
    code = "90020" if q == 2 else "90220"
    cover = (
        "<TABLE ACLASS=\"NORMAL\" WIDTH=\"600\" BORDER=\"0\"><COLGROUP><COL WIDTH=\"600\"/></COLGROUP><TBODY>"
        + F.tr([F.cell(f"제 60 (당) {word}")])
        + F.tr([F.cell(f"{year}년 01월 01일", "TU", AUNIT="PERIODFROM", AUNITVALUE=f"{year}0101")])
        + F.tr([F.cell(end_txt, "TU", AUNIT="PERIODTO", AUNITVALUE=end)])
        + "</TBODY></TABLE>"
    )
    bs = F.table(
        [["과  목", f"제 60(당) {word}말", "제 59(전) 기말"]],
        [["현금및현금성자산", cash_cur, cash_prev],
         ["대손충당금", "(100)", "(90)"],
         ["자산총계", cash_cur, cash_prev]],
        unit="원",
    )
    pl = F.table(
        [["과  목", f"||COLSPAN=2||제 60(당) {word}", f"||COLSPAN=2||제 59(전) {word}"],
         ["", "3개월", "누적", "3개월", "누적"]],
        [["매출액", sales_3m, sales_cum, sales_3m_prev, sales_cum_prev],
         ["매출원가", "(500)", "(900)", "(400)", "(700)"],
         [("반기" if q == 2 else "분기") + "순이익(손실)", "100", "200", "80", "150"]],
        unit="원",
    )
    cf = F.table(
        [["과  목", f"제 60(당) {word}", f"제 59(전) {word}"]],
        [["영업활동현금흐름", "300", "250"],
         ["유형자산의 취득", "(50)", "(40)"],
         ["기말현금및현금성자산", cash_cur, "777"]],
        unit="원",
    )
    sce = F.table(
        [["구  분", "자본금", "이익잉여금", "총계"]],
        [[f"{year - 1}.01.01(전기초)", "1,000", "500", "1,500"],
         [f"{jeon}순이익", "-", "80", "80"],
         [f"{pend_txt}({jeon}말)", "1,000", "580", "1,580"],
         [f"{year}.01.01(당기초)", "1,000", "600", "1,600"],
         [f"{dang}순이익", "-", "100", "100"],
         [f"{end_txt}({dang}말)", "1,000", "700", "1,700"]],
        unit="원",
    )
    note1 = F.table(
        [["구  분", dang, jeon]],
        [["임대수익", note_flow_cur, note_flow_prev],
         ["합  계", note_flow_cur, note_flow_prev]],
        unit="원",
    )
    note2 = F.table(
        [["구  분", f"{dang}말", "전기말"]],
        [["보증금", "10", "9"],
         ["합  계", "10", "9"]],
        unit="원",
    )
    return F.document(
        _header(doc_name, code),
        F.p("표지"),
        cover,
        F.title(f"{word}재무제표 검토보고서", 2),
        F.p(f"본인은 {end_txt}로 종료되는 3개월 보고기간의 재무제표를 검토하였습니다."),
        F.title("재무제표", 3),
        F.statement_title_table("재 무 상 태 표"),
        bs,
        F.statement_title_table("포 괄 손 익 계 산 서"),
        pl,
        F.statement_title_table("현 금 흐 름 표"),
        cf,
        F.statement_title_table("자 본 변 동 표"),
        sce,
        F.title("주석", 4),
        F.p(f"1. 임대수익\n{dang} 중 임대수익 내역은 다음과 같습니다."),
        note1,
        F.p(f"2. 보증금\n{dang}말과 전기말 현재 보증금 내역"),
        note2,
    )


def _doc_1q_2026() -> bytes:
    return interim_doc(1, 2026, cash_cur="1,200", cash_prev="1,000",
                       sales_3m="2,000", sales_cum="2,000",
                       sales_3m_prev="1,800", sales_cum_prev="1,800",
                       note_flow_cur="30", note_flow_prev="25")


def _doc_2q_2025() -> bytes:
    return interim_doc(2, 2025, cash_cur="950", cash_prev="800",
                       sales_3m="1,900", sales_cum="3,700",
                       sales_3m_prev="1,500", sales_cum_prev="3,000",
                       note_flow_cur="55", note_flow_prev="48")


def _texts(doc):
    return [c.text for t in doc.tables.values() for r in t.rows for c in r.cells]


def _row(doc, label_part: str, table_hint: str | None = None):
    for t in doc.tables.values():
        for r in t.rows:
            if r.cells and label_part in r.cells[0].text:
                return [c.text for c in r.cells]
    raise AssertionError(f"행 없음: {label_part}")


# ── 기간 판정 ────────────────────────────────────────────────────────

def test_detect_period_and_target():
    doc = parse_contents(_doc_1q_2026())
    p = P.detect_period(doc)
    assert (p.year, p.quarter, p.corp_code) == (2026, 1, "00000001")
    assert not p.is_annual and not p.consolidated
    assert p.next_interim().label == "2026년 반기"
    assert P.Period(2026, 3).next_interim() == P.Period(2027, 1)
    assert P.parse_target("3Q", p).quarter == 3
    assert P.parse_target("2027Q1", p) == P.Period(2027, 1, "00000001", False, "90220", "분기검토보고서")


def test_vocab_rules_1q_to_2q_and_2q_to_3q():
    r = P.vocab_rules(1, 2)
    assert P.apply_vocab("제 60(당) 1분기", r) == "제 60(당) 반기"
    assert P.apply_vocab("당분기말과 전기말", r) == "당반기말과 전기말"
    assert P.apply_vocab("분 기 재 무 상 태 표", r) == "반 기 재 무 상 태 표"
    assert P.apply_vocab("분기별 실적, 4분기", r) == "분기별 실적, 4분기"
    r = P.vocab_rules(2, 3)
    assert P.apply_vocab("제 60(당) 반기말", r) == "제 60(당) 3분기말"
    assert P.apply_vocab("반기 총포괄이익", r) == "분기 총포괄이익"
    assert P.apply_vocab("요약반기재무제표", r) == "요약분기재무제표"
    assert P.apply_vocab("6개월", r) == "9개월"


def test_date_map_same_year_and_year_roll():
    m = P.date_map(P.Period(2026, 1), P.Period(2026, 2))
    assert P.apply_dates("2026년 03월 31일 현재 / 2025년 3월 31일 / 2025-12-31", m) \
        == "2026년 06월 30일 현재 / 2025년 6월 30일 / 2025-12-31"
    m = P.date_map(P.Period(2025, 3), P.Period(2026, 1))
    assert P.apply_dates("2025.09.30 / 2024년 12월 31일 / 2025년 1월 1일", m) \
        == "2026.03.31 / 2025년 12월 31일 / 2026년 1월 1일"


# ── 밀기 ─────────────────────────────────────────────────────────────

def _run(base_bytes: bytes, b_bytes: bytes | None = None, use_dart: bool = False):
    doc = parse_contents(base_bytes)
    base = P.detect_period(doc)
    inp = RI.RollInputs(doc=doc, base=base, target=base.next_interim(),
                        doc_b=parse_contents(b_bytes) if b_bytes else None,
                        use_dart=use_dart)
    ops, roll, drep = RI.roll_interim_ops(inp)
    rolled = splice.apply(doc.data, ops)
    return parse_contents(rolled), roll, drep


def test_1q_to_2q_with_prior_same_file():
    doc2, roll, drep = _run(_doc_1q_2026(), _doc_2q_2025())
    texts = _texts(doc2)
    # 어휘·날짜: 1분기 → 반기, 3.31 → 6.30, 전기말 12.31 유지
    assert any("제 60(당) 반기말" in t for t in texts)
    assert any("2026년 06월 30일" in t for t in texts)
    assert not any("1분기" in t for t in texts)
    # BS: 당기말 비움, 전기말 유지
    assert _row(doc2, "현금및현금성자산")[1:] == ["", "1,000"]
    assert _row(doc2, "대손충당금")[1:] == ["", "(90)"]
    # PL: 당기 비움, 전기 ← B 의 당기(25년 반기 3개월/누적)
    assert _row(doc2, "매출액")[1:] == ["", "", "1,900", "3,700"]
    assert _row(doc2, "매출원가")[1:] == ["", "", "(500)", "(900)"]
    # CF
    assert _row(doc2, "영업활동현금흐름")[1:] == ["", "300"]
    # 주석 흐름형: 전기 ← B 당기, 재무상태형: 전기말 유지
    assert _row(doc2, "임대수익")[1:] == ["", "55"]
    assert _row(doc2, "보증금")[1:] == ["", "9"]
    # SCE: 전기 블록 ← B 당기 블록(라벨 플립), 당기초 유지, 당기 변동·기말 비움
    assert _row(doc2, "(전반기말)")[1:] == ["1,000", "700", "1,700"]
    assert _row(doc2, "(당기초)")[1:] == ["1,000", "600", "1,600"]
    assert _row(doc2, "(당반기말)")[1:] == ["", "", ""]
    assert roll.count("미처리") == 0
    # 문단도 치환
    para = next(n for n in doc2.nodes.values() if "임대수익 내역" in n.text)
    assert "당반기 중" in para.text
    # 검토보고서 문단은 확인필요로만
    assert any("검토보고서 문단" in i.item for i in roll.items)
    # TU 자동필드
    tu = [n for n in doc2.nodes.values() if n.tag == "TU"]
    assert any(n.attrs.get("AUNITVALUE") == "20260630" and "06월 30일" in n.text for n in tu)


def test_1q_to_2q_without_sources_clears_prev_flow_and_flags():
    doc2, roll, drep = _run(_doc_1q_2026())
    assert _row(doc2, "매출액")[1:] == ["", "", "", ""]
    assert _row(doc2, "임대수익")[1:] == ["", ""]
    assert _row(doc2, "보증금")[1:] == ["", "9"]           # 전기말은 유지
    assert _row(doc2, "(전반기말)")[1:] == ["", "", ""]   # 전기 블록 비움
    assert roll.count("확인필요") >= 3
    assert drep.skipped_reason == "--no-dart"


def _fake_dart(monkeypatch):
    def items(kind_rows):
        out = []
        for sj, nm, th, add in kind_rows:
            out.append(D.Item(sj, f"id_{nm}", nm, len(out), th, add, None, None, None))
        return out

    def fetch(corp, year, reprt, fs):
        if reprt == "11012":   # 2025 반기
            return D.DartData(corp, year, reprt, fs, items([
                ("CIS", "수익(매출액)", 1900, 3700), ("CIS", "매출원가", 500, 900),
                ("CIS", "당기순이익(손실)", 100, 200),
                ("CF", "영업활동현금흐름", 300, None),
                ("CF", "유형자산의 취득", 50, None),
                ("CF", "기말현금및현금성자산", 950, None)]))
        if reprt == "11013":   # 2025 1분기 (앵커용)
            return D.DartData(corp, year, reprt, fs, items([
                ("CIS", "수익(매출액)", 1800, 1800), ("CIS", "매출원가", 400, 400),
                ("CIS", "당기순이익(손실)", 80, 80),
                ("CF", "영업활동현금흐름", 250, None), ("CF", "유형자산의 취득", 40, None),
                ("CF", "기말현금및현금성자산", 777, None)]))
        if reprt == "11011":   # 2025 사업보고서
            return D.DartData(corp, year, reprt, fs, items([
                ("BS", "현금및현금성자산", 1000, None), ("BS", "대손충당금", 90, None),
                ("BS", "자산총계", 1000, None)]))
        raise D.DartUnavailable("없음")

    monkeypatch.setattr(D, "fetch_statements", fetch)


def test_1q_to_2q_dart_fill_and_compare(monkeypatch):
    _fake_dart(monkeypatch)
    doc2, roll, drep = _run(_doc_1q_2026(), use_dart=True)
    # PL 전기 ← DART(라벨 매칭, 매출원가는 문서 관례(괄호)에 맞춰 음수 표기)
    assert _row(doc2, "매출액")[1:] == ["", "", "1,900", "3,700"]
    assert _row(doc2, "매출원가")[1:] == ["", "", "(500)", "(900)"]
    # CF: 방향 라벨(취득)은 문서 옛 값 부호(음수) 유지
    assert _row(doc2, "유형자산의 취득")[1:] == ["", "(50)"]
    assert _row(doc2, "기말현금및현금성자산")[1:] == ["", "950"]
    # BS 전기말 유지 + 대사 일치 (대손충당금은 절대값 비교)
    assert drep.count("불일치") == 0
    assert drep.count("일치") >= 3
    assert drep.count("채움") >= 6
    assert not drep.skipped_reason


def test_dart_unavailable_is_reported_not_fatal(monkeypatch):
    def fetch(*a, **k):
        raise D.DartUnavailable("비상장")
    monkeypatch.setattr(D, "fetch_statements", fetch)
    doc2, roll, drep = _run(_doc_1q_2026(), use_dart=True)
    assert "비상장" in drep.skipped_reason
    assert _row(doc2, "보증금")[1:] == ["", "9"]


def test_annual_rejected_by_cli(tmp_path):
    from dsdxlsx.cli import main
    F.write_dsd(tmp_path / "a.dsd", F.standard_contents())
    assert main(["roll-interim", str(tmp_path / "a.dsd"), "--no-dart"]) == 2


def test_cli_end_to_end(tmp_path):
    from openpyxl import load_workbook

    from dsdxlsx.cli import main
    a = tmp_path / "q1.dsd"
    b = tmp_path / "prev.dsd"
    F.write_dsd(a, _doc_1q_2026())
    F.write_dsd(b, _doc_2q_2025())
    rc = main(["roll-interim", str(a), "--prior-same", str(b), "--no-dart"])
    assert rc in (0, 1)
    out_dsd = tmp_path / "q1_당전기밀기.dsd"
    out_xlsx = tmp_path / "q1_당전기밀기.xlsx"
    assert out_dsd.exists() and out_xlsx.exists()
    wb = load_workbook(out_xlsx)
    assert "롤포워드" in wb.sheetnames and RI.DART_SHEET in wb.sheetnames


# ── dart_client 단위 ─────────────────────────────────────────────────

def test_canon_and_match():
    assert D.canon("매출액") == D.canon("수익(매출액)")
    assert D.canon("반기순이익(손실)") == D.canon("당기순이익(손실)")
    assert D.canon("기본주당이익(손실) (단위 : 원)") == D.canon("기본주당이익(손실)")
    assert D.canon("금융비용") == D.canon("금융원가")
    items = [D.Item("BS", "a", "유동자산", 1, 10, None, None, None, None),
             D.Item("BS", "b", "기타", 2, 5, None, None, None, None),
             D.Item("BS", "c", "기타", 3, 7, None, None, None, None)]
    m = D.match_items(["Ⅰ. 유동자산", "기타", "기타", "없는계정"], items)
    assert [m[i].account_id for i in (0, 1, 2)] == ["a", "b", "c"] and 3 not in m
    assert D.anchor_match([7, 5, 99], items, "thstrm")[0].account_id == "c"


# ── 3Q → 다음 해 1Q (해 넘김): 기수 +1, 전기말 ← 전기 감사보고서 ─────

def _doc_annual_2025() -> bytes:
    """감사보고서 축소판(2025 연차): BS 당기말=25.12.31, 재무상태형 주석."""
    cover = (
        "<TABLE ACLASS=\"NORMAL\" WIDTH=\"600\" BORDER=\"0\"><COLGROUP><COL WIDTH=\"600\"/></COLGROUP><TBODY>"
        + F.tr([F.cell("제 60 기")])
        + F.tr([F.cell("2025년 12월 31일", "TU", AUNIT="PERIODTO", AUNITVALUE="20251231")])
        + "</TBODY></TABLE>"
    )
    bs = F.table(
        [["과  목", "제 60(당) 기말", "제 59(전) 기말"]],
        [["현금및현금성자산", "1,000", "800"], ["대손충당금", "(90)", "(70)"],
         ["자산총계", "1,000", "800"]], unit="원")
    sce = F.table(
        [["구  분", "자본금", "이익잉여금", "총계"]],
        [["2024.01.01(전기초)", "1,000", "400", "1,400"],
         ["당기순이익", "-", "100", "100"],
         ["2024.12.31(전기말)", "1,000", "500", "1,500"],
         ["2025.01.01(당기초)", "1,000", "500", "1,500"],
         ["당기순이익", "-", "100", "100"],
         ["2025.12.31(당기말)", "1,000", "600", "1,600"]], unit="원")
    note2 = F.table([["구  분", "당기말", "전기말"]],
                    [["보증금", "9", "7"], ["합  계", "9", "7"]], unit="원")
    return F.document(
        _header("감사보고서", "00760"), F.p("표지"), cover,
        F.title("재무제표", 3),
        F.statement_title_table("재 무 상 태 표"), bs,
        F.statement_title_table("자 본 변 동 표"), sce,
        F.title("주석", 4),
        F.p("1. 임대수익\n당기 중 임대수익"),
        F.table([["구  분", "당기", "전기"]], [["임대수익", "120", "110"]], unit="원"),
        F.p("2. 보증금\n당기말과 전기말 현재 보증금 내역"), note2,
    )


def test_3q_to_next_1q_with_prior_annual():
    base = interim_doc(3, 2025, cash_cur="900", cash_prev="800",
                       sales_3m="1,000", sales_cum="3,000",
                       sales_3m_prev="900", sales_cum_prev="2,700",
                       note_flow_cur="70", note_flow_prev="60")
    doc = parse_contents(base)
    p = P.detect_period(doc)
    assert (p.year, p.quarter) == (2025, 3) and p.next_interim() == P.Period(2026, 1, "00000001", False, "90220", "분기검토보고서")
    inp = RI.RollInputs(doc=doc, base=p, target=p.next_interim(),
                        doc_c=parse_contents(_doc_annual_2025()), use_dart=False)
    ops, roll, drep = RI.roll_interim_ops(inp)
    doc2 = parse_contents(splice.apply(doc.data, ops))
    texts = _texts(doc2)
    # 기수 +1, 3분기 → 1분기, 당기말 2026.3.31, 전기말 2025.12.31
    assert any("제 61(당) 1분기말" in t for t in texts)
    assert any("제 60(전) 기말" in t for t in texts)
    assert any("2026년 03월 31일" in t for t in texts)
    tu = [n for n in doc2.nodes.values() if n.tag == "TU"]
    assert any(n.attrs.get("AUNITVALUE") == "20260331" for n in tu)
    # BS 전기말 ← 감사보고서 당기말(25.12.31)
    assert _row(doc2, "현금및현금성자산")[1:] == ["", "1,000"]
    assert _row(doc2, "대손충당금")[1:] == ["", "(90)"]
    # 재무상태형 주석 전기말 ← 감사보고서 당기말
    assert _row(doc2, "보증금")[1:] == ["", "9"]
    # SCE 당기초 ← 감사보고서 당기말 행
    assert _row(doc2, "(당기초)")[1:] == ["1,000", "600", "1,600"]
    # 흐름 전기: 전년 동기 없음 → 비움
    assert _row(doc2, "매출액")[1:] == ["", "", "", ""]
    assert roll.count("미처리") == 0
