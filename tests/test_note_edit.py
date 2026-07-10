"""A2 주석 편집 — 재번호 엔진, #새주석/#주석삭제, C4 동형성 게이트."""
import pytest
from openpyxl import load_workbook

from dsdxlsx import renumber, splice
from dsdxlsx.import_xlsx import ImportError_, import_xlsx
from dsdxlsx.parse import parse_contents

from fixtures import factory


def _doc(**kw):
    return parse_contents(factory.standard_contents(**kw))


# ── 재번호 엔진 ───────────────────────────────────────────────────────

def test_note_infos_and_own_heading():
    infos = renumber.note_infos(_doc())
    assert [i.no for i in infos] == [1, 4, 5, 7]
    assert all(i.own_heading for i in infos)


def test_renumber_headings_and_fs_refs():
    doc = _doc()
    ops, rep = renumber.renumber_ops(doc, {5: 6, 7: 8})
    new = splice.apply(doc.data, ops)
    doc2 = parse_contents(new)
    assert [i.no for i in renumber.note_infos(doc2)] == [1, 4, 6, 8]
    # FS 매출채권의 참조 열 '5' → '6' (구분자·정렬 보존)
    ar_ref = next(n.text for n in doc2.nodes.values() if n.text.strip() == "6"
                  and n.tag == "TD")
    assert ar_ref.strip() == "6"
    assert any("매출채권" in line for line in rep.ref_cells)
    assert len(rep.headings) == 2


def test_renumber_mentions_reported_not_substituted():
    contents = factory.standard_contents()
    # 서술 멘션이 있는 문단 추가 fixture: 주석7 문단에 '주석 5 참조' 삽입
    contents = contents.replace(
        "담보로 제공된 자산의".encode(),
        "주석 5 참조. 담보로 제공된 자산의".encode())
    doc = parse_contents(contents)
    ops, rep = renumber.renumber_ops(doc, {5: 6})
    assert any("주석 5" in m for m in rep.mentions)
    new = splice.apply(doc.data, ops)
    assert "주석 5 참조".encode() in new  # 서술문은 자동 치환하지 않음


def test_renumber_rejects_mixed_heading():
    # 주석 7 헤딩을 앞 문단 꼬리에 병합 (혼합 P — 흥구 4·13 유형)
    contents = factory.document(
        factory.p("표지"),
        factory.title("재무제표", 2),
        factory.statement_title_table("재 무 상 태 표"),
        factory.table([["과목", "제 2(당) 기말", "제 1(전) 기말"]],
                      [["현금", "100", "90"]], unit="원"),
        factory.title("주석", 3),
        factory.p("1. 일반사항\n내용"),
        factory.p("5. 매출채권\n내용입니다.\n7. 담보제공자산"),  # 7이 혼합
        factory.table([["구분", "당기말"]], [["매출채권", "100"]], unit="원"),
    )
    doc = parse_contents(contents)
    infos = renumber.note_infos(doc)
    assert not next(i for i in infos if i.no == 7).own_heading
    with pytest.raises(renumber.RenumberError, match="혼합 문단"):
        renumber.renumber_ops(doc, {7: 8})
    # 삭제도 거부 — 5의 꼬리에 7 헤딩이 붙어 있음
    five = next(i for i in infos if i.no == 5)
    assert renumber.deletable(infos, five) is not None


def test_close_gap_mapping():
    assert renumber.build_close_gap_mapping([1, 2, 3, 5, 6]) == {5: 4, 6: 5}
    assert renumber.build_close_gap_mapping([1, 2, 3]) == {}


def test_c4_signature_isomorphic_after_renumber():
    doc = _doc()
    before = renumber.c4_signature(doc)
    ops, _ = renumber.renumber_ops(doc, {4: 5, 5: 6, 7: 8})
    doc2 = parse_contents(splice.apply(doc.data, ops))
    assert renumber.c4_signature(doc2) == before


# ── E2E: #새주석 / #주석삭제 ─────────────────────────────────────────

def _exported(tmp_path):
    from dsdxlsx.export_xlsx import export_xlsx

    contents = factory.standard_contents()
    dsd_path = tmp_path / "s.dsd"
    dsd_path.write_bytes(factory.dsd_bytes(contents))
    xlsx = str(tmp_path / "s.xlsx")
    export_xlsx(parse_contents(contents), dsd_path.read_bytes(), str(dsd_path), xlsx)
    return xlsx


def _new_sheet(xlsx, name, rows):
    wb = load_workbook(xlsx)
    ws = wb.create_sheet(name)
    for r, row in enumerate(rows, start=1):
        for col, val in row.items():
            ws.cell(row=r, column=col, value=val)
    wb.save(xlsx)


def test_new_note_append(tmp_path):
    xlsx = _exported(tmp_path)
    _new_sheet(xlsx, "8", [
        {2: "#새주석 8 리스"},
        {2: "#문단 회사가 리스이용자인 경우의 정보는 다음과 같습니다."},
        {2: "#새표 원본=t004 헤더=1"},
        {2: "구  분", 3: "당기말", 4: "전기말"},
        {2: "사용권자산", 3: 800, 4: 700},
    ])
    _, new_contents, rep = import_xlsx(xlsx)
    assert rep.note_inserts and rep.structural
    doc2 = parse_contents(new_contents)
    infos = renumber.note_infos(doc2)
    assert [i.no for i in infos] == [1, 4, 5, 7, 8]
    eight = next(i for i in infos if i.no == 8)
    assert eight.own_heading and "리스" in eight.title
    texts = [c.text for b in eight.blocks if b.kind == "table"
             for r in b.item.rows for c in r.cells]
    assert "사용권자산" in texts and "800" in texts
    # 신설 표 위치 재대응 정합
    for tid, positions in rep.row_positions.items():
        assert len(positions) == len(doc2.tables[tid].rows), tid


def test_new_note_mid_insert_with_cascade(tmp_path):
    xlsx = _exported(tmp_path)
    _new_sheet(xlsx, "새5", [
        {2: "#새주석 5 재고자산"},
        {2: "#문단 재고자산 내역은 다음과 같습니다."},
    ])
    _, new_contents, rep = import_xlsx(xlsx)
    assert rep.note_inserts and rep.renumbered
    doc2 = parse_contents(new_contents)
    infos = renumber.note_infos(doc2)
    assert [i.no for i in infos] == [1, 4, 5, 6, 8]  # 구5→6, 구7→8
    assert "재고자산" in next(i for i in infos if i.no == 5).title
    # FS 매출채권 참조 열 5 → 6 캐스케이드
    assert any("매출채권" in line and "6" in line for line in rep.renumbered)
    # 재번호 후에도 C4 클린 (참조와 주석 번호가 함께 이동)
    from dsdxlsx.check import run_checks
    report = run_checks(doc2, only={"C4"})
    assert not [f for f in report.findings if f.level == "warn"]


def test_delete_note(tmp_path):
    xlsx = _exported(tmp_path)
    wb = load_workbook(xlsx)
    ws = wb["4"]  # 주석4 현금및현금성자산
    prow = next(r for r in ws.iter_rows()
                if r[0].value and str(r[0].value).startswith("p"))
    ws.cell(row=prow[0].row, column=2, value="#주석삭제")
    wb.save(xlsx)
    _, new_contents, rep = import_xlsx(xlsx)
    assert rep.note_deletes
    # 잔존 참조 경고 (FS 현금 행이 주석 4를 참조 중)
    assert any("주석 4" in w and "close-gap" in w for w in rep.warnings)
    doc2 = parse_contents(new_contents)
    assert [i.no for i in renumber.note_infos(doc2)] == [1, 5, 7]
    assert not any("보통예금" in c.text for t in doc2.tables.values()
                   for r in t.rows for c in r.cells)


def test_guard_new_note_needs_fresh_sheet_and_single_use(tmp_path):
    xlsx = _exported(tmp_path)
    _new_sheet(xlsx, "8", [{2: "내용"}, {2: "#새주석 8 리스"}])
    with pytest.raises(ImportError_, match="첫 행"):
        import_xlsx(xlsx)


def test_guard_new_note_number_gap(tmp_path):
    xlsx = _exported(tmp_path)
    _new_sheet(xlsx, "99", [{2: "#새주석 99 잘못된번호"}])
    with pytest.raises(ImportError_, match="불연속"):
        import_xlsx(xlsx)
