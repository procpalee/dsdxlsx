"""refoot — 합계 자동 재계산 (구조 고정 방식).

DART 편집기는 본문 합계를 자동 재계산하지 않는다. 롤포워드로 당기 열에 결산
수치(leaf)를 입력한 뒤, 상위 합계(유동자산·자산총계·매출총이익 등)를 이 명령이
채우거나 고친다.

안전 원칙 — 검증 휴리스틱을 '쓰기'에 쓸 때의 오기입 방어 (플랜 확정):
1. **구조 고정(pinning)**: 재계산 대상 열의 합산 구조를 추측하지 않는다.
   같은 표에서 **검증을 통과한 전기 열**의 footing 결과(C1/C2/C3/C7, 세로축)가
   확정한 항 구조(행+부호)만 당기 열에 적용한다.
2. **xlsx 셀만 기입** — DSD 쓰기 경로에 자동 숫자 생성이 섞이지 않는다.
   사용자는 엑셀에서 눈으로 확인하고, check의 살아있는 수식으로 즉시 재검증.
3. **미리보기 기본**, `--apply`로만 실제 기입. 구성 셀이 미입력이면 제안하지
   않고 사유를 표시. 심리 완료 원본에서는 제안 0건이어야 한다(회귀 게이트).

중첩 소계(유동자산 → 자산총계)는 표 내 고정점 반복(≤4패스)으로 수렴시킨다.
가로합(합계 열·ROWPART)은 v1 비대상.
"""
from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal

from .check import CheckReport, build_views, use_profile
from .model import Document

_FOOTING = {"C1", "C2", "C3", "C7"}
_MAX_PASSES = 4


@dataclass
class RefootItem:
    sheet: str
    xlsx_row: int
    xlsx_col: int
    context: str
    label: str
    check: str
    old: object          # None = 비어 있던 셀
    new: object
    n_terms: int

    def describe(self) -> str:
        old = "(비어 있음)" if self.old is None else f"{self.old:,}"
        return (f"[{self.context}] '{self.label}' {self.sheet}!"
                f"{self.xlsx_row}행: {old} → {self.new:,} "
                f"(전기 구조 {self.check}, 항 {self.n_terms}개)")


@dataclass
class RefootSkip:
    context: str
    label: str
    reason: str


def _cell_positions(doc: Document, row_positions: dict) -> dict:
    from .annotate_xlsx import _cell_positions as f
    return f(doc, row_positions)


@dataclass
class _Pin:
    """고정된 합산 구조 — check 결과 또는 단일 자식 항등에서 유도."""
    row_id: str
    grid_col: int                    # 전기 열
    terms: list                     # [(row_id, 전기 grid_col, 부호)]
    row_label: str
    check: str


def _single_child_pins(tv, prev_cols: list[int]) -> list[_Pin]:
    """단일 자식 그룹의 항등 고정 — C2는 오탐 방지로 자식 2개 이상만 검사하므로
    (이익잉여금 = 미처분이익잉여금 단독 같은) 항등 관계가 결과에 없다.
    전기 열에서 부모값 == 유일한 하위 자식값이 '정확히' 확인될 때만 고정."""
    pins: list[_Pin] = []
    body = tv.body
    for gi, i in enumerate(body):
        d = tv.label_depth(i)
        span = []
        for j in body[gi + 1:]:
            if not tv.row_label(j):
                continue
            if tv.label_depth(j) <= d:
                break
            span.append(j)
        valued = [j for j in span
                  if any(tv.value(j, c) is not None for c in prev_cols)]
        if len(valued) != 1:
            continue
        child = valued[0]
        for pc in prev_cols:
            pv = tv.value(i, pc)
            if pv is None:
                continue
            if pv == 0:
                continue  # 0 == 0 항등은 구조 정보가 없음 — 고정 근거 불가
            for cc2 in prev_cols:
                cv = tv.value(child, cc2)
                if cv is not None and cv == pv:
                    pins.append(_Pin(
                        row_id=tv.table.rows[i].id, grid_col=pc,
                        terms=[(tv.table.rows[child].id, cc2, 1)],
                        row_label=tv.row_label(i)[:30], check="C2(단일)"))
                    break
            else:
                continue
            break
    return pins


def refoot_plan(doc: Document, report: CheckReport, row_positions: dict,
                profile=None) -> tuple[list[RefootItem], list[RefootSkip]]:
    items: list[RefootItem] = []
    skips: list[RefootSkip] = []
    pos_of = _cell_positions(doc, row_positions)

    with use_profile(profile):
        views, _ = build_views(doc)
        by_table: dict[str, list] = {}
        for x in report.results:
            if x.check in _FOOTING and x.axis == "col" and x.ok and x.terms:
                by_table.setdefault(x.table_id, []).append(x)

        for tv in views:
            results = by_table.get(tv.table.id)
            if not results:
                continue
            cur = [c for c in range(tv.ncols) if tv.period_bucket(c) == "당기"]
            prev = [c for c in range(tv.ncols) if tv.period_bucket(c) == "전기"]
            if not cur or len(cur) != len(prev):
                continue  # 페어 불가 표 — 롤포워드와 동일 기준
            pair = dict(zip(prev, cur))
            row_index = {r.id: i for i, r in enumerate(tv.table.rows)}

            # 전기 열에서 통과한 검산만 — 당기 열에 같은 구조 적용. 제외 규칙:
            # - 전기 구성값이 전부 0 → 부호·구조 미결정(0=0은 아무 후보나 매치)
            # - 무들여쓰기(C2b) footing → 열마다 트리가 정당하게 다를 수 있어
            #   (기수 간 구성 변동 실측) 열 간 전이 금지 — 그 문서 유형은
            #   check의 기준열 적용·강등 로직이 담당
            def exact_and_determinate(x) -> bool:
                """구조 고정 자격 — 전기 열 자체 재검산이 '정확히' 일치해야 한다.
                check는 ±1 절사·손실 절대값 매치 같은 보정 경로로도 통과시키지만,
                그 결과의 구조는 부호·단수가 모호해 쓰기 근거로 부적합."""
                total = 0
                nonzero = False
                for rid, gc, sign in x.terms:
                    if rid not in row_index:
                        return False
                    v = tv.value(row_index[rid], gc)
                    if v is None:
                        return False
                    if v != 0:
                        nonzero = True
                    total += sign * v
                i = row_index.get(x.row_id)
                return (nonzero and i is not None
                        and tv.value(i, x.grid_col) == total)

            pinned: list = [x for x in results if x.grid_col in pair
                            and all(gc in pair for _, gc, _ in x.terms)
                            and "무들여쓰기" not in x.note
                            and exact_and_determinate(x)]
            # 보충: 단일 자식 항등(전기에서 정확 확인) — check 결과에 없는 고리
            seen_keys = {(x.row_id, x.grid_col) for x in pinned}
            for pin in _single_child_pins(tv, list(pair)):
                if (pin.row_id, pin.grid_col) not in seen_keys:
                    pinned.append(pin)
            if not pinned:
                continue
            pinned.sort(key=lambda x: row_index.get(x.row_id, 0))

            # 당기 열 값 상태 (제안이 확정한 값을 반영해가며 고정점 반복)
            vals: dict[tuple[str, int], object] = {}
            for i, row in enumerate(tv.table.rows):
                if row.in_thead:
                    continue
                for cc in cur:
                    v = tv.value(i, cc)
                    if v is not None:
                        vals[(row.id, cc)] = v

            proposed: dict[tuple[str, int], object] = {}
            for _ in range(_MAX_PASSES):
                changed = False
                for x in pinned:
                    cc = pair[x.grid_col]
                    key = (x.row_id, cc)
                    i = row_index[x.row_id]
                    if tv.cell(i, cc) is None or tv.cell(i, cc).readonly:
                        continue
                    term_vals = []
                    missing = 0
                    for rid, gc, sign in x.terms:
                        v = vals.get((rid, pair[gc]))
                        if v is None:
                            missing += 1
                        else:
                            term_vals.append(sign * v)
                    if missing:
                        if key not in proposed:
                            skips.append(RefootSkip(
                                tv.context, x.row_label,
                                f"구성 {len(x.terms)}개 중 {missing}개 미입력"))
                        continue
                    new = sum(term_vals)
                    if vals.get(key) == new:
                        continue
                    vals[key] = new
                    proposed[key] = (x, new)
                    changed = True
                if not changed:
                    break

            for (row_id, cc), (x, new) in proposed.items():
                loc = pos_of.get(row_id)
                if loc is None:
                    continue
                i = row_index[row_id]
                old = tv.value(i, cc)
                if isinstance(new, Decimal) and new == new.to_integral_value():
                    new = int(new)
                items.append(RefootItem(
                    sheet=loc[0], xlsx_row=loc[1], xlsx_col=2 + cc,
                    context=tv.context, label=x.row_label, check=x.check,
                    old=old, new=new, n_terms=len(x.terms)))

    # 같은 셀을 다뤘던 스킵은 제안으로 해소됐으면 제거
    resolved = {(it.context, it.label) for it in items}
    skips = [s for s in skips if (s.context, s.label) not in resolved]
    return items, skips


def apply_refoot(xlsx_path: str, items: list[RefootItem]) -> int:
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path)
    n = 0
    for it in items:
        cell = wb[it.sheet].cell(row=it.xlsx_row, column=it.xlsx_col)
        cell.value = float(it.new) if isinstance(it.new, Decimal) else it.new
        n += 1
    wb.save(xlsx_path)
    return n
