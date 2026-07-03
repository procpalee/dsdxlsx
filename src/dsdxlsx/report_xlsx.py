"""check 결과 → 엑셀 리포트.

시트 구성:
- 요약: 검사별 검사 수·일치·불일치
- 검증결과: 검사한 셀 전 항목 — 일치는 TRUE, 불일치는 차이 금액(빨강). 대상이
  편집용 xlsx면 해당 셀로 가는 하이퍼링크 포함 (불일치 우선 정렬)
- 대사·표기: C5(주석 간 대사)·C6(표기) 발견 사항 목록

편집용 xlsx 안에 시트를 심지 않고 별도 파일로 쓴다 — 편집 파일이 열려 있어도
동작하고, import 스캔을 오염시키지 않는다.
"""
from __future__ import annotations

from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter

from .check import CHECK_NAMES, CheckReport
from .export_xlsx import (
    FILL_HEAD, FONT, FONT_HEAD, FONT_LINK, FONT_TITLE, MAP_SHEET, THIN, _C, _T,
)

_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
_FILL_BAD = PatternFill("solid", fgColor="FDE9E9")
_FONT_BAD = Font(name=_T["font"], size=_T["size_body"], bold=True, color=_C["accent"])
_NUMFMT = "#,###,##0;[Red](#,###,##0);-"


def build_id_map(xlsx_path: str) -> dict[str, tuple[str, int]]:
    """편집용 xlsx의 숨김 A열 스캔: 행 ID → (시트, 행 번호)."""
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    id_map: dict[str, tuple[str, int]] = {}
    for ws in wb.worksheets:
        if ws.title == MAP_SHEET:
            continue
        for rnum, row in enumerate(ws.iter_rows(min_col=1, max_col=1), start=1):
            v = row[0].value if row else None
            if v:
                id_map.setdefault(str(v).strip(), (ws.title, rnum))
    wb.close()
    return id_map


def _num(value):
    if isinstance(value, Decimal):
        return float(value)
    return value


def _write_header(ws, row: int, headers: list[str]):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = FONT_HEAD
        cell.fill = FILL_HEAD
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")


def write_report(report: CheckReport, out_path: str, target_xlsx: str | None = None) -> dict:
    id_map = build_id_map(target_xlsx) if target_xlsx else {}
    wb = Workbook()

    # ── 요약 ─────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "요약"
    ws.column_dimensions["B"].width = 34
    for col, w in (("C", 10), ("D", 10), ("E", 10), ("F", 12)):
        ws.column_dimensions[col].width = w
    title = ws.cell(row=2, column=2, value="DSD 정합성 검증 결과")
    title.font = FONT_TITLE
    _write_header(ws, 4, ["", "검사", "검사 수", "일치", "불일치", "참고(info)"])
    ws.cell(row=4, column=1).value = None
    r = 5
    for key, name in CHECK_NAMES.items():
        results = [x for x in report.results if x.check == key]
        infos = [f for f in report.findings if f.check == key and f.level == "info"]
        checked = report.counts.get(key, 0)
        bad = sum(1 for x in results if not x.ok)
        row_vals = [f"{key} {name}", checked, len(results) - bad if results else "",
                    bad if results else "", len(infos)]
        for c, v in enumerate(row_vals, start=2):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = _FONT_BAD if (c == 5 and isinstance(v, int) and v > 0) else FONT
            cell.border = _BORDER
            if c > 2:
                cell.alignment = Alignment(horizontal="center")
        r += 1
    note = ws.cell(row=r + 1, column=2,
                   value="모든 결과는 참고용 — 최종 확인은 DART 편집기에서.")
    note.font = FONT

    # ── 검증결과 (C1~C4 전 항목) ─────────────────────────────────
    ws = wb.create_sheet("검증결과")
    widths = {"A": 6, "B": 22, "C": 7, "D": 30, "E": 22, "F": 18, "G": 18, "H": 16, "I": 18, "J": 26}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    _write_header(ws, 1, ["검사", "구역", "표", "행", "열", "기재값", "계산값", "결과", "셀", "비고"])
    ws.freeze_panes = "A2"
    ordered = sorted(report.results, key=lambda x: (x.ok, x.check, x.table_id))
    r = 2
    for x in ordered:
        vals = [x.check, x.context, x.table_id, x.row_label, x.col_header,
                _num(x.stated), _num(x.computed)]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = FONT
            cell.border = _BORDER
            if c in (6, 7) and isinstance(v, (int, float)):
                cell.number_format = _NUMFMT
        # 결과: 일치 → TRUE / 불일치 → 차이 금액(계산 가능할 때) 또는 '미확인'
        res = ws.cell(row=r, column=8)
        res.border = _BORDER
        res.alignment = Alignment(horizontal="center")
        if x.ok:
            res.value = True
            res.font = FONT
        else:
            if x.computed is not None:
                res.value = _num(x.stated) - _num(x.computed)
                res.number_format = _NUMFMT
            else:
                res.value = "미확인"
            res.font = _FONT_BAD
            for c in range(1, 11):
                ws.cell(row=r, column=c).fill = _FILL_BAD
        # 셀 하이퍼링크 (편집용 xlsx 대상일 때)
        loc = id_map.get(x.row_id)
        cell_ref = ws.cell(row=r, column=9)
        cell_ref.border = _BORDER
        if loc:
            sheet, rownum = loc
            addr = f"{get_column_letter(2 + x.grid_col)}{rownum}"
            cell_ref.value = f"{sheet}!{addr}"
            cell_ref.hyperlink = f"{target_xlsx}#'{sheet}'!{addr}"
            cell_ref.font = FONT_LINK
        else:
            cell_ref.value = x.row_id
            cell_ref.font = FONT
        note_cell = ws.cell(row=r, column=10, value=x.note or None)
        note_cell.font = FONT
        note_cell.border = _BORDER
        r += 1
    ws.auto_filter.ref = f"A1:J{max(r - 1, 1)}"

    # ── 대사·표기 (C5·C6 + 기타 발견 사항) ───────────────────────
    ws = wb.create_sheet("대사·표기")
    for col, w in (("A", 6), ("B", 8), ("C", 46), ("D", 90)):
        ws.column_dimensions[col].width = w
    _write_header(ws, 1, ["검사", "수준", "위치", "내용"])
    ws.freeze_panes = "A2"
    r = 2
    for f in report.findings:
        if f.check in ("C1", "C2", "C3", "C4"):
            continue  # 검증결과 시트가 담당
        for c, v in enumerate([f.check, f.level, f.where, f.message], start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = _FONT_BAD if f.level == "warn" else FONT
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(c == 4))
        r += 1
    ws.auto_filter.ref = f"A1:D{max(r - 1, 1)}"

    if _T.get("modernize"):
        _T["modernize"](wb)
    wb.save(out_path)
    bad = sum(1 for x in report.results if not x.ok)
    return {"results": len(report.results), "bad": bad}
