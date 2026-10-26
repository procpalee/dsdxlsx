"""DSD ↔ 조서 대사(recon) — 라벨 키, 값 비교, 합성 분기보고서 × 합성 조서."""
from __future__ import annotations

import openpyxl
import pytest

from dsdxlsx import recon as R
from dsdxlsx.parse import parse_contents

from test_rollinterim import _doc_1q_2026


def _wp(tmp_path, cash_cur=1200, split_cash=False):
    wb = openpyxl.Workbook()
    bs = wb.active
    bs.title = "BS"
    bs.append(["계정과목", None, "2026.03.31", "2025.12.31"])
    if split_cash:
        bs.append(["현금및현금성자산(보통예금)", None, 700, 600])
        bs.append(["현금및현금성자산(정기예금)", None, 500, 400])
    else:
        bs.append(["현금및현금성자산", None, cash_cur, 1000])
    bs.append(["대손충당금", None, -100, -90])
    bs.append(["자산총계", None, 1200, 1000])
    pl = wb.create_sheet("IS")
    pl.append(["계정과목", None, "2026.03.31", "2025.03.31"])
    pl.append([None, None, "누적", "누적"])
    pl.append(["매출액", None, 2000, 1800])
    pl.append(["매출원가", None, -900, -700])
    pl.append(["분기순이익", None, 200, 150])
    nt = wb.create_sheet("주석")
    nt.append([None, "1. 임대수익"])
    nt.append([None, "구분", "당분기"])
    nt.append([None, "임대수익", 30])
    nt.append([None, "2. 보증금"])
    nt.append([None, "보증금", 10])
    path = tmp_path / "8110.xlsx"
    wb.save(path)
    return str(path)


def _find(rep, sheet, label, period, col_part=""):
    for x in rep.items:
        if x.sheet.startswith(sheet) and label in x.label and x.period == period and col_part in x.col:
            return x
    raise AssertionError(f"칸 없음: {sheet} {label} {period} {col_part}")


def test_label_keys_strip_notes_dates_and_loss_alias():
    assert "지분법적용투자주식" in R.label_keys(["지분법적용투자주식(주석4)"])
    assert R.label_keys(["Ⅹ. 분기순손실"]) & R.label_keys(["당기순이익"])
    assert R.label_keys(["2026.01.01(당기초)"]) & R.label_keys(["당기초(보고금액)"])
    assert R.label_keys(["미수금의 감소(증가)"]) & R.label_keys(["미수금"])


@pytest.mark.parametrize("v,w,unit,want", [
    (1623, 1622855, 1000, "반올림"),      # 천원 DSD ↔ 원 조서
    (6334710, 6334709, 1000, "반올림"),   # 천원 반올림 ±1
    (27.27, 0.2727, 1, "같음"),           # 지분율
    (-56, -55.538, 1, "반올림"),          # 주당손익
    (1382453464, -1382453464, 1, "부호"),  # 결손금 양수 표기
    (0, 0.2727, 1, None),                 # 0 은 반올림으로 맞추지 않음
    (100, 101, 1, None),
])
def test_compare(v, w, unit, want):
    assert R.compare(v, w, unit, "") == want


def test_reconcile_matches_statements_and_notes(tmp_path):
    doc = parse_contents(_doc_1q_2026())
    rep = R.reconcile(doc, R.load_workpaper(_wp(tmp_path)))
    assert rep.period == "2026년 1분기"
    assert _find(rep, "재무상태표", "현금및현금성자산", "당기").verdict == "일치"
    assert _find(rep, "재무상태표", "현금및현금성자산", "전기").verdict == "일치"
    assert _find(rep, "재무상태표", "대손충당금", "전기").where == "BS!D3"
    assert _find(rep, "포괄손익계산서", "매출액", "전기", "누적").verdict == "일치"
    assert _find(rep, "주석1", "임대수익", "당기").verdict == "일치"
    # 조서 주석 시트에는 당기만 — 전기 칸은 대사 대상 아님
    assert _find(rep, "주석1", "임대수익", "전기").verdict == "대사 대상 아님"


def test_reconcile_flags_mismatch_with_workpaper_value(tmp_path):
    doc = parse_contents(_doc_1q_2026())
    rep = R.reconcile(doc, R.load_workpaper(_wp(tmp_path, cash_cur=1199)))
    x = _find(rep, "재무상태표", "현금및현금성자산", "당기")
    assert x.verdict == "불일치" and x.wp == 1199 and x.where == "BS!C2"


def test_reconcile_sums_sub_accounts(tmp_path):
    doc = parse_contents(_doc_1q_2026())
    rep = R.reconcile(doc, R.load_workpaper(_wp(tmp_path, split_cash=True)))
    x = _find(rep, "재무상태표", "현금및현금성자산", "당기")
    assert x.verdict == "일치" and x.how.startswith("하위 합")


def _wp_pa(tmp_path):
    """PA 정산표 모양: (정산) · (공시) 시트, 기수 머리글, 병합 '3개월 | 누적', 주석마다 시트 + 소제목 기간."""
    wb = openpyxl.Workbook()
    bad = wb.active
    bad.title = "BS(정산)"                    # 공시 시트가 있으면 정산 시트는 안 봄
    bad.append(["과목", "2026.03.31", "2025.12.31"])
    bad.append(["현금및현금성자산", 1, 1])
    bs = wb.create_sheet("BS(공시)")
    bs.append([None, "BS | 재무상태표(공시)"])
    bs.append([None, "제 60(당) 1분기말  2026년 03월 31일 현재"])
    bs.append([None, "회사", None, None, "(단위: 원)"])
    bs.append([None, "과목", "주 석", "제 60(당) 1분기말", "제 59(전) 기말"])
    bs.append([None, "현금및현금성자산", 3, 1200, 1000])     # 주석 번호 3 은 값이 아님
    bs.append([None, "대손충당금", None, -100, -90])
    pl = wb.create_sheet("IS(공시)")
    pl.append([None, "과목", "제 60(당) 1분기", None, "제 59(전) 1분기", None])
    pl.append([None, None, "3개월", "누적", "3개월", "누적"])
    pl.append([None, "매출액", 2000, 2000, 1800, 1800])
    pl.append([None, "매출원가", 500, 900, 400, 700])
    idx = wb.create_sheet("주석 Index")
    idx.append([None, 1, "임대수익"])
    n1 = wb.create_sheet("1. 임대수익")
    n1.append([None, "1. 임대수익"])
    n1.append([None, "구분", "당분기", "전분기"])
    n1.append([None, "임대수익", 30, 25])
    n2 = wb.create_sheet("2. 보증금")
    n2.append([None, "① 당분기말", None, "(단위: 원)"])
    n2.append([None, "구분", "금액"])
    n2.append([None, "보증금", 10])
    n2.append([None, "② 전기말"])
    n2.append([None, "구분", "금액"])
    n2.append([None, "보증금", 9])
    path = tmp_path / "정산표.xlsx"
    wb.save(path)
    return str(path)


def test_classify_sheets_public_over_settlement_and_note_each(tmp_path):
    sheets = R.load_workpaper(_wp_pa(tmp_path))
    kinds = {s.name: (s.kind, s.fs_kind) for s in sheets}
    assert kinds == {"BS(공시)": ("fs", "BS"), "IS(공시)": ("fs", "PL"),
                     "1. 임대수익": ("note", ""), "2. 보증금": ("note", "")}
    cfg = R.sheet_config(sheets)
    assert cfg["재무제표"] == {"BS": ["BS(공시)"], "PL": ["IS(공시)"]}
    assert cfg["주석"] == ["1. 임대수익", "2. 보증금"]


def test_reconcile_pa_layout(tmp_path):
    doc = parse_contents(_doc_1q_2026())
    rep = R.reconcile(doc, R.load_workpaper(_wp_pa(tmp_path)))
    x = _find(rep, "재무상태표", "현금및현금성자산", "당기")
    assert x.verdict == "일치" and x.where == "BS(공시)!D5"
    assert _find(rep, "재무상태표", "대손충당금", "전기").where == "BS(공시)!E6"
    assert _find(rep, "포괄손익계산서", "매출원가", "전기", "누적").where == "IS(공시)!F4"
    assert _find(rep, "포괄손익계산서", "매출원가", "당기", "3개월").where == "IS(공시)!C4"
    assert _find(rep, "주석1", "임대수익", "전기").verdict == "일치"
    # 소제목 '② 전기말' 아래 표 = 전기
    assert _find(rep, "주석2", "보증금", "전기").where == "2. 보증금!C6"
    assert _find(rep, "주석2", "보증금", "당기").where == "2. 보증금!C3"


def test_consolidated_ignores_entity_columns(tmp_path):
    """연결정산표: 법인별 열(머리글에 날짜 없음)은 대사하지 않고 연결(날짜) 열만."""
    wb = openpyxl.Workbook()
    bs = wb.active
    bs.title = "BS"
    bs.append(["계정과목", "갑", "을", "조정", "2026.03.31", "2025.12.31"])
    bs.append(["현금및현금성자산", 1200, 300, 0, 1199, 1000])
    path = tmp_path / "연결정산표.xlsx"
    wb.save(path)
    doc = parse_contents(_doc_1q_2026())
    x = _find(R.reconcile(doc, R.load_workpaper(str(path))), "재무상태표", "현금및현금성자산", "당기")
    assert x.verdict == "불일치" and x.where == "BS!E2"


def test_config_excludes_and_aliases(tmp_path):
    doc = parse_contents(_doc_1q_2026())
    cfg = R.ReconConfig.from_dict({"빼는 칸": [{"구획": "재무상태표", "행": "대손충당금"}],
                                   "이름 다른 행": {"현금및현금성자산": "현금"}})
    wb = openpyxl.Workbook()
    bs = wb.active
    bs.title = "BS"
    bs.append(["계정과목", "2026.03.31", "2025.12.31"])
    bs.append(["현금", 1200, 1000])
    path = tmp_path / "w.xlsx"
    wb.save(path)
    rep = R.reconcile(doc, R.load_workpaper(str(path)), config=cfg)
    assert _find(rep, "재무상태표", "대손충당금", "당기").verdict == "제외(설정)"
    x = _find(rep, "재무상태표", "현금및현금성자산", "당기")
    assert x.verdict == "일치" and x.how == "같은 행"


@pytest.mark.parametrize("text,want", [
    ("제 40(당) 기말", "당기"), ("전분기", "전기"), ("당기순이익", None), ("3개월 미만", None),
])
def test_header_token(text, want):
    t = R._hdr_token(text)
    assert (t[1] if t else None) == want


def test_report_written(tmp_path):
    doc = parse_contents(_doc_1q_2026())
    rep = R.reconcile(doc, R.load_workpaper(_wp(tmp_path, cash_cur=1199)))
    out = tmp_path / "out.xlsx"
    stats = R.write_recon_report(rep, str(out), "a.dsd", _wp(tmp_path))
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["요약", "확인 필요", "대사결과"]
    assert stats["issues"] >= 1
