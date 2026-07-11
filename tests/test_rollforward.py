"""A3 롤포워드 — 기수 +1, 당기→전기 값 이동, 연도 선택 치환, 3등급 리포트."""
from openpyxl import load_workbook

from dsdxlsx import rollforward, splice
from dsdxlsx.check import run_checks
from dsdxlsx.cli import main
from dsdxlsx.parse import parse_contents

from fixtures import factory


def _rolled(contents):
    doc = parse_contents(contents)
    ops, report = rollforward.roll_ops(doc)
    return parse_contents(splice.apply(doc.data, ops)), report


def test_roll_moves_values_and_bumps_gisu():
    doc2, report = _rolled(factory.standard_contents())
    texts = [c.text for t in doc2.tables.values() for r in t.rows for c in r.cells]
    # 기수 +1
    assert any("제 3(당) 기말" in t for t in texts)
    assert any("제 2(전) 기말" in t for t in texts)
    # BS 현금 행: 전기열 = 구 당기값(1,000), 당기열 = 빈 값
    bs = doc2.tables["t002"]
    cash = next(r for r in bs.rows if r.cells and "현금및현금성자산" in r.cells[0].text)
    assert cash.cells[2].text == "" and cash.cells[3].text == "1,000"
    # 주석 표도 동일 (주석4 합계: 전기 ← 1,000)
    n4 = doc2.tables["t003"]
    total = next(r for r in n4.rows if r.cells and "합" in r.cells[0].text)
    assert total.cells[1].text == "" and total.cells[2].text == "1,000"
    # 리포트: 값 이동 자동확정 + 미처리 없음
    assert report.count("자동확정") >= 4
    assert report.count("미처리") == 0

    # 이동된 전기열은 구 당기값이므로 footing·대사가 그대로 클린이어야 한다
    check = run_checks(doc2)
    assert not [f for f in check.findings if f.level == "warn"]


def test_roll_year_bump_is_selective():
    contents = factory.document(
        factory.p("표지"),
        factory.title("재무제표", 2),
        factory.statement_title_table("재 무 상 태 표"),
        factory.table(
            [["과목", "제 2(당) 기말", "제 1(전) 기말"]],
            [["현금", "1,000", "900"], ["자산총계", "1,000", "900"]],
            unit="원",
        ),
        factory.title("주석", 3),
        # 기수+보고기간 연도+역사 연도가 한 문단에
        factory.p("1. 일반사항\n제 2(당) 기는 2025년 12월 31일로 종료하며, "
                  "회사는 1994년 11월 상장하였습니다."),
    )
    doc2, report = _rolled(contents)
    para = next(n for n in doc2.nodes.values() if "상장" in n.text)
    assert "제 3(당) 기" in para.text          # 기수 +1
    assert "2026년 12월 31일" in para.text      # 보고기간 연도 +1
    assert "1994년 11월 상장" in para.text      # 역사적 연도 불변
    # 기수 없는 노드의 날짜는 손대지 않고 확인필요로
    assert any(i.grade == "확인필요" for i in report.items)


def test_roll_unpaired_table_untouched():
    contents = factory.document(
        factory.p("표지"),
        factory.title("재무제표", 2),
        factory.statement_title_table("재 무 상 태 표"),
        factory.table(  # 당기 열만 있는 표 — 미처리
            [["과목", "제 2(당) 기말"]],
            [["현금", "1,000"], ["자산총계", "1,000"]],
            unit="원",
        ),
    )
    doc = parse_contents(contents)
    ops, report = rollforward.roll_ops(doc)
    doc2 = parse_contents(splice.apply(doc.data, ops))
    cash = next(r for t in doc2.tables.values() for r in t.rows
                if r.cells and r.cells[0].text == "현금")
    assert cash.cells[1].text == "1,000"  # 값 무변경
    assert any(i.grade == "미처리" and "대칭" in i.note for i in report.items)


def test_roll_cli_outputs_and_report_sheet(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "전기.dsd").write_bytes(factory.dsd_bytes(factory.standard_contents()))
    rc = main(["rollforward", "전기.dsd"])
    out = capsys.readouterr().out
    assert rc == 0 and "자동확정" in out
    assert (tmp_path / "전기_차기기초.dsd").exists()
    xlsx = tmp_path / "전기_차기기초.xlsx"
    assert xlsx.exists()
    wb = load_workbook(str(xlsx))
    assert wb.sheetnames[0] == "롤포워드"
    vals = [c.value for row in wb["롤포워드"].iter_rows() for c in row if c.value]
    assert any("자동확정" == v for v in vals)

    # 리포트 시트(!r ID)가 있어도 무수정 import는 변경 없음
    from dsdxlsx.import_xlsx import import_xlsx
    _, _, rep = import_xlsx(str(xlsx))
    assert not rep.changed
