"""분·반기 당전기 밀기 — 직전 중간보고서 DSD(골격)로 다음 분기 DSD 를 만든다.

입력
- A 골격: 직전 중간보고서(1Q→2Q, 2Q→3Q, 3Q→다음 해 1Q). 최근 풋팅에서 고친
  표 구조·문구가 그대로 살아야 하므로 이 파일이 뼈대다.
- B 전년 동기(선택): 흐름표(손익·현금흐름·자본변동·흐름형 주석)의 전기 값 원천.
  B 의 '당기' 열 → A 의 '전기' 열.
- C 전기 감사보고서(선택, 3Q→1Q 해 넘김일 때): 재무상태표·재무상태형 주석의
  전기말(직전 연말) 값 원천. 같은 해 전환에서는 A 의 전기말이 이미 직전 연말이라 유지.
- D DART(상장사): 본 재무제표(BS/PL/CF) 전기 값 대사. B 가 없으면 PL/CF 전기 값을
  DART 로 채운다(라벨 + 값 앵커 매칭). 비상장·키 없음이면 생략하고 사유만 남긴다.

원칙 (rollforward.py 와 동일)
- 바이트 스플라이스: 손대는 셀만 inner 치환, 행 블록 교체는 원시 바이트 클론.
- 3등급 리포트: 자동확정 / 확인필요 / 미처리. 값이 틀린 채 남는 것보다 빈 칸이
  안전하므로, 원천을 못 찾은 전기 칸은 비우고 확인필요로 남긴다.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field

from . import dart_client as D
from . import period as P
from . import splice
from .check import TableView, cell_value, is_total_label, norm, use_profile
from .export_xlsx import STATEMENT_KEYWORDS, Sec, refined_sections
from .model import Document, Node
from .numbers import classify, to_dsd_text
from .rollforward import RollReport, _column_pairs, _final_painted, _numeric_cells

# 기수 라벨: '제 60(당) 반기' / '제 60 (당) 3분기' / '제 21(당) 기 3분기' / '제 59(전) 기말'
# (rollforward._GISU_RE 는 '(당) 기' 바로 이어지는 꼴만 잡아 분·반기 표기를 놓친다)
_GISU_RE = re.compile(r"(제\s*)(\d{1,3})(?=\s*\(?[당전]?\)?\s*(?:기|\d\s*분\s*기|반\s*기))")

# 시트 종류
KIND_OF_TITLE = {"재무상태표": "BS", "손익계산서": "PL", "포괄손익계산서": "PL",
                 "자본변동표": "SCE", "현금흐름표": "CF",
                 "이익잉여금처분계산서": "RE", "결손금처리계산서": "RE"}
_BALANCE_HDR_RE = re.compile(r"전기말|전\)기말|\(전\)기$|제\d+\(전\)기$")
# 캡션 '전기' 표 = 연간(직전 사업연도 - 분 · 반기 보고서는 전년 동기를 '전분기 · 전반기'로 따로 씀)
_ANNUAL_MARK_RE = re.compile(r"[<(\[]?전\s*기[>)\]]?")
_ACCT_COL_RE = re.compile(r"당기순|당기손익|전기순|전기손익|전기오류")
_CUR_MARK_RE = re.compile(r"당\s*\d?\s*(분|반)?\s*기\s*말?\s*$")
_PREV_MARK_RE = re.compile(r"전\s*\d?\s*(분|반)?\s*기\s*말?\s*$")
_BAL_MARK_RE = re.compile(r"전\s*기\s*말\s*$")
_OPEN_RE = re.compile(r"기초")
_CLOSE_RE = re.compile(r"기말")
_LEAF_3M_RE = re.compile(r"3\s*개월")
_LEAF_CUM_RE = re.compile(r"누\s*적")
# 자본변동표 블록 라벨 플립 (B 당기 블록 → A 전기 블록)
_FLIP = [("당분기말", "전분기말"), ("당반기말", "전반기말"), ("당기초", "전기초"),
         ("당기말", "전기말"), ("당분기", "전분기"), ("당반기", "전반기")]


@dataclass
class DartRow:
    kind: str
    where: str
    label: str
    col: str
    file_value: int | None
    dart_value: int | None
    dart_name: str
    how: str          # 라벨 | 값앵커 | 미매칭
    verdict: str      # 일치 | 불일치 | 미매칭 | 채움


@dataclass
class DartReport:
    period_desc: str = ""
    skipped_reason: str = ""
    rows: list[DartRow] = field(default_factory=list)
    unmatched_dart: list[tuple[str, str, int | None]] = field(default_factory=list)

    def count(self, verdict: str) -> int:
        return sum(1 for r in self.rows if r.verdict == verdict)


@dataclass
class TableRec:
    """구획 문맥이 붙은 표 뷰."""
    sec: Sec
    tv: TableView
    kind: str            # BS/PL/CF/SCE/RE/NOTE/OTHER
    cur: list[int]
    prev: list[int]
    caption: str = ""    # 세로형 표: 바로 앞 문단 마지막 줄('(1) 당반기말')
    data: bytes = b""    # 이 표가 속한 문서의 원시 바이트 (행 블록 클론용)

    @property
    def table(self):
        return self.tv.table

    @property
    def caption_is_cur(self) -> bool:
        return bool(self.caption) and bool(_CUR_MARK_RE.search(norm(self.caption)))

    @property
    def where(self) -> str:
        return self.tv.context

    @property
    def is_balance(self) -> bool:
        if self.kind == "BS":
            return True
        if self.kind in ("PL", "CF", "SCE", "RE"):
            return False
        if self.prev:
            h = norm(self.tv.col_header(self.prev[0]))
            # '전기'만 쓴 열 = 직전 사업연도(분 · 반기 보고서는 전년 동기를 '전분기 · 전반기'로 씀)
            return bool(_BALANCE_HDR_RE.search(h)) or h.split("/")[-1] in ("전기", "전기(주1)")
        if self.caption:
            return bool(_BAL_MARK_RE.search(norm(self.caption)))
        return False


# ── 문서 스캔 ────────────────────────────────────────────────────────

def _skip_sec(sec: Sec) -> bool:
    t = norm(sec.title)
    return sec.note_no is None and not sec.notes_intro and t not in STATEMENT_KEYWORDS


def scan(doc: Document) -> tuple[list[TableRec], list[Sec]]:
    secs = refined_sections(doc)
    recs: list[TableRec] = []
    for sec in secs:
        if _skip_sec(sec):
            continue
        kind = KIND_OF_TITLE.get(sec.title, "NOTE" if (sec.note_no or sec.notes_intro) else "OTHER")
        label = f"주석{sec.note_no} {sec.note_title[:12]}" if sec.note_no else sec.title
        last_mark = ""
        for block in sec.blocks:
            if block.kind == "p":
                lines = [l.strip() for l in block.item.text.split("\n") if l.strip()]
                last_mark = lines[-1] if lines else last_mark
                continue
            if block.kind != "table":
                continue
            t = block.item
            tv = TableView(t, label, sec_title=sec.title)
            # 기간 열: 0열(이름 열 - 왼쪽 위 '전기' · '당분기말'은 표 전체 기간)과
            # 계정 이름 머리글('당기순손익' · '당기손익')은 빼고
            cur = [c for c in range(1, tv.ncols) if tv.period_bucket(c) == "당기"
                   and not _ACCT_COL_RE.search(norm(tv.col_header(c)))]
            prev = [c for c in range(1, tv.ncols) if tv.period_bucket(c) == "전기"
                    and not _ACCT_COL_RE.search(norm(tv.col_header(c)))]
            caption = ""
            top_left = norm(tv.col_header(0)) if tv.ncols else ""
            if not cur and not prev and (_CUR_MARK_RE.search(top_left) or _PREV_MARK_RE.search(top_left)):
                caption = top_left                 # '전기 | 자산총액 | 부채총액 …' 표
            elif not cur and not prev and (_CUR_MARK_RE.search(last_mark)
                                           or _PREV_MARK_RE.search(last_mark)):
                caption = last_mark
            recs.append(TableRec(sec, tv, kind, cur, prev, caption, doc.data))
            if tv.n > 2:
                last_mark = ""
    return recs, secs


def _has_tu(rec: TableRec) -> bool:
    return any(c.tag == "TU" for r in rec.table.rows for c in r.cells)


def _row_labels(tv: TableView) -> list[str]:
    return [tv.row_label(i) for i in tv.body]


def _leaf(tv: TableView, c: int) -> str:
    """열 머리글의 마지막 조각('제 60(당) 반기/3개월' → '3개월')."""
    parts = tv.col_header(c).split("/")
    return parts[-1].strip() if parts else ""


# ── 표 매칭 (A ↔ B/C) ────────────────────────────────────────────────

def _sig(tv: TableView) -> tuple:
    heads = tuple(D.norm_label(tv.col_header(c)) for c in range(tv.ncols))
    return (tv.ncols, heads)


def _label_overlap(a: list[str], b: list[str]) -> float:
    sa = {D.canon(x) for x in a if x.strip()}
    sb = {D.canon(x) for x in b if x.strip()}
    if not sa or not sb:
        return 0.0
    return len(sa & sb) / max(len(sa), 1)


def match_notes(a_recs: list[TableRec], b_recs: list[TableRec]) -> dict[int, TableRec]:
    """A 주석 표 → B 주석 표. 주석 제목(정규화) → 표 순서·헤더 서명·행 라벨 겹침."""
    def by_note(recs):
        out: dict[str, list[TableRec]] = {}
        for r in recs:
            if r.sec.note_no is None:
                continue
            key = D.norm_label(r.sec.note_title)
            out.setdefault(key, []).append(r)
        return out

    bmap = by_note(b_recs)
    result: dict[int, TableRec] = {}
    for key, a_list in by_note(a_recs).items():
        cands = list(bmap.get(key, []))
        for a in a_list:
            if not (a.cur or a.caption or _blocks(a.tv)):
                continue
            if a.caption_is_cur:
                continue  # 당기 캡션 표는 비우기만 — 원천 불필요
            best, best_score = None, 0.0
            pool = cands
            if a.caption:
                # 전기 캡션 표의 원천 = B 의 '당기' 캡션 표
                pool = [b for b in cands if b.caption_is_cur] or cands
            for b in pool:
                score = 0.0
                if a.caption and b.tv.ncols > 1 and a.tv.ncols > 1:
                    ha = {D.norm_label(a.tv.col_header(c)) for c in range(1, a.tv.ncols)}
                    hb = {D.norm_label(b.tv.col_header(c)) for c in range(1, b.tv.ncols)}
                    if len(ha & hb) * 2 < len(ha | hb):
                        continue                   # 캡션 표는 열 머리글(0열 빼고)이 절반 이상 같아야
                if b.tv.ncols == a.tv.ncols:
                    score += 1.0
                if _sig(b.tv) == _sig(a.tv):
                    score += 1.0
                score += _label_overlap(_row_labels(a.tv), _row_labels(b.tv))
                if score > best_score and score >= 1.6:
                    best, best_score = b, score
            if best is not None:
                result[id(a)] = best
                cands.remove(best)
    return result


def match_statement(a: TableRec, b_recs: list[TableRec]) -> TableRec | None:
    cands = [b for b in b_recs if b.kind == a.kind and (b.cur or b.kind == "SCE")]
    if not cands:
        return None
    if len(cands) == 1:
        return cands[0]
    return max(cands, key=lambda b: _label_overlap(_row_labels(a.tv), _row_labels(b.tv)))


def _row_map(a_tv: TableView, b_tv: TableView) -> dict[int, int]:
    """A 행 → B 행 (정규화 라벨, 같은 라벨은 순서대로). 값 없는 구획 행도 포함."""
    def key(label: str) -> str:                    # '지급임차료(주석9,12)' ↔ '지급임차료'
        return D.canon(re.sub(r"\(\s*주석[^)]*\)", "", label))
    pool: dict[str, list[int]] = {}
    for j in b_tv.body:
        pool.setdefault(key(b_tv.row_label(j)), []).append(j)
    out: dict[int, int] = {}
    for i in a_tv.body:
        k = key(a_tv.row_label(i))
        if not k:
            continue
        lst = pool.get(k)
        if lst:
            out[i] = lst.pop(0)
    # 라벨이 어긋난 행('분기말 장부가액'↔'기말 장부가액', '당기지급액'↔'기중 지급액')은
    # 두 표의 행 수가 같고 라벨 매칭이 순서를 지킬 때만 남은 행끼리 **위치로** 잇는다.
    a_body, b_body = a_tv.body, b_tv.body
    if len(a_body) == len(b_body) and len(out) < len(a_body):
        if all(a_body.index(i) == b_body.index(j) for i, j in out.items()):
            for i, j in zip(a_body, b_body):
                if i not in out and j not in out.values() and a_tv.row_label(i).strip():
                    out[i] = j
    return out


# ── 값 셀 조작 ───────────────────────────────────────────────────────

def _clear_cells(rec: TableRec, cols: list[int], ops: list, touched: set) -> int:
    n = 0
    for i in rec.tv.body:
        for c in cols:
            cell = rec.tv.cell(i, c)
            if cell is None or cell.tag == "TU" or cell.id in touched:
                continue
            if cell.text.strip() and (cell.text.strip() == "-" or cell_value(cell.text) is not None):
                ops.append(splice.replace_inner(cell, ""))
                touched.add(cell.id)
                n += 1
    return n


def _clear_all_numeric(rec: TableRec, ops: list, touched: set) -> int:
    n = 0
    for i in rec.tv.body:
        for _, cell in _numeric_cells(rec.tv, i):
            if cell.id in touched:
                continue
            ops.append(splice.replace_inner(cell, ""))
            touched.add(cell.id)
            n += 1
    return n


def _copy_cols(a: TableRec, a_cols: list[int], b: TableRec, b_cols: list[int],
               rowmap: dict[int, int], ops: list, touched: set, painted_cells: list
               ) -> tuple[int, list[int]]:
    """B 행의 b_cols 텍스트 → A 행의 a_cols. 반환 (복사 칸 수, 원천 없는 A 행)."""
    n = 0
    missing: list[int] = []
    for i in a.tv.body:
        j = rowmap.get(i)
        row_has_num = any(a.tv.cell(i, c) is not None and a.tv.cell(i, c).id not in touched
                          and (cell_value(a.tv.cell(i, c).text) is not None
                               or a.tv.cell(i, c).text.strip() == "-")
                          for c in a_cols)
        if j is None:
            if row_has_num:
                missing.append(i)
                for c in a_cols:
                    cell = a.tv.cell(i, c)
                    if cell is not None and cell.id not in touched and cell.text.strip():
                        ops.append(splice.replace_inner(cell, ""))
                        touched.add(cell.id)
            continue
        for ac, bc in zip(a_cols, b_cols):
            cell = a.tv.cell(i, ac)
            src = b.tv.cell(j, bc)
            if cell is None or cell.tag == "TU" or cell.id in touched:
                continue
            new = src.text if src is not None else ""
            if cell.text != new:
                ops.append(splice.replace_inner(cell, new))
            touched.add(cell.id)
            if new.strip() and (cell_value(new) is not None or new.strip() == "-"):
                painted_cells.append((a.table.rows[i].start, ac))
                n += 1
    return n, missing


def _insert_b_only_rows(a: TableRec, b: TableRec, a_cols: list[int], b_cols: list[int],
                        rowmap: dict[int, int], ops: list) -> list[str]:
    """전년 동기 표에만 있고 값이 있는 행(예: 현금흐름표 '지분법이익')을 A 에 넣는다 - 전년 동기에서
    바로 위에 있던 행 다음에, 같은 열(세부 · 소계)을 쓰는 가까운 A 행을 복제해 이름 · 전기 값만 바꿈."""
    inv = {j: i for i, j in rowmap.items()}
    out: list[str] = []
    for j in b.tv.body:
        if j in inv:
            continue
        lab_cell = b.tv.cell(j, 0)
        if lab_cell is None or not lab_cell.text.strip():
            continue
        vals = [(k, b.tv.cell(j, bc)) for k, bc in enumerate(b_cols)
                if b.tv.cell(j, bc) is not None and b.tv.value(j, bc) not in (None, 0)]
        if not vals:
            continue                               # 전년 동기에도 값이 없는 행은 넣지 않음
        ups = [jj for jj in b.tv.body if jj < j and jj in inv]
        if not ups:
            continue
        anchor = inv[ups[-1]]
        cols_used = {a_cols[k] for k, _c in vals}
        order = [i for i in a.tv.body if i > anchor] + sorted((i for i in a.tv.body if i <= anchor), reverse=True)
        tpl = next((i for i in order if not is_total_label(a.tv.row_label(i))
                    and all(a.tv.cell(i, c) is not None and (a.tv.value(i, c) is not None
                                                             or a.tv.cell(i, c).text.strip() == "")
                            for c in cols_used)
                    and any(a.tv.value(i, c) is not None for c in cols_used)), None)
        if tpl is None:
            continue
        trow = a.table.rows[tpl]

        def idx(c, _row=trow, _i=tpl):
            cell = a.tv.cell(_i, c)
            return next((k for k, x in enumerate(_row.cells) if x is cell), None)
        new_texts: dict[int, str] = {}
        if (k := idx(0)) is not None:
            new_texts[k] = lab_cell.text
        for c in range(1, a.tv.ncols):
            k = idx(c)
            if k is not None and (a.tv.value(tpl, c) is not None or a.tv.cell(tpl, c).text.strip() == "-"):
                new_texts[k] = ""
        for kk, bcell in vals:
            if (k := idx(a_cols[kk])) is not None:
                new_texts[k] = bcell.text
        ops.append(splice.insert_after_row(a.table.rows[anchor], splice.clone_row(a.data, trow, new_texts)))
        out.append(lab_cell.text.strip())
    return out


# ── DART ─────────────────────────────────────────────────────────────

def _dart_field(rec: TableRec, c: int, target: P.Period) -> str:
    """A 의 전기 열 c 에 대응하는 DART 항목 필드."""
    if rec.kind == "BS":
        return "thstrm"
    if rec.kind == "PL":
        leaf = _leaf(rec.tv, c)
        if _LEAF_3M_RE.search(leaf):
            return "thstrm"
        if _LEAF_CUM_RE.search(leaf) or target.quarter != 1:
            return "thstrm_add"
        return "thstrm"
    return "thstrm"


def _fval(v, unit: int) -> int | None:
    if v is None:
        return None
    return int(round(float(v) * unit))


def dart_compare(rec: TableRec, cols: list[int], data: D.DartData, target: P.Period,
                 report: DartReport, values: dict[tuple[int, int], int | None] | None = None):
    """A 전기 열 값 ↔ DART. values 를 주면(밀린 뒤 값) 그것을 파일 값으로 쓴다."""
    items = data.by_kind(rec.kind)
    if not items:
        return
    labels = [rec.tv.row_label(i) for i in rec.tv.body]
    lm = D.match_items(labels, items)
    unit = rec.tv.unit
    used: set[int] = set()
    for idx, i in enumerate(rec.tv.body):
        for c in cols:
            if values is not None:
                fv = values.get((i, c))
            else:
                fv = _fval(rec.tv.value(i, c), unit)
            if fv is None:
                continue
            fld = _dart_field(rec, c, target)
            it = lm.get(idx)
            how = "라벨"
            if it is None:
                cand = D.anchor_match([fv], items, fld, 1)
                it = cand.get(0)
                how = "값앵커" if it else "미매칭"
            if it is None:
                if fv != 0:  # '-'(0) 행은 DART 에 없는 게 정상 — 잡음 제외
                    report.rows.append(DartRow(rec.kind, rec.where, labels[idx],
                                               _leaf(rec.tv, c), fv, None, "", "미매칭", "미매칭"))
                continue
            used.add(id(it))
            dv = getattr(it, fld)
            if dv is None and fld == "thstrm_add":
                dv = it.thstrm
            if dv is None:
                dv = 0
            ok = abs(abs(fv) - abs(dv)) <= max(unit // 2, 0)
            report.rows.append(DartRow(rec.kind, rec.where, labels[idx], _leaf(rec.tv, c),
                                       fv, dv, it.account_nm, how, "일치" if ok else "불일치"))
    for it in items:
        if id(it) not in used and (it.thstrm is not None or it.thstrm_add is not None):
            report.unmatched_dart.append((rec.kind, it.account_nm, it.thstrm))


# 손익 비용성 행(원가·비용·판관비·법인세·금융원가·기타손실) — 결과 행(이익·손익)은 제외
_EXPENSE_RE = re.compile(r"(원가|비용|판매비|판관비|법인세|손실)")
_RESULT_RE = re.compile(r"(이익|손익|순손실|주당)")
# 현금흐름표: 방향이 라벨에 있는 행(DART 는 양수 크기) vs 순액 행(DART 부호 그대로)
_CF_DIRECTIONAL_RE = re.compile(r"(취득|처분|지급|상환|증가|감소|배당|차입|납부|수취|유입|유출|회수|예치|인출)")
_CF_NET_RE = re.compile(r"(현금흐름|순증가|순감소|변동|조정|가감|순이익|순손실|기초|기말|환율)")


def _is_expense_row(label: str) -> bool:
    n = norm(label)
    return bool(_EXPENSE_RE.search(n)) and not _RESULT_RE.search(n)


def _expense_negative(rec: TableRec, cols: list[int]) -> bool:
    """손익표가 비용을 음수(괄호)로 적는 관례인가 — 비용성 행의 옛 값 부호 다수결."""
    if rec.kind != "PL":
        return False
    neg = pos = 0
    for i in rec.tv.body:
        if not _is_expense_row(rec.tv.row_label(i)):
            continue
        for c in cols:
            v = rec.tv.value(i, c)
            if v is None or v == 0:
                continue
            if v < 0:
                neg += 1
            else:
                pos += 1
    return neg > pos


def _present_sign(rec: TableRec, label: str, old_v, dv: int, expense_neg: bool) -> int:
    """DART 값(dv) → 문서 표기 부호.
    PL: DART 는 비용을 양수(비용 크기, 환입·수익은 음수)로 준다. 비용을 괄호로 적는
        문서면 부호를 뒤집고, 크기로 적는 문서면 그대로. 결과 행은 DART 부호 그대로.
    CF: DART 가 음수면 순액이므로 그대로; 양수·방향 라벨 행은 문서 옛 값의 부호를 따른다.
    BS: 옛 값이 음수(대손충당금 등)면 음수."""
    n = norm(label)
    if rec.kind == "PL":
        if _is_expense_row(label):
            return -dv if expense_neg else dv
        return dv
    if rec.kind == "CF":
        if dv < 0 or _CF_NET_RE.search(n):
            return dv
        if _CF_DIRECTIONAL_RE.search(n) and old_v is not None and old_v < 0:
            return -abs(dv)
        return dv
    # BS 등
    if old_v is not None and old_v < 0:
        return -abs(dv)
    return dv


def dart_fill(rec: TableRec, cols: list[int], data_t: D.DartData, data_base: D.DartData | None,
              base_q: int, target: P.Period, ops: list, touched: set, painted_cells: list,
              report: DartReport, roll: RollReport) -> int:
    """B 가 없을 때 A 전기 열을 DART 로 채운다.
    행 → 항목: 라벨 매칭, 실패 시 A 의 밀기 전 전기 값(전년 base 분기)과
    DART(base 분기, 전년) 값 앵커로 account_id 를 찾아 목표 분기 항목으로 잇는다."""
    items_t = data_t.by_kind(rec.kind)
    if not items_t:
        return 0
    labels = [rec.tv.row_label(i) for i in rec.tv.body]
    lm = D.match_items(labels, items_t)
    by_id = {it.account_id: it for it in items_t if it.account_id}
    items_b = data_base.by_kind(rec.kind) if data_base else []
    unit = rec.tv.unit
    expense_neg = _expense_negative(rec, cols)
    n = 0
    for idx, i in enumerate(rec.tv.body):
        old = {c: rec.tv.value(i, c) for c in cols}
        if all(v is None for v in old.values()):
            continue
        it = lm.get(idx)
        how = "라벨"
        if it is None and items_b:
            # 값 앵커: A 의 옛 전기 값(전년 base 분기) == DART base 분기 값 → 같은 account_id
            for c in cols:
                fld_b = _dart_field(rec, c, P.Period(target.year, base_q))
                cand = D.anchor_match([_fval(old[c], unit)], items_b, fld_b, 1)
                if cand:
                    hit = by_id.get(cand[0].account_id)
                    if hit is not None:
                        it, how = hit, "값앵커"
                        break
        for c in cols:
            cell = rec.tv.cell(i, c)
            if cell is None or cell.tag == "TU" or cell.id in touched:
                continue
            if it is None:
                if old.get(c) == 0:
                    # 옛 값이 '-'(0)인 행 — 그대로 두고 확인만 (DART 미매칭 목록과 대조)
                    touched.add(cell.id)
                    report.rows.append(DartRow(rec.kind, rec.where, labels[idx],
                                               _leaf(rec.tv, c), 0, None, "", "값0 유지", "미매칭"))
                    continue
                if cell.text.strip():
                    ops.append(splice.replace_inner(cell, ""))
                touched.add(cell.id)
                report.rows.append(DartRow(rec.kind, rec.where, labels[idx], _leaf(rec.tv, c),
                                           None, None, "", "미매칭", "미매칭"))
                continue
            fld = _dart_field(rec, c, target)
            dv = getattr(it, fld)
            if dv is None and fld == "thstrm_add":
                dv = it.thstrm
            if dv is None:
                new = "-" if cell.text.strip() else ""
            else:
                signed = _present_sign(rec, labels[idx], old[c], dv, expense_neg)
                val = round(signed / unit) if unit != 1 else signed
                new = to_dsd_text(val, like=cell.text) if val != 0 else "-"
            if cell.text != new:
                ops.append(splice.replace_inner(cell, new))
            touched.add(cell.id)
            painted_cells.append((rec.table.rows[i].start, c))
            n += 1
            report.rows.append(DartRow(rec.kind, rec.where, labels[idx], _leaf(rec.tv, c),
                                       None, dv, it.account_nm, how, "채움"))
    return n


# ── 행 블록(자본변동표·주석 <전기>/<당기>) ───────────────────────────

def _blocks(tv: TableView) -> tuple[tuple[int, int], tuple[int, int]] | None:
    """[전기초..전기말], [당기초..당기말] 두 블록의 (open, close) 인덱스."""
    body = tv.body
    opens = [i for i in body if _OPEN_RE.search(norm(tv.row_label(i)))
             and any(tv.value(i, c) is not None for c in range(tv.ncols))]
    closes = [i for i in body if _CLOSE_RE.search(norm(tv.row_label(i)))
              and any(tv.value(i, c) is not None for c in range(tv.ncols))]
    if len(opens) != 2 or len(closes) != 2:
        return None
    o1, o2 = opens
    c1, c2 = closes
    if not (o1 < c1 < o2 < c2):
        return None
    return (o1, c1), (o2, c2)


def _cur_block(tv: TableView) -> tuple[int, int] | None:
    """B 문서의 당기 블록 (두 블록이면 둘째, 하나면 그것)."""
    body = tv.body
    opens = [i for i in body if _OPEN_RE.search(norm(tv.row_label(i)))
             and any(tv.value(i, c) is not None for c in range(tv.ncols))]
    closes = [i for i in body if _CLOSE_RE.search(norm(tv.row_label(i)))
              and any(tv.value(i, c) is not None for c in range(tv.ncols))]
    if not opens or not closes:
        return None
    o, c = opens[-1], closes[-1]
    return (o, c) if o < c else None


def roll_block_table(a: TableRec, b: TableRec | None, same_year: bool, ops: list,
                     touched: set, painted: list, roll: RollReport,
                     open_src: dict[int, str] | None = None) -> bool:
    """자본변동표형 표. 반환 True 면 처리됨(미처리 포함 보고 완료)."""
    blk = _blocks(a.tv)
    if blk is None:
        return False
    (o1, c1), (o2, c2) = blk
    t = a.table
    if any(cell.tag == "TU" for i in a.tv.body for cell in t.rows[i].cells):
        roll.add("미처리", a.where, f"{t.id} 변동표 TU 열 포함", note="자동 필드 표는 손대지 않음")
        return True
    rows1 = t.rows[o1:c1 + 1]
    # 전기 블록 ← B 당기 블록 클론
    if b is not None and b.tv.ncols == a.tv.ncols and (cb := _cur_block(b.tv)) is not None:
        rows_b = b.table.rows[cb[0]:cb[1] + 1]
        payload = b.data[rows_b[0].start:rows_b[-1].end]
        for x, y in _FLIP:
            payload = payload.replace(x.encode(), y.encode())
        ops.append(splice.Op(pos=rows1[0].start, start=rows1[0].start,
                             end=rows1[-1].end, payload=payload))
        for row in rows1:
            for cell in row.cells:
                touched.add(cell.id)
        painted.append(("socie", t.start))
        src_note = f"전기 블록 ← 전년 동기 DSD {b.table.id} 당기 블록"
    else:
        n = 0
        for i in a.tv.body:
            if o1 <= i <= c1:
                for _, cell in _numeric_cells(a.tv, i):
                    if cell.id not in touched:
                        ops.append(splice.replace_inner(cell, ""))
                        touched.add(cell.id)
                        n += 1
        src_note = f"전기 블록 값 {n}칸 비움(전년 동기 DSD 없음 — 수기 입력)"
        roll.add("확인필요", a.where, f"{t.id} 변동표 전기 블록 원천 없음",
                 note="전년 동기 DSD 를 함께 넣으면 자동으로 채워진다")
    # 당기 블록: 기초 행 유지(같은 해) 또는 open_src 로 교체, 나머지 비움
    n_clear = 0
    for i in a.tv.body:
        if o2 < i <= c2:
            for _, cell in _numeric_cells(a.tv, i):
                if cell.id not in touched:
                    ops.append(splice.replace_inner(cell, ""))
                    touched.add(cell.id)
                    n_clear += 1
    if not same_year:
        for c, cell in _numeric_cells(a.tv, o2):
            if cell.id in touched:
                continue
            new = (open_src or {}).get(c)
            if new is None:
                ops.append(splice.replace_inner(cell, ""))
            elif cell.text != new:
                ops.append(splice.replace_inner(cell, new))
            touched.add(cell.id)
        if not open_src:
            roll.add("확인필요", a.where, f"{t.id} 당기초 잔액 원천 없음",
                     note="전기 감사보고서 DSD 를 함께 넣거나 수기 입력")
    roll.add("자동확정", a.where,
             f"{t.id} 변동표: {src_note}, 당기 변동·기말 {n_clear}칸 비움",
             note="당기초는 " + ("유지" if same_year else "전기 감사보고서/DART 기준"))
    return True


# ── 메인 ─────────────────────────────────────────────────────────────

@dataclass
class RollInputs:
    doc: Document
    base: P.Period
    target: P.Period
    doc_b: Document | None = None
    doc_c: Document | None = None
    use_dart: bool = True
    company: str = ""


def _fetch(corp: str, year: int, reprt: str, fs: str, dreport: DartReport):
    try:
        return D.fetch_statements(corp, year, reprt, fs)
    except D.DartUnavailable as e:
        if not dreport.skipped_reason:
            dreport.skipped_reason = str(e)
        return None


def roll_interim_ops(inp: RollInputs, profile=None
                     ) -> tuple[list[splice.Op], RollReport, DartReport]:
    doc, base, target = inp.doc, inp.base, inp.target
    same_year = target.year == base.year
    ops: list[splice.Op] = []
    roll = RollReport()
    dreport = DartReport(period_desc=f"{base.label} → {target.label}")
    touched: set[str] = set()
    painted: list[tuple] = []
    painted_cells: list[tuple[int, int]] = []

    with use_profile(profile):
        recs, secs = scan(doc)
        b_recs: list[TableRec] = []
        c_recs: list[TableRec] = []
        if inp.doc_b is not None:
            b_recs, _ = scan(inp.doc_b)
        if inp.doc_c is not None:
            c_recs, _ = scan(inp.doc_c)
        note_map_b = match_notes(recs, b_recs) if b_recs else {}
        note_map_c = match_notes(recs, c_recs) if c_recs else {}

        # ── DART 준비 ──
        fs = "CFS" if base.consolidated else "OFS"
        corp = base.corp_code
        if inp.use_dart and not corp and inp.company:
            corp = D.resolve_corp_code(inp.company) or ""
        d_annual = d_target = d_base = None
        if inp.use_dart:
            d_annual = _fetch(corp, target.year - 1, "11011", fs, dreport)
            d_target = _fetch(corp, target.year - 1, target.reprt_code, fs, dreport)
            if not b_recs:
                d_base = _fetch(corp, base.year - 1, base.reprt_code, fs, dreport)
            if d_annual is not None:
                dreport.period_desc += (
                    f" / 전기말 ← {target.year - 1} 사업보고서(11011)"
                    f" / 전기 ← {target.year - 1} {P.QUARTER_WORD[target.quarter]}"
                    f"({target.reprt_code}) / {fs} / corp {corp}")
        else:
            dreport.skipped_reason = "--no-dart"

        # ── ① 재무제표 본표 ──
        for rec in recs:
            if rec.kind not in ("BS", "PL", "CF", "SCE", "RE"):
                continue
            t = rec.table
            if rec.kind == "SCE":
                b = match_statement(rec, b_recs) if b_recs else None
                open_src = None
                if not same_year:
                    open_src = _sce_open_from(rec, c_recs, d_annual)
                if not roll_block_table(rec, b, same_year, ops, touched, painted, roll,
                                        open_src):
                    if any(rec.tv.value(i, c) is not None for i in rec.tv.body
                           for c in range(rec.tv.ncols)):
                        roll.add("미처리", rec.where, f"{t.id} 자본변동표 블록 경계 인식 실패",
                                 note="당기 실적을 수기 작성")
                continue
            if not rec.cur:
                continue
            if _has_tu(rec):
                roll.add("미처리", rec.where, f"{t.id} TU 열 포함", note="자동 필드 표는 손대지 않음")
                continue
            paired = _column_pairs(rec.tv, rec.cur, rec.prev)
            n_cur = _clear_cells(rec, rec.cur, ops, touched)
            if paired is None:
                roll.add("미처리", rec.where, f"{t.id} 당기 {len(rec.cur)}열/전기 {len(rec.prev)}열",
                         note=f"열 대응 실패 — 당기 {n_cur}칸만 비움")
                continue
            pairs, _ = paired
            prev_cols = [pc for _, pc in pairs]
            if rec.kind == "BS":
                if same_year:
                    roll.add("자동확정", rec.where,
                             f"{t.id} 당기말 {n_cur}칸 비움, 전기말 유지(직전 연말)")
                    if d_annual is not None:
                        dart_compare(rec, prev_cols, d_annual, target, dreport)
                else:
                    _fill_balance_from_c(rec, prev_cols, c_recs, ops, touched, painted_cells,
                                         roll, n_cur, d_annual, target, dreport)
                continue
            # PL / CF / RE — 흐름
            b = match_statement(rec, b_recs) if b_recs else None
            if b is not None:
                b_pairs = _column_pairs(b.tv, b.cur, b.prev)
                if b_pairs is None or len(b_pairs[0]) != len(pairs):
                    b = None
                    roll.add("확인필요", rec.where,
                             f"{t.id} 전년 동기 표 열 구성 불일치 — DART/수기로 채움")
            if b is not None:
                b_cur = [cc for cc, _ in _column_pairs(b.tv, b.cur, b.prev)[0]]
                rowmap = _row_map(rec.tv, b.tv)
                n, missing = _copy_cols(rec, prev_cols, b, b_cur, rowmap, ops, touched,
                                        painted_cells)
                roll.add("자동확정", rec.where,
                         f"{t.id} 당기 {n_cur}칸 비움, 전기 {n}칸 ← 전년 동기 DSD 당기 열")
                for lab in _insert_b_only_rows(rec, b, prev_cols, b_cur, rowmap, ops):
                    roll.add("자동확정", rec.where,
                             f"{t.id} 전년 동기에만 있는 행 '{lab[:20]}' 넣음(전기 값, 당기 칸은 비움)")
                for i in missing:
                    roll.add("확인필요", rec.where,
                             f"{t.id} '{rec.tv.row_label(i)[:20]}' 전년 동기 표에 없는 행 — 전기 비움")
                if d_target is not None:
                    vals = _values_after(rec, prev_cols, b, b_cur, rowmap)
                    dart_compare(rec, prev_cols, d_target, target, dreport, vals)
            elif d_target is not None:
                n = dart_fill(rec, prev_cols, d_target, d_base, base.quarter, target, ops,
                              touched, painted_cells, dreport, roll)
                roll.add("자동확정", rec.where,
                         f"{t.id} 당기 {n_cur}칸 비움, 전기 {n}칸 ← DART "
                         f"{target.year - 1} {P.QUARTER_WORD[target.quarter]}")
            else:
                n = _clear_cells(rec, prev_cols, ops, touched)
                roll.add("확인필요", rec.where,
                         f"{t.id} 당기 {n_cur}칸·전기 {n}칸 비움 — 전년 동기 DSD·DART 모두 없음",
                         note="전기 값 수기 입력")

        # ── ② 주석 표 ──
        for rec in recs:
            if rec.kind != "NOTE":
                continue
            t = rec.table
            if _has_tu(rec):
                continue
            # 행 블록형(<전기>/<당기> 기초..기말)
            if not rec.cur and not rec.prev and not rec.caption and _blocks(rec.tv):
                b = note_map_b.get(id(rec))
                if b is None and b_recs:
                    # 블록형은 헤더에 기간이 없어 match_notes 대상이 아니었을 수 있음 → 제목만
                    b = _block_partner(rec, b_recs)
                open_src = None
                if not same_year:
                    open_src = _note_open_from(rec, note_map_c.get(id(rec)))
                roll_block_table(rec, b, same_year, ops, touched, painted, roll, open_src)
                continue
            if rec.caption:
                _roll_caption_table(rec, recs, note_map_b, note_map_c, same_year, ops,
                                    touched, painted_cells, roll,
                                    b_present=bool(b_recs), c_present=bool(c_recs))
                continue
            if not rec.cur:
                continue
            paired = _column_pairs(rec.tv, rec.cur, rec.prev)
            n_cur = _clear_cells(rec, rec.cur, ops, touched)
            if not rec.prev:
                roll.add("자동확정", rec.where, f"{t.id} 당기 전용 표 {n_cur}칸 비움")
                continue
            if paired is None:
                roll.add("미처리", rec.where, f"{t.id} 당기 {len(rec.cur)}열/전기 {len(rec.prev)}열",
                         note=f"열 대응 실패 — 당기 {n_cur}칸만 비움, 전기 열 확인")
                continue
            pairs, _ = paired
            prev_cols = [pc for _, pc in pairs]
            if rec.is_balance:
                if same_year:
                    roll.add("자동확정", rec.where, f"{t.id} 당기 {n_cur}칸 비움, 전기말 유지")
                else:
                    c = note_map_c.get(id(rec))
                    _fill_note_from(rec, prev_cols, c, "전기 감사보고서 DSD", ops, touched,
                                    painted_cells, roll, n_cur, use_cur_of_src=True,
                                    src_present=bool(c_recs))
                continue
            b = note_map_b.get(id(rec))
            _fill_note_from(rec, prev_cols, b, "전년 동기 DSD", ops, touched, painted_cells,
                            roll, n_cur, use_cur_of_src=True, src_present=bool(b_recs))

        # ── ③ 어휘·날짜 치환 (셀·문단·제목, TU 제외) ──
        rules = P.vocab_rules(base.quarter, target.quarter)
        dmap = P.date_map(base, target)
        gisu = _gisu_mapping(doc) if not same_year else {}
        review_secs = {id(s) for s in secs if _skip_sec(s) and "검토보고서" in norm(s.title)}
        review_nodes: set[str] = set()
        for s in secs:
            if id(s) in review_secs:
                for blk in s.blocks:
                    if blk.kind == "p" or blk.kind == "label":
                        review_nodes.add(blk.item.id)
                    elif blk.kind == "table":
                        for r in blk.item.rows:
                            for c in r.cells:
                                review_nodes.add(c.id)
        n_vocab = n_date = 0
        for node in doc.nodes.values():
            if node.tag == "TU" or node.id in touched:
                continue
            text = node.text
            if node.id in review_nodes:
                # 검토보고서 문단은 종류별로 문장이 달라 자동 치환하지 않는다
                if any(p.search(text) for p, _ in rules) or any(
                        f"{y}년" in text for (y, _, _) in dmap):
                    roll.add("확인필요", node.id, "검토보고서 문단 — 기간 문구·날짜 직접 수정",
                             before=text.strip()[:36].replace("\n", " "))
                continue
            new = P.apply_vocab(text, rules)
            new = P.apply_dates(new, dmap)
            if gisu:
                new = _bump_gisu(new, gisu)
            if new == text:
                continue
            ops.append(splice.replace_inner(node, new))
            touched.add(node.id)
            if node.tag == "P":
                roll.add("자동확정", node.id, "문단 기간 문구 치환",
                         before=text.strip()[:30].replace("\n", " "),
                         after=new.strip()[:30].replace("\n", " "))
            n_vocab += 1
        if n_vocab:
            roll.add("자동확정", "문서 전체", f"기간 어휘·날짜 치환 {n_vocab}곳",
                     note=f"{P.QUARTER_WORD[base.quarter]}→{P.QUARTER_WORD[target.quarter]}")

        # ── ④ TU 자동필드 (AUNITVALUE + 표시 텍스트) ──
        tmap = P.tu_value_map(dmap)
        n_tu = 0
        for node in doc.nodes.values():
            if node.tag != "TU":
                continue
            frag = doc.data[node.start:node.end]
            new = frag
            for k, v in tmap.items():
                new = new.replace(k, v)
            inner = doc.data[node.inner_start:node.inner_end].decode("utf-8", "replace")
            new_inner = P.apply_dates(inner, dmap)
            if gisu:
                new_inner = _bump_gisu(new_inner, gisu)
            if new_inner != inner:
                head = new[: node.inner_start - node.start]
                tail = new[node.inner_end - node.start:]
                new = head + new_inner.encode("utf-8") + tail
            if new != frag:
                ops.append(splice.Op(pos=node.start, start=node.start, end=node.end,
                                     payload=new))
                n_tu += 1
        if n_tu:
            roll.add("자동확정", "자동 필드", f"TU 기간 {n_tu}곳 → {target.label}")

        roll.add("확인필요", "문서정보", "보고서 종류·기간 (DOCUMENT-NAME·meta)",
                 note="DART 편집기 [문서정보]에서 문서 종류(분기/반기)와 기간을 확인 — 도구는 미변경")

    if painted_cells:
        painted.append(("cells", painted_cells))
    roll.painted = _final_painted(ops, painted)
    return ops, roll, dreport


# ── 보조: 값 원천 채움 ───────────────────────────────────────────────

def _gisu_mapping(doc: Document) -> dict[int, int]:
    """문서에 등장하는 기수 중 상위 3개만 +1 — '제 1 기'류 서술(설립 연혁) 오치환 방지."""
    nums = set()
    for node in doc.nodes.values():
        for m in _GISU_RE.finditer(node.text):
            nums.add(int(m.group(2)))
    return {n: n + 1 for n in sorted(nums, reverse=True)[:3]}


def _bump_gisu(text: str, mapping: dict[int, int]) -> str:
    """기수 라벨만 +1 (연도는 date_map 이 따로 처리하므로 rollforward._roll_text 와 달리
    연도 치환은 하지 않는다 — 이중 치환 방지)."""
    return _GISU_RE.sub(
        lambda m: m.group(1) + str(mapping.get(int(m.group(2)), int(m.group(2)))), text)


def _values_after(a: TableRec, a_cols, b: TableRec, b_cols, rowmap) -> dict:
    """복사 후 A 전기 열 값(원 환산) — DART 대사용."""
    unit = a.tv.unit
    out = {}
    for i in a.tv.body:
        j = rowmap.get(i)
        for ac, bc in zip(a_cols, b_cols):
            v = None
            if j is not None:
                src = b.tv.cell(j, bc)
                v = cell_value(src.text) if src is not None else None
            out[(i, ac)] = _fval(v, unit) if v is not None else None
    return out


def _fill_note_from(rec: TableRec, prev_cols: list[int], src: TableRec | None,
                    src_name: str, ops: list, touched: set, painted_cells: list,
                    roll: RollReport, n_cur: int, use_cur_of_src: bool,
                    src_present: bool = True) -> None:
    t = rec.table
    if src is None:
        n = _clear_cells(rec, prev_cols, ops, touched)
        why = f"{src_name} 없음" if not src_present else f"{src_name}에서 짝 표를 못 찾음"
        roll.add("확인필요", rec.where,
                 f"{t.id} 당기 {n_cur}칸 비움, 전기 {n}칸 비움 — {why}",
                 note="전기 값 수기 입력" + ("" if src_present else f" 또는 {src_name}를 함께 넣어 재실행"))
        return
    s_paired = _column_pairs(src.tv, src.cur, src.prev) if src.prev else ([(c, c) for c in src.cur], [])
    if s_paired is None or len(s_paired[0]) != len(prev_cols):
        n = _clear_cells(rec, prev_cols, ops, touched)
        roll.add("확인필요", rec.where,
                 f"{t.id} 당기 {n_cur}칸 비움, 전기 {n}칸 비움 — {src_name} 표 열 구성 불일치")
        return
    src_cols = [cc for cc, _ in s_paired[0]] if use_cur_of_src else [pc for _, pc in s_paired[0]]
    rowmap = _row_map(rec.tv, src.tv)
    n, missing = _copy_cols(rec, prev_cols, src, src_cols, rowmap, ops, touched, painted_cells)
    roll.add("자동확정", rec.where,
             f"{t.id} 당기 {n_cur}칸 비움, 전기 {n}칸 ← {src_name} {src.table.id}")
    for i in missing:
        roll.add("확인필요", rec.where,
                 f"{t.id} '{rec.tv.row_label(i)[:20]}' {src_name} 표에 없는 행 — 전기 비움")


def _fill_balance_from_c(rec: TableRec, prev_cols, c_recs, ops, touched, painted_cells,
                         roll, n_cur, d_annual, target, dreport) -> None:
    """3Q→1Q: 재무상태표 전기말 ← 전기 감사보고서 당기말, 없으면 DART 사업보고서."""
    t = rec.table
    c = match_statement(rec, c_recs) if c_recs else None
    if c is not None:
        c_paired = _column_pairs(c.tv, c.cur, c.prev)
        if c_paired is not None and len(c_paired[0]) == len(prev_cols):
            c_cur = [cc for cc, _ in c_paired[0]]
            rowmap = _row_map(rec.tv, c.tv)
            n, missing = _copy_cols(rec, prev_cols, c, c_cur, rowmap, ops, touched,
                                    painted_cells)
            roll.add("자동확정", rec.where,
                     f"{t.id} 당기말 {n_cur}칸 비움, 전기말 {n}칸 ← 전기 감사보고서 당기말")
            for i in missing:
                roll.add("확인필요", rec.where,
                         f"{t.id} '{rec.tv.row_label(i)[:20]}' 전기 감사보고서에 없는 행 — 전기말 비움")
            if d_annual is not None:
                vals = _values_after(rec, prev_cols, c, c_cur, rowmap)
                dart_compare(rec, prev_cols, d_annual, target, dreport, vals)
            return
        roll.add("확인필요", rec.where, f"{t.id} 전기 감사보고서 재무상태표 열 구성 불일치")
    if d_annual is not None:
        n = dart_fill(rec, prev_cols, d_annual, None, 4, target, ops, touched, painted_cells,
                      dreport, roll)
        roll.add("자동확정", rec.where,
                 f"{t.id} 당기말 {n_cur}칸 비움, 전기말 {n}칸 ← DART {target.year - 1} 사업보고서")
        return
    n = _clear_cells(rec, prev_cols, ops, touched)
    roll.add("확인필요", rec.where,
             f"{t.id} 당기말 {n_cur}칸·전기말 {n}칸 비움 — 전기 감사보고서·DART 모두 없음",
             note="전기말 값 수기 입력")


def _sce_open_from(rec: TableRec, c_recs, d_annual) -> dict[int, str] | None:
    """3Q→1Q 자본변동표 당기초 행 값(열별 텍스트): 전기 감사보고서 SCE 의 당기말 행,
    없으면 DART 사업보고서 BS 자본 항목을 열 머리글로 매칭."""
    c = match_statement(rec, c_recs) if c_recs else None
    if c is not None and c.tv.ncols == rec.tv.ncols:
        cb = _cur_block(c.tv)
        if cb is not None:
            close_row = cb[1]
            return {col: c.tv.cell(close_row, col).text
                    for col, _ in _numeric_cells(c.tv, close_row)
                    if c.tv.cell(close_row, col) is not None}
    if d_annual is not None:
        items = d_annual.by_kind("BS")
        out: dict[int, str] = {}
        blk = _blocks(rec.tv)
        if blk is None:
            return None
        o2 = blk[1][0]
        for col, cell in _numeric_cells(rec.tv, o2):
            hdr = rec.tv.col_header(col)
            m = D.match_items([hdr], items)
            if 0 in m and m[0].thstrm is not None:
                out[col] = to_dsd_text(round(m[0].thstrm / rec.tv.unit), like=cell.text)
        return out or None
    return None


def _note_open_from(rec: TableRec, c: TableRec | None) -> dict[int, str] | None:
    if c is None or c.tv.ncols != rec.tv.ncols:
        return None
    cb = _cur_block(c.tv)
    if cb is None:
        return None
    close_row = cb[1]
    return {col: c.tv.cell(close_row, col).text
            for col, _ in _numeric_cells(c.tv, close_row)}


def _block_partner(rec: TableRec, b_recs: list[TableRec]) -> TableRec | None:
    key = D.norm_label(rec.sec.note_title)
    cands = [b for b in b_recs
             if b.sec.note_no is not None and D.norm_label(b.sec.note_title) == key
             and b.tv.ncols == rec.tv.ncols and _cur_block(b.tv) is not None]
    if not cands:
        return None
    return max(cands, key=lambda b: _label_overlap(_row_labels(rec.tv), _row_labels(b.tv)))


def _roll_caption_table(rec: TableRec, recs: list[TableRec], note_map_b, note_map_c,
                        same_year: bool, ops, touched, painted_cells, roll,
                        b_present: bool = False, c_present: bool = False) -> None:
    """세로형 표(바로 앞 문단 '(1) 당반기말' / '(2) 전기말' / '<당반기>' / '<전반기>')."""
    t = rec.table
    cap = norm(rec.caption)
    has_num = any(rec.tv.value(i, c) is not None for i in rec.tv.body for c in range(rec.tv.ncols))
    if not has_num:
        return  # 캡션 아래 제목·단위 표 — 값 없음
    if _CUR_MARK_RE.search(cap):
        n = _clear_all_numeric(rec, ops, touched)
        roll.add("자동확정", rec.where, f"{t.id} 당기 표('{rec.caption[:12]}') 값 {n}칸 비움")
        return
    if _BAL_MARK_RE.search(cap) or _ANNUAL_MARK_RE.fullmatch(cap):
        # '전기말' · '전기'(연간 - 요약재무정보 등)는 같은 해 전환에서 그대로, 해 넘김은 전기 감사보고서
        if same_year:
            roll.add("자동확정", rec.where, f"{t.id} 전기(말) 표('{rec.caption[:12]}') 유지")
            return
        src, src_name, present = note_map_c.get(id(rec)), "전기 감사보고서 DSD", bool(note_map_c) or c_present
    else:
        src, src_name, present = note_map_b.get(id(rec)), "전년 동기 DSD", b_present
    if src is None:
        n = _clear_all_numeric(rec, ops, touched)
        why = f"{src_name} 없음" if not present else f"{src_name}에서 원천 표를 못 찾음"
        roll.add("확인필요", rec.where,
                 f"{t.id} 전기 표('{rec.caption[:12]}') 값 {n}칸 비움 — {why}",
                 note="전기 값 수기 입력" + ("" if present else f" 또는 {src_name}를 함께 넣어 재실행"))
        return
    # 원천은 '당기' 캡션 표여야 한다 — 짝 표는 match_notes 가 헤더 서명으로 골랐으므로 그대로 복사
    if src.tv.ncols != rec.tv.ncols:
        n = _clear_all_numeric(rec, ops, touched)
        roll.add("확인필요", rec.where, f"{t.id} 전기 표 열 수 불일치 — 값 {n}칸 비움")
        return
    rowmap = _row_map(rec.tv, src.tv)
    cols = list(range(rec.tv.ncols))
    n, missing = _copy_cols(rec, cols, src, cols, rowmap, ops, touched, painted_cells)
    roll.add("자동확정", rec.where, f"{t.id} 전기 표 값 {n}칸 ← 원천 {src.table.id}")
    for i in missing:
        roll.add("확인필요", rec.where, f"{t.id} '{rec.tv.row_label(i)[:20]}' 원천 표에 없는 행 — 비움")


# ── 리포트 시트 ──────────────────────────────────────────────────────

DART_SHEET = "DART대사"


def write_dart_sheet(xlsx_path: str, rep: DartReport) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, PatternFill

    from .export_xlsx import FILL_HEAD, FONT, FONT_HEAD, FONT_TITLE

    wb = load_workbook(xlsx_path)
    if DART_SHEET in wb.sheetnames:
        del wb[DART_SHEET]
    ws = wb.create_sheet(DART_SHEET, 1)
    ws.column_dimensions["A"].hidden = True
    for col, w in (("B", 8), ("C", 22), ("D", 30), ("E", 10), ("F", 18), ("G", 18),
                   ("H", 16), ("I", 30), ("J", 8), ("K", 8)):
        ws.column_dimensions[col].width = w
    rid = iter(range(1, 100000))

    def put(r, c, v, font=FONT, fill=None, numfmt=None):
        ws.cell(row=r, column=1, value=f"!r{next(rid):05d}").font = FONT
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = font
        if fill:
            cell.fill = fill
        if numfmt:
            cell.number_format = numfmt
        cell.alignment = Alignment(vertical="center")
        return cell

    put(2, 2, "DART 대사 — 본 재무제표 전기 값 ↔ OpenDART 공시 값", FONT_TITLE)
    put(3, 2, rep.period_desc or "")
    if rep.skipped_reason:
        put(4, 2, f"DART 생략/제한: {rep.skipped_reason}")
    r = 6
    heads = ["표", "위치", "행 라벨", "열", "파일 값(원)", "DART 값(원)", "차이", "DART 계정명",
             "매칭", "판정"]
    for c, h in enumerate(heads, start=2):
        cell = put(r, c, h, FONT_HEAD)
        cell.fill = FILL_HEAD
    r += 1
    fill_bad = PatternFill("solid", fgColor="FDE9E9")
    fill_warn = PatternFill("solid", fgColor="FFF4DE")
    fill_fill = PatternFill("solid", fgColor="DDEBF7")
    order = {"불일치": 0, "미매칭": 1, "채움": 2, "일치": 3}
    for row in sorted(rep.rows, key=lambda x: order.get(x.verdict, 9)):
        f = {"불일치": fill_bad, "미매칭": fill_warn, "채움": fill_fill}.get(row.verdict)
        put(r, 2, row.kind, fill=f)
        put(r, 3, row.where)
        put(r, 4, row.label)
        put(r, 5, row.col)
        put(r, 6, row.file_value, numfmt="#,##0;[Red](#,##0);-")
        put(r, 7, row.dart_value, numfmt="#,##0;[Red](#,##0);-")
        diff = (abs(row.file_value) - abs(row.dart_value)
                if row.file_value is not None and row.dart_value is not None else None)
        put(r, 8, diff, numfmt="#,##0;[Red](#,##0);-")
        put(r, 9, row.dart_name)
        put(r, 10, row.how)
        put(r, 11, row.verdict, fill=f)
        r += 1
    if rep.unmatched_dart:
        r += 1
        put(r, 2, "파일에서 짝을 못 찾은 DART 항목 (재분류·표시 차이 확인용)", FONT_HEAD)
        r += 1
        for kind, name, val in rep.unmatched_dart:
            put(r, 2, kind)
            put(r, 4, name)
            put(r, 7, val, numfmt="#,##0;[Red](#,##0);-")
            r += 1
    wb.save(xlsx_path)
