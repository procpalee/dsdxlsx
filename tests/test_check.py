"""check(정합성 검증) 테스트 — 원본은 깨끗해야 하고, 주입한 오류는 잡아야 한다."""
import pytest

from dsdxlsx import splice
from dsdxlsx.check import cell_value, run_checks
from dsdxlsx.dsd_io import read_dsd
from dsdxlsx.parse import parse_contents


@pytest.fixture(scope="module")
def doc(sample_dsd):
    return parse_contents(read_dsd(sample_dsd).contents)


def test_cell_value():
    from decimal import Decimal

    assert cell_value("12,345,678,900") == 12345678900
    assert cell_value("(232,116,145)") == -232116145
    assert cell_value("-") == 0
    assert cell_value("6,28,31") is None
    assert cell_value("9.56%") == Decimal("9.56")  # 지분율 합계 검증 대상
    assert cell_value("2,978,480주") == 2978480


def test_pristine_sample_sums_are_clean(doc):
    """심리 완료 원본은 모든 합계 검산(C1·C2·C3·C7)을 전수 통과해야 한다."""
    report = run_checks(doc, only={"C1", "C2", "C3", "C7"})
    issues = [f for f in report.findings]
    assert issues == [], "\n".join(f"{f.check} {f.where}: {f.message}" for f in issues)
    assert report.counts["C1"] == 194  # AUTOSUM 전수
    assert report.counts["C3"] >= 170  # 합계행 + 합계 헤더 열(가로합)
    assert report.counts["C7"] >= 20  # 손익 계산 라인 + 자본변동(기초+변동=기말) + CF 합산식


def _find_autosum_component(doc):
    """AUTOSUM(세로 합계) 셀의 구성 항목 하나(수치 본문 셀)를 동적으로 찾는다.
    특정 회사 수치에 의존하지 않도록 구조로 탐색."""
    from dsdxlsx.check import build_views

    views, _ = build_views(doc)
    for tv in views:
        for i in tv.body:
            for c in range(tv.ncols):
                cell = tv.cell(i, c)
                if cell is None or not cell.attrs.get("AUTOSUM", "").startswith("COL"):
                    continue
                # 이 합계 위쪽에서 수치 본문 셀(합계·소계 아님) 하나
                for j in reversed([b for b in tv.body if b < i]):
                    comp = tv.cell(j, c)
                    if (comp is not None and tv.value(j, c) not in (None, 0)
                            and not tv.row_is_total(j)):
                        return comp, tv.value(j, c)
    return None, None


def test_injected_sum_error_is_detected(sample_dsd, doc):
    """합계 구성 항목을 변조하면 C1이 합계 불일치를 잡아야 한다."""
    target, value = _find_autosum_component(doc)
    if target is None:
        pytest.skip("이 샘플에 AUTOSUM 합계가 없음")
    from decimal import Decimal

    bumped = int(Decimal(str(value))) + 1_000_000
    new_contents = splice.apply(doc.data, [splice.replace_inner(target, f"{bumped:,}")])
    report = run_checks(parse_contents(new_contents), only={"C1"})
    assert any(not x.ok for x in report.results), "변조된 합계 불일치를 감지하지 못함"


def test_fs_note_crosscheck_runs(doc):
    report = run_checks(doc, only={"C4"})
    assert report.counts.get("C4", 0) > 100  # 참조 없는 행도 전역 검색으로 대사됨
    # 참조번호 오류 의심(warn)과 미확인(info)만 존재
    assert all(f.level in ("warn", "info") for f in report.findings)
    # 참조 주석에 없고 다른 주석에서 발견되는 케이스가 감지됨 (참조번호 오류 후보)
    assert any("참조번호" in f.message for f in report.findings if f.level == "warn")


def test_format_check_finds_known_issues(doc):
    report = run_checks(doc, only={"C6"})
    assert any("여분 공백" in f.message for f in report.findings)  # 실측 12건 존재
    assert not any(f.level == "warn" for f in report.findings)  # 깨진 숫자는 없어야 함


def test_annotate_in_place(sample_dsd, tmp_path):
    """검증 셀이 편집용 xlsx 안에 엑셀 수식으로 기록되고(살아있는 검증),
    주석이 있어도 import 결과는 오염되지 않아야 한다."""
    import re as re_

    from openpyxl import load_workbook

    from dsdxlsx.cli import main
    from dsdxlsx.dsd_io import read_dsd
    from dsdxlsx.import_xlsx import import_xlsx

    doc0 = parse_contents(read_dsd(sample_dsd).contents)
    edit_xlsx = str(tmp_path / "edit.xlsx")
    assert main(["export", sample_dsd, "-o", edit_xlsx]) == 0

    # 1) 원본 상태 check → 검증 열이 IF 수식으로 기록 + 검증요약 시트 생성
    assert main(["check", edit_xlsx, "--only", "C1,C2,C3,C7"]) == 0
    wb = load_workbook(edit_xlsx)
    assert wb.sheetnames[0] == "검증요약"
    formulas = []
    for ws in wb.worksheets:
        if ws.title == "검증요약":
            continue
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("=IF("):
                    formulas.append(cell.value)
    assert len(formulas) > 150  # C1/C2/C3/C7 검증이 수식으로 존재
    # 수식 형태: 단일 =IF(대상=(…),TRUE,…) 또는 결합 =IF(AND(…),TRUE,…)
    assert all(
        re_.match(r"^=IF\((AND\()?[A-Z]+\d+=\([A-Z\d+\-]+\)", f) for f in formulas
    )
    # C7 부호 검산: 손익 계산 라인 수식에 뺄셈 항 존재 (매출원가/판관비 등)
    assert any("-" in f for f in formulas)
    wb.close()

    # 2) 수식이 붙은 파일도 무수정 import가 그대로 성립 (오염 없음)
    _, _, imp = import_xlsx(edit_xlsx)
    assert not imp.changed

    # 3) 값 수정 후 check 재실행 → 요약에 불일치 목록 생성 (수식은 엑셀에서도 재계산됨)
    comp, value = _find_autosum_component(doc0)
    assert comp is not None
    numeric = int(value)
    wb = load_workbook(edit_xlsx)
    hit = False
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value == numeric and not hit:
                    cell.value = numeric + 1_000_000
                    hit = True
    assert hit
    wb.save(edit_xlsx)
    assert main(["check", edit_xlsx, "--only", "C1"]) == 0
    wb = load_workbook(edit_xlsx)
    summary = wb["검증요약"]
    labels = [c.value for row in summary.iter_rows() for c in row if isinstance(c.value, str)]
    assert any("불일치 목록" in s for s in labels)
    wb.close()


def test_separate_report_option(doc, tmp_path):
    """-o 옵션: 별도 검증 리포트 xlsx (불일치 차이 금액, 상단 정렬)."""
    from openpyxl import load_workbook

    from dsdxlsx.report_xlsx import write_report

    target, value = _find_autosum_component(doc)
    if target is None:
        pytest.skip("이 샘플에 AUTOSUM 합계가 없음")
    new_contents = splice.apply(
        doc.data, [splice.replace_inner(target, f"{int(value) + 1_000_000:,}")])
    report = run_checks(parse_contents(new_contents), only={"C1"})
    bad_xlsx = str(tmp_path / "bad_검증.xlsx")
    stats = write_report(report, bad_xlsx)
    assert stats["bad"] >= 1
    wb = load_workbook(bad_xlsx)
    first = list(wb["검증결과"].iter_rows(min_row=2, max_row=1 + stats["bad"], values_only=True))
    assert all(isinstance(r[7], (int, float)) and r[7] != 0 for r in first)  # 차이 금액
    wb.close()
