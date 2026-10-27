"""A1 표 편집 — 마커 문법, clone-only 합성, import 통합, 가드, ID 재대응."""
import pytest
from openpyxl import load_workbook

from dsdxlsx import compose, markers
from dsdxlsx import grid as grid_mod
from dsdxlsx.import_xlsx import ImportError_, import_xlsx
from dsdxlsx.parse import parse_contents

from fixtures import factory


# ── 마커 문법 ─────────────────────────────────────────────────────────

def test_marker_syntax():
    m = markers.parse_marker("#새표 헤더=2 [원본=t014] 단위=원")
    assert m.kind == "새표" and m.args == {"헤더": "2", "원본": "t014", "단위": "원"}
    m = markers.parse_marker("#문단 새 설명 문단입니다.")
    assert m.kind == "문단" and m.text == "새 설명 문단입니다."
    assert markers.parse_marker("#표삭제").kind == "표삭제"
    assert not markers.is_marker("##텍스트")  # 이스케이프
    assert markers.unescape("##텍스트") == "#텍스트"
    with pytest.raises(markers.MarkerError):
        markers.parse_marker("#새표오타")


# ── parse: TABLE-GROUP / INSERTION 추적 (가드 근거) ──────────────────

def test_parse_tracks_group_and_insertion():
    t = factory.table([["구분", "당기"]], [["항목", "1,000"]])
    contents = factory.document(
        factory.p("표지"),
        factory.group(t, aclass="PBS-2:A"),
        factory.insertion(factory.table([["a", "b"]], [["c", "1"]])),
        t,
    )
    doc = parse_contents(contents)
    tables = list(doc.tables.values())
    assert tables[0].group_aclass == "PBS-2:A" and not tables[0].in_insertion
    assert tables[1].in_insertion and tables[1].group_aclass is None
    assert tables[2].group_aclass is None and not tables[2].in_insertion

    from dsdxlsx.blockops import table_delete_guard
    assert "TABLE-GROUP" in table_delete_guard(tables[0])
    assert "INSERTION" in table_delete_guard(tables[1])
    assert table_delete_guard(tables[2]) is None


# ── E2E: export → 마커 편집 → import ────────────────────────────────

def _exported(tmp_path):
    from dsdxlsx.export_xlsx import export_xlsx

    contents = factory.standard_contents()
    dsd_path = tmp_path / "s.dsd"
    dsd_path.write_bytes(factory.dsd_bytes(contents))
    xlsx = str(tmp_path / "s.xlsx")
    export_xlsx(parse_contents(contents), dsd_path.read_bytes(), str(dsd_path), xlsx)
    return xlsx


def _append_rows(xlsx, sheet, rows):
    """rows = [dict(col번호→값)] — 시트 끝에 이어 붙임."""
    wb = load_workbook(xlsx)
    ws = wb[sheet]
    r = ws.max_row
    for row in rows:
        r += 1
        for col, val in row.items():
            ws.cell(row=r, column=col, value=val)
    wb.save(xlsx)


def test_new_table_and_paragraph(tmp_path):
    xlsx = _exported(tmp_path)
    _append_rows(xlsx, "5", [
        {2: "#새표 헤더=1 단위=원"},
        {2: "구  분", 3: "당기말", 4: "전기말"},
        {2: "신규항목", 3: 1234, 4: 999},
        {2: "합  계", 3: 1234, 4: 999},
        {2: "#표끝"},
        {2: "#문단 새로 추가된 설명 문단입니다."},
    ])
    _, new_contents, rep = import_xlsx(xlsx)
    assert len(rep.table_inserts) == 1 and len(rep.para_inserts) == 1
    assert rep.structural and rep.changed

    doc2 = parse_contents(new_contents)
    assert len(doc2.tables) == 6  # 5(제목표·BS·주석4·5·7) + 신설 1
    # 신설 표: 주석5 표(t004) 뒤 → t005
    new_t = doc2.tables["t005"]
    texts = [c.text for r in new_t.rows for c in r.cells]
    assert "(단위 : 원)" in texts and "구  분" in texts
    assert "신규항목" in texts and "1,234" in texts  # 숫자 표기 변환
    # 삽입 관례: <TABLE 앞 빈 줄
    assert new_contents[new_t.start - 4:new_t.start] == b"\r\n\r\n"
    # 새 문단
    assert any("새로 추가된 설명 문단" in n.text for n in doc2.nodes.values())

    # row_positions 재대응: 모든 tid가 새 문서의 행 수와 정합
    for tid, positions in rep.row_positions.items():
        assert len(positions) == len(doc2.tables[tid].rows), tid
    # 신설 표 위치가 실제 시트 행을 가리킴
    assert all(s == "5" for s, _ in rep.row_positions["t005"])

    # 결정론: 같은 xlsx 재import → 동일 바이트
    _, again, _ = import_xlsx(xlsx)
    assert again == new_contents


def test_table_delete(tmp_path):
    xlsx = _exported(tmp_path)
    wb = load_workbook(xlsx)
    ws = wb["7"]  # 주석7 담보제공 표
    target = next(r for r in ws.iter_rows()
                  if r[0].value and str(r[0].value).startswith("t"))
    ws.cell(row=target[0].row, column=2, value="#표삭제")
    wb.save(xlsx)

    _, new_contents, rep = import_xlsx(xlsx)
    assert rep.table_deletes and rep.structural
    doc2 = parse_contents(new_contents)
    assert len(doc2.tables) == 4  # 5 - 1
    labels = [c.text for t in doc2.tables.values() for r in t.rows for c in r.cells]
    # 주석7 표(담보 매출채권 5,500)는 사라지고 주석5 표(5,500)는 남는다
    assert sum(1 for x in labels if x == "5,500") == 1


def test_paragraph_delete(tmp_path):
    xlsx = _exported(tmp_path)
    wb = load_workbook(xlsx)
    ws = wb["1"]  # 주석1 — 문단 하나뿐
    prow = next(r for r in ws.iter_rows()
                if r[0].value and str(r[0].value).startswith("p"))
    ws.cell(row=prow[0].row, column=2, value="#문단삭제")
    wb.save(xlsx)
    _, new_contents, rep = import_xlsx(xlsx)
    assert rep.para_deletes
    doc2 = parse_contents(new_contents)
    assert not any("일반사항" in n.text for n in doc2.nodes.values())


# ── 가드 ─────────────────────────────────────────────────────────────

def test_guard_new_table_on_statement_sheet(tmp_path):
    xlsx = _exported(tmp_path)
    _append_rows(xlsx, "재무상태표", [
        {2: "#새표"}, {2: "구분", 3: "당기"}, {2: "x", 3: 1},
    ])
    with pytest.raises(ImportError_, match="주석 시트"):
        import_xlsx(xlsx)


def test_guard_unknown_marker_and_escape(tmp_path):
    xlsx = _exported(tmp_path)
    _append_rows(xlsx, "5", [{2: "#없는마커"}])
    with pytest.raises(ImportError_, match="알 수 없는 마커"):
        import_xlsx(xlsx)


def test_guard_too_many_columns(tmp_path):
    xlsx = _exported(tmp_path)
    _append_rows(xlsx, "5", [
        {2: "#새표"},
        {2: "구분", 3: "당기", 4: "전기", 5: "추가열"},  # 템플릿 3열 초과
        {2: "x", 3: 1, 4: 2, 5: 3},
    ])
    with pytest.raises(ImportError_, match="열"):
        import_xlsx(xlsx)


def test_vanished_table_error_mentions_marker(tmp_path):
    """표를 시트에서 그냥 지우면 여전히 에러 — #표삭제 안내 포함."""
    xlsx = _exported(tmp_path)
    wb = load_workbook(xlsx)
    ws = wb["7"]
    for row in list(ws.iter_rows()):
        if row[0].value and str(row[0].value).startswith("t"):
            ws.delete_rows(row[0].row, 1)
            break
    # 남은 행도 전부 제거해 '표 전체 소실' 상황을 만든다
    wb2rows = [r[0].row for r in ws.iter_rows()
               if r[0].value and str(r[0].value).startswith("t")]
    for rn in reversed(wb2rows):
        ws.delete_rows(rn, 1)
    wb.save(xlsx)
    with pytest.raises(ImportError_, match="#표삭제"):
        import_xlsx(xlsx)


def test_compose_harvest_rejects_grouped_template():
    t = factory.table([["구분", "당기"]], [["항목", "1,000"]])
    contents = factory.document(factory.p("표지"), factory.group(t))
    doc = parse_contents(contents)
    table = next(iter(doc.tables.values()))
    with pytest.raises(compose.ComposeError, match="TABLE-GROUP"):
        compose.harvest(doc, table)


# ── 행 추가 본보기: 가로 병합 줄 · 세로 병합 묶음 ───────────────────────

def _span_doc(tmp_path):
    """본문이 모두 세로 병합 묶음 + 합계 줄(구분 · 이름 가로 병합)뿐인 표 - 보통 행이 없다."""
    from dsdxlsx.export_xlsx import export_xlsx

    span = factory.table(
        [["구  분", "특수관계자명", "매출 등"]],
        [
            ["||ROWSPAN=2||종속기업", "A사", "100"],
            ["B사", "200"],
            ["||ROWSPAN=2||기타", "C사", "300"],
            ["D사", "400"],
            ["||COLSPAN=2||합  계", "1,000"],
        ],
        ncols=3,
    )
    base = factory.standard_contents()
    tail = "\r\n" + factory.p("8. 특수관계자\n거래 내역은 다음과 같습니다.") + "\r\n\r\n" + span + "\r\n"
    contents = base.replace(b"\r\n</DOCUMENT>", tail.encode("utf-8") + b"</DOCUMENT>")
    dsd_path = tmp_path / "span.dsd"
    dsd_path.write_bytes(factory.dsd_bytes(contents))
    xlsx = str(tmp_path / "span.xlsx")
    export_xlsx(parse_contents(contents), dsd_path.read_bytes(), str(dsd_path), xlsx)
    return contents, xlsx


def test_row_insert_never_clones_colspan_total_row(tmp_path):
    """합계 줄(COLSPAN 2)을 본보기로 쓰면 새 행의 이름 칸이 사라진다 - 묶음 안 행을 본보기로
    전체 폭 행을 만들고, 세로 병합 자리는 ROWSPAN 없는 칸으로 채운다."""
    contents, xlsx = _span_doc(tmp_path)
    wb = load_workbook(xlsx)
    ws = wb["8"]
    rows = {ws.cell(r, 3).value: r for r in range(1, ws.max_row + 1)}
    ws.insert_rows(rows["D사"] + 1)
    ws.cell(rows["D사"] + 1, 3, "F사"); ws.cell(rows["D사"] + 1, 4, 60)
    ws.insert_rows(rows["B사"] + 1)
    ws.cell(rows["B사"] + 1, 2, "관계기업"); ws.cell(rows["B사"] + 1, 3, "E사"); ws.cell(rows["B사"] + 1, 4, 50)
    wb.save(xlsx)

    _, new_contents, rep = import_xlsx(xlsx)
    assert len(rep.row_inserts) == 2
    doc = parse_contents(contents)
    doc2 = parse_contents(new_contents)
    tid = next(t for t, tb in doc2.tables.items() if any(c.text == "E사" for r in tb.rows for c in r.cells))
    t2 = doc2.tables[tid]
    assert len(t2.rows) == len(doc.tables[tid].rows) + 2
    texts = {tuple(c.text for c in r.cells) for r in t2.rows}
    assert ("관계기업", "E사", "50") in texts
    assert ("", "F사", "60") in texts
    new_rows = [r for r in t2.rows if any(c.text in ("E사", "F사") for c in r.cells)]
    assert all(len(r.cells) == 3 and all(c.rowspan == 1 and c.colspan == 1 for c in r.cells) for r in new_rows)
    g = grid_mod.expand(t2)
    assert not t2.irregular and all(gc is not None for row in g for gc in row)
