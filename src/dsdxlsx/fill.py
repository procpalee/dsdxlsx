"""DSD 당기 칸 채우기 — fill (recon 의 짝짓기를 거꾸로).

roll-interim 으로 당전기를 민 DSD 는 당기 칸이 비어 있고 전기 칸만 차 있다. 이 모듈은 비어 있는
당기 칸마다 조서(8110 최종분석적검토 · PA 정산표 · 연결정산표)에서 값을 찾아 DSD 원문 바이트에 쓴다.

채울 칸: 당기 열 칸 중 짝 전기 칸(같은 행 · 짝 열)에 값이 있는 칸 - 밀기 뒤 남은 이전 분기 값도 덮어씀.
  전기 칸이 모두 빈 행(새 계정)은 바로 위 행이 쓴 열(세부 · 소계)에. 자본변동표 · 변동표 주석처럼 행이
  기간을 정하는 표는 당기 행의 값 열. 짝 전기 칸이 비어 있고 조서에 전기 값이 있으면 전기 칸도.
  조서에 합계 행이 없으면 위 칸 합.
조서 값 찾기 (recon 과 같은 규칙):
  같은 이름 행 × 당기 열(머리글 날짜 · 기간 말 · 열 이름) → 후보가 여럿이면 그 행의 전기 값이
  DSD 전기 값과 같은 행(앵커). 앵커가 단위(원/천원 · %) · 부호(유출 섹션 양수 관례)도 정한다.
  앵커가 없으면 같은 표(없으면 같은 주석) 값 자릿수로 단위를 고르고, 손익계산서 결과 행 '…손실'은 양수 관례로.
못 채운 칸은 이유와 함께 보고(값이 틀린 채 채우는 것보다 빈 칸이 안전).
"""
from __future__ import annotations

import math
from dataclasses import dataclass, field

from . import period as P
from . import splice
from .check import is_total_label, norm
from .model import Document
from .numbers import to_dsd_text
from .recon import (_3M_RE, _RESULT_RE, _UNIT_MULT, ReconConfig, WpRow, WpSheet, _cell_periods, _row_cols,
                    _target_date, col_keys_of, common_keys, rec_scope, row_keys,
                    same_name_rows, table_strict)
from .rollforward import _column_pairs
from .rollinterim import _leaf, scan

_RATE_HDR_RE = __import__("re").compile(r"%|율")
_UNIT_TXT_RE = __import__("re").compile(r"단위\s*[:：]\s*([^)]+)\)")


@dataclass
class FillItem:
    sheet: str           # 구획 (재무상태표 / 주석4 …)
    table: str
    row_id: str
    label: str
    col: str
    value: float | None  # 쓴 값(DSD 단위)
    verdict: str         # 채움 | 못 채움
    how: str = ""        # 앵커 · 하나뿐 · 단위 · 부호 / 못 채운 이유
    where: str = ""      # 조서 위치


@dataclass
class FillReport:
    period: str = ""
    items: list[FillItem] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def count(self, verdict: str) -> int:
        return sum(1 for x in self.items if x.verdict == verdict)


def _ref(row: WpRow, col: int) -> str:
    from openpyxl.utils import get_column_letter
    return f"{row.sheet}!{get_column_letter(col + 1)}{row.row}"


def _scales(unit: int, wp_unit: int, rate: bool) -> list[float]:
    """조서 값 → DSD 단위 배수 후보."""
    out = [wp_unit / unit, 1.0, 1.0 / unit if unit > 1 else 1.0, 1 / 1000, 1 / 1_000_000]   # DSD 표에 단위 표시가 없어도
    if rate:
        out = [100.0, 1.0]
    seen, uniq = set(), []
    for s in out:
        if s not in seen:
            seen.add(s)
            uniq.append(s)
    return uniq


def _fit(pv: float, w: float, scales: list[float]) -> tuple[float, int] | None:
    """앵커: DSD 전기 값 pv 와 조서 전기 값 w 가 맞는 (배수, 부호)."""
    for s in scales:
        x = w * s
        for sg in (1, -1):
            y = sg * x
            if y == pv or (pv and (abs(y - pv) <= 1 if s != 1.0 or abs(pv) > 100 else abs(y - pv) <= 0.006)):
                return s, sg
            if pv == int(pv) and round(y) == pv and pv != 0:
                return s, sg
    return None


def _round(v: float, rate: bool, like: float | None) -> float:
    if rate or (like is not None and like != int(like)):
        return round(v, 2)
    return float(round(v))


def fill_ops(doc: Document, sheets: list[WpSheet], config: ReconConfig | None = None,
             keep_filled: bool = False) -> tuple[list[splice.Op], FillReport]:
    """당기 칸을 조서 값으로(기본: 값이 있어도 덮어씀 - 밀기 뒤 당기 칸에는 이전 분기 값이 남아 있을 수 있음).
    짝 전기 칸이 비어 있고(밀기가 이름이 달라 못 채움) 조서에 전기 값이 있으면 전기 칸도 채운다."""
    per = P.detect_period(doc)
    cfg = config or ReconConfig()
    rep = FillReport(period=per.label if per else "(기간 미상)")
    if per is None:
        rep.warnings.append("DSD 보고기간을 못 읽음 - 조서 머리글 날짜를 쓰지 않음")
    one_q = bool(per and per.quarter == 1)
    fs_rows = [r for sh in sheets if sh.kind == "fs" for r in sh.rows]
    ops: list[splice.Op] = []
    recs, _ = scan(doc)
    sec_mags: dict[int, list[float]] = {}         # 구획 → 값 크기(표에 값이 없을 때 자릿수 기준)
    for rec in recs:
        sec_mags.setdefault(id(rec.sec), []).extend(
            abs(float(rec.tv.value(i, c))) for i in rec.tv.body for c in range(rec.tv.ncols)
            if rec.tv.value(i, c) not in (None, 0))
    for k, rec in enumerate(recs):
        if rec.kind in ("OTHER", "RE"):
            continue
        tv = rec.tv
        scope, ctx = rec_scope(rec, sheets, fs_rows)
        unit = _eff_unit(recs, k)
        common = common_keys(scope)
        strict = table_strict(rec)
        pairs: dict[int, int] | None = None
        if rec.cur and rec.prev:
            got = _column_pairs(tv, list(rec.cur), list(rec.prev))
            pairs = dict(got[0]) if got else None
        periods = _cell_periods(rec, per.year if per else None)
        value_cols = sorted({c for i in tv.body for c in range(1, tv.ncols) if tv.value(i, c) is not None})
        if not value_cols and not pairs:            # 밀기가 비운 당기 표 - 같은 머리글 짝 표의 값 열
            value_cols = _partner_cols(rec, recs)
            if not value_cols and rec.caption_is_cur:   # 짝 표도 비면 0열 밖 모두(조서에 값이 있는 칸만 채워짐)
                value_cols = list(range(1, tv.ncols))
        mags = [abs(float(tv.value(i, c))) for i in tv.body for c in range(tv.ncols)
                if tv.value(i, c) not in (None, 0)] or sec_mags.get(id(rec.sec), [])
        mag = sorted(mags)[len(mags) // 2] if mags else None
        tgt = _target_date(per, "당기", rec)
        ptgt = _target_date(per, "전기", rec)
        last_row: dict[str, int] = {}            # 시트 → 바로 앞에서 채운 조서 행 (같은 이름 행이 둘일 때)
        last_pair: dict[str, int] = {}           # 머리글 열 묶음 → 바로 위 행이 쓴 당기 열
        new_vals: dict[tuple[int, int], float] = {}   # 이번에 쓴 값(합계 행 계산)
        for i in tv.body:
            label, keys, own = row_keys(tv, i)
            keys |= cfg.alias_keys(keys)
            if pairs:
                row_prev = {c: tv.value(i, pc) for c, pc in pairs.items()}
                no_prev = all(v is None for v in row_prev.values())
                had_cur = any(tv.value(i, c) is not None for c in pairs)
                slots = []
                for c in sorted(pairs):
                    pv = row_prev[c]
                    if pv is None and not no_prev:
                        continue                   # 같은 행 다른 열(소계 열 등)에 전기 값 - 이 열 자리가 아님
                    if no_prev and had_cur and tv.value(i, c) is None:
                        continue                   # 전기는 없고 당기 값이 다른 열에 있음
                    if no_prev and not had_cur and last_pair.get(
                            _grp(tv, c), min(x for x in pairs if _grp(tv, x) == _grp(tv, c))) != c:
                        continue                   # 새 행: 바로 위 행이 쓴 열만(세부 · 소계 열 중 하나)
                    slots.append((c, pairs[c], pv))
            else:                                  # 행이 기간을 정하는 표(자본변동표 · 변동표 주석)
                no_prev = False
                if rec.prev and not rec.cur:
                    slots = []                     # 전기만 있는 표(전기 변동표) - 밀기 몫
                elif rec.cur and not rec.prev:     # '당기초 | 지분법손익 | 당분기말' - 비어 있는 당기 열도
                    slots = [(c, None, None) for c in sorted(set(rec.cur) | set(value_cols)) if c > 0]
                else:
                    slots = [(c, None, None) for c in value_cols if periods.get((i, c)) == "당기"]
            for c, pc, pv in slots:
                cell = tv.cell(i, c)
                if cell is None:
                    continue
                cur_v = tv.value(i, c)
                if cur_v is not None and keep_filled:
                    last_pair[_grp(tv, c)] = c
                    continue
                if cur_v is None and cell.text.strip():
                    continue                       # 글 칸
                item = FillItem(ctx, tv.table.id, tv.table.rows[i].id, label[:40],
                                tv.col_header(c)[:30].replace("\n", " "), None, "못 채움")
                if not label.strip() or (cfg.exclude and cfg.excluded(_as_recon_item(item, ctx))):
                    continue
                stale = " · 이전 값 남음" if cur_v is not None else ""
                if not scope:
                    item.how = "조서에 이 구획이 없음" + stale
                    rep.items.append(item)
                    continue
                col_keys, neutral = col_keys_of(rec, c)
                rate = bool(_RATE_HDR_RE.search(norm(tv.col_header(c))))
                is3m = bool(_3M_RE.search(_leaf(tv, c)))
                cands = same_name_rows(keys, own, scope, common, col_keys)
                if not cands:
                    tot = _total_of(tv, i, c, new_vals) if is_total_label(label) else None
                    if tot is not None:
                        _put(ops, tv, i, c, pc, tot, new_vals)
                        item.value, item.verdict, item.how = tot, "채움", "합계(위 칸 합 - 조서에 합계 행 없음)"
                        rep.items.append(item)
                        ptot = _total_of(tv, i, pc, new_vals, True) if pc is not None and tv.value(i, pc) is None else None
                        if ptot is not None:
                            _put(ops, tv, i, pc, None, ptot, new_vals)
                    elif not no_prev or cur_v is not None:   # 머리 행('자 산')은 조용히
                        item.how = "조서에 같은 이름 행 없음" + stale
                        rep.items.append(item)
                    continue
                opts = []                          # (행, 열, 당기 값, 전기 값 | None)
                for row in cands:
                    cols, _ = _row_cols(row, tgt, "당기", is3m, one_q, col_keys, True, neutral, strict)
                    vals = [(cc, row.nums[cc]) for cc in cols if cc in row.nums]
                    if not vals:
                        continue
                    pcols, _ = _row_cols(row, ptgt, "전기", is3m, one_q, col_keys, False, neutral, strict)
                    pvals = [pcc for pcc in pcols if pcc in row.nums]
                    for cc, w in vals:
                        # 짝 전기 열 = 이 당기 열 오른쪽 가장 가까운 전기 열(원 · 천원 두 판이 나란히 있어도 같은 판)
                        right = [pcc for pcc in pvals if pcc > cc]
                        pcc = min(right) if right else (max(pvals) if pvals else None)
                        opts.append((row, cc, w, None if pcc is None else row.nums[pcc]))
                # 금액 칸에 율 열('이자율(%)' · '지분율') 후보는 빼고, 율 칸에는 율 열만
                named_rate = [o for o in opts if any(_RATE_HDR_RE.search(k) for k in o[0].names.get(o[1], ()))]
                if named_rate and len(named_rate) < len(opts):
                    opts = named_rate if rate else [o for o in opts if o not in named_rate]
                if any(o[2] for o in opts) and not all(o[2] == 0 for o in opts):
                    opts = [o for o in opts if o[2]]   # 표 오른쪽 점검 칸(0)은 다른 후보가 있으면 버림
                if not opts:
                    if cur_v is not None:          # 조서 칸이 빈 곳(자본변동표 자본금 열의 순손실 등)은 DSD 도 빈 칸
                        item.how = "같은 이름 행에 당기 값 없음" + stale
                        rep.items.append(item)
                    continue
                scales = _scales(unit, opts[0][0].unit, rate)
                got = _choose(opts, None if pv is None else float(pv), scales, last_row, mag, unit)
                if got is None:
                    item.how = ("후보 %d개 값이 다름(%s)" % (len(opts), ", ".join(_ref(o[0], o[1]) for o in opts[:4]))
                                if len(opts) > 1 else "조서 전기 값이 DSD 전기 값과 다름(%s)" % _ref(opts[0][0], opts[0][1]))
                    item.how += stale
                    rep.items.append(item)
                    continue
                w, used, fit, how = got
                if fit:
                    sc, sg = fit
                else:
                    sc = _mag_scale(w, scales, mag)
                    sg = -1 if (rec.kind == "PL" and "손실" in label and "이익" not in label
                                and _RESULT_RE.search(keys_text(label))) else 1
                val = _round(sg * w * sc, rate, None if pv is None else float(pv))
                _put(ops, tv, i, c, pc, val, new_vals)
                last_pair[_grp(tv, c)] = c
                last_row[used[0][0].sheet] = max(r.row for r, _c in used)
                item.value, item.verdict = val, "채움"
                item.where = ", ".join(_ref(r, cc) for r, cc in used[:3])
                item.how = how + (" · 부호 반대 관례" if fit and fit[1] < 0 else "") + (" · 새 행(위 행 열)" if no_prev else "")
                rep.items.append(item)
                # 짝 전기 칸이 비어 있고 조서에 전기 값이 있으면 같은 단위 · 부호로
                if pc is not None and tv.value(i, pc) is None and len(used) == 1:
                    o = next(o for o in opts if o[0] is used[0][0] and o[1] == used[0][1])
                    if o[3] is not None and tv.cell(i, pc) is not None and not tv.cell(i, pc).text.strip():
                        pval = _round(sg * o[3] * sc, rate, None)
                        _put(ops, tv, i, pc, None, pval, new_vals)
                        rep.items.append(FillItem(ctx, tv.table.id, tv.table.rows[i].id, label[:40],
                                                  tv.col_header(pc)[:30].replace("\n", " "), pval, "채움",
                                                  "전기 빈칸 - 조서 전기 값", _ref(o[0], o[1])))
        if rec.kind == "PL":
            relabel_pl(rec, new_vals, ops, rep, ctx)
    return ops, rep


_PL_RESULT_RE = __import__("re").compile(
    r"(매출총|영업|법인세비용차감전\s*순|법인세차감전\s*순|당기\s*순|분기\s*순|반기\s*순|순|총포괄)\s*(이익|손실)(?!\s*\()")


def relabel_pl(rec, new_vals: dict, ops: list, rep: "FillReport", ctx: str) -> None:
    """손익계산서 결과 행('매출총이익' · '영업손실' · '분기순이익')의 값 부호가 섞이면 라벨을 '…이익(손실)'로.
    값은 그대로(이익 +, 손실 -). '…손실' 라벨에 값이 모두 양수(손실 양수 관례)면 그대로 둔다."""
    tv = rec.tv
    for i in tv.body:
        cell = tv.cell(i, 0)
        if cell is None:
            continue
        m = _PL_RESULT_RE.search(cell.text)
        if not m:
            continue
        vals = [new_vals.get((i, c), tv.value(i, c)) for c in range(1, tv.ncols)]
        vals = [float(v) for v in vals if v not in (None, 0)]
        if not vals:
            continue
        pos, neg = any(v > 0 for v in vals), any(v < 0 for v in vals)
        if not (pos and neg):
            continue
        new = cell.text[:m.start()] + m.group(1) + "이익(손실)" + cell.text[m.end():]
        ops.append(splice.replace_inner(cell, new))
        rep.items.append(FillItem(ctx, tv.table.id, tv.table.rows[i].id, cell.text.strip()[:40], "(라벨)",
                                  None, "채움", "라벨 '%s' → '%s'(이익 · 손실 칸이 섞임)" % (
                                      cell.text.strip(), new.strip())))


def _put(ops, tv, i: int, c: int, like_col: int | None, val: float, new_vals: dict) -> None:
    cell = tv.cell(i, c)
    like = tv.cell(i, like_col).text if like_col is not None and tv.cell(i, like_col) is not None else cell.text
    new = to_dsd_text(int(val) if val == int(val) else val, like=like) if val != 0 else "-"
    ops.append(splice.replace_inner(cell, new))
    new_vals[(i, c)] = val


def _total_of(tv, i: int, c: int, new_vals: dict, existing: bool = False) -> float | None:
    """합계 행(조서에 합계 행이 없음) = 바로 위 합계 행 다음부터 이 행 앞까지 같은 열 값의 합. 위 칸이 다 있을 때만."""
    rows = []
    for j in reversed([x for x in tv.body if x < i]):
        if is_total_label(tv.row_label(j)):
            break
        rows.append(j)
    vals = []
    for j in rows:
        if (j, c) in new_vals:
            vals.append(new_vals[(j, c)])
        elif tv.value(j, c) is not None:
            if not existing:
                return None                        # 이번에 안 쓴(이전 값일 수 있는) 칸이 섞임
            vals.append(float(tv.value(j, c)))
    return sum(vals) if vals else None


def _partner_cols(rec, recs) -> list[int]:
    """같은 구획에서 머리글(0열 빼고)이 같은 다른 표의 값 열('당분기말 | 지분율 | …' ↔ '전기말 | 지분율 | …')."""
    tv = rec.tv
    hdr = [norm(tv.col_header(c)) for c in range(1, tv.ncols)]
    for o in recs:
        if o is rec or o.sec is not rec.sec or o.tv.ncols != tv.ncols:
            continue
        other = [norm(o.tv.col_header(c)) for c in range(1, o.tv.ncols)]
        if sum(a == b for a, b in zip(hdr, other)) * 2 < len(hdr):
            continue                               # 머리글 절반 이상 같아야('분기순손익' ↔ '당기순손익')
        cols = sorted({c for i in o.tv.body for c in range(1, o.tv.ncols) if o.tv.value(i, c) is not None})
        if cols:
            return cols
    return []


def _grp(tv, c: int) -> str:
    """열 묶음 이름 = 머리글 글('제48기 (당)3분기/3 개 월'). 세부 · 소계 두 열이 병합 머리글을 같이 쓴다."""
    return norm(tv.col_header(c))


def keys_text(label: str) -> str:
    return norm(label).replace("(", "").replace(")", "")


def _choose(opts, pv, scales, last_row, mag=None, unit: int = 1):
    """후보 → (조서 값, 쓴 (행, 열)들, 앵커(배수, 부호) | None, 방법) 또는 None(못 고름).
    1) 전기 값이 DSD 전기 값과 맞는 행(앵커) 2) 같은 이름 행 여럿의 전기 합이 맞으면 하위 합
    3) 값이 하나로 모이면 그것 4) 바로 앞에서 채운 조서 행 다음 행."""
    if pv is not None:
        hit = [(o, f) for o in opts if o[3] is not None and (f := _fit(pv, o[3], scales))]
        if hit:
            nd = 0 if pv == int(pv) else 2
            vals = {round(o[2] * f[0] * f[1], nd) for o, f in hit}    # 앵커 단위 · 부호로 맞춘 값(DSD 자릿수)
            if len(vals) == 1:
                return hit[0][0][2], [(hit[0][0][0], hit[0][0][1])], hit[0][1], "앵커(전기 값 일치)"
            hit = _after(hit, last_row, key=lambda x: x[0][0])
            if hit:
                o = hit[0][0]                  # _after 는 (후보, 배수) 짝을 돌려준다
                return o[2], [(o[0], o[1])], None, "앵커 · 다음 행"
        rows = {}
        for o in opts:                         # 행마다 첫 값
            rows.setdefault(id(o[0]), o)
        sub = list(rows.values())
        if len(sub) > 1 and all(o[3] is not None for o in sub):
            f = _fit(pv, sum(o[3] for o in sub), scales)
            if f:
                return sum(o[2] for o in sub), [(o[0], o[1]) for o in sub], f, "하위 합 %d행(전기 합 일치)" % len(sub)
        if any(o[3] is not None for o in opts) and not any(o[3] is None for o in opts):
            if len({o[2] for o in opts}) == 1:   # 후보 하나 - 채우고 전기 칸 확인 표시(전기 기준이 다를 수 있음)
                return opts[0][2], [(opts[0][0], opts[0][1])], None, "전기 값이 조서와 다름 - 전기 칸 확인"
            return None                        # 전기 값이 있는데 아무것도 안 맞음 - 다른 행
    got = _pick_same(opts, unit, mag)
    if got:
        return got
    near = _after(opts, last_row, key=lambda x: x[0])
    if near:                                   # 바로 앞에서 채운 조서 행(같은 행 포함) 다음 - 가장 가까운 행의 후보만
        grp = [o for o in near if o[0] is near[0][0]]
        got = _pick_same(grp, unit, mag)
        if got:
            return got[0], got[1], None, "바로 앞 행 다음"
    return None


def _eff_unit(recs, k: int) -> int:
    """DSD 표 단위 - 표 머리에 없으면 같은 구획 바로 앞 한 칸 표('(단위: 천원)')."""
    tv = recs[k].tv
    if tv.unit != 1 or k == 0:
        return tv.unit
    prev = recs[k - 1]
    if prev.sec is recs[k].sec and len(prev.tv.table.rows) == 1:
        texts = " ".join(prev.tv.text(0, c) for c in range(prev.tv.ncols))
        m = _UNIT_TXT_RE.search(texts)
        if m:
            return _UNIT_MULT.get(norm(m.group(1)).split(",")[-1], 1)
    return 1


def _pick_same(opts, unit: int, mag):
    """후보 값이 하나로 모이면 그것. 원 · 천원 두 판(같은 숫자 1,000배 차이)이면 DSD 단위에 맞는 판
    (DSD 천원 = 작은 쪽, 단위 모름 = 다른 값 자릿수에 가까운 쪽) - 반올림 ±1."""
    vals = {o[2] for o in opts}
    if len(vals) == 1:
        return opts[0][2], [(opts[0][0], opts[0][1])], None, ("하나뿐" if len(opts) == 1 else "같은 값 %d개" % len(opts))
    if not all(o[2] for o in opts):
        return None
    small = min(opts, key=lambda o: abs(o[2]))
    if not all(any(abs(o[2] / 1000 ** k - small[2]) <= 1 for k in (0, 1, 2)) for o in opts):
        return None
    if unit >= 1000:
        ref = small
    elif mag:
        ref = min(opts, key=lambda o: abs(math.log10(abs(o[2])) - math.log10(mag)))
    else:
        ref = max(opts, key=lambda o: abs(o[2]))
    return ref[2], [(ref[0], ref[1])], (1.0, 1), "같은 값 %d개(원 · 천원 판 - DSD 단위)" % len(opts)


def _mag_scale(w: float, scales: list[float], mag: float | None) -> float:
    """앵커가 없을 때 단위 배수: 1배가 기본, 다른 배수가 표(또는 주석) 값 자릿수에 1.5자리 이상 더 가까울 때만."""
    if not mag or not w:
        return 1.0 if 1.0 in scales else scales[0]

    def dist(x):
        return abs(math.log10(abs(w * x)) - math.log10(mag))
    base = 1.0 if 1.0 in scales else scales[0]
    best = min(scales, key=dist)
    return best if dist(best) + 1.5 <= dist(base) else base


def _after(items, last_row, key):
    """바로 앞에서 채운 조서 행부터 아래에 있는 후보를 가까운 순으로."""
    out = []
    for x in items:
        r = key(x)
        lr = last_row.get(r.sheet)
        if lr is not None and r.row >= lr:          # 같은 행(다음 열)도
            out.append((r.row - lr, x))
    out.sort(key=lambda t: t[0])
    return [x for _d, x in out]


def _as_recon_item(x: FillItem, ctx: str):
    from .recon import Item
    return Item(ctx, x.table, x.row_id, 0, x.label, x.col, "당기", 0.0, None, "")


def write_fill_report(rep: FillReport, out_path: str, dsd_path: str, wp_path: str) -> None:
    import os

    from openpyxl import Workbook

    from .report_xlsx import _BORDER, _FILL_BAD, _NUMFMT, FONT, FONT_TITLE, _write_header

    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    ws.cell(row=2, column=2, value="DSD 당기 칸 채우기 결과").font = FONT_TITLE
    for r, (k, v) in enumerate([("DSD", os.path.basename(dsd_path)), ("조서", os.path.basename(wp_path)),
                                ("보고기간", rep.period), ("채움", rep.count("채움")),
                                ("못 채움", rep.count("못 채움"))], start=4):
        ws.cell(row=r, column=2, value=k).font = FONT
        ws.cell(row=r, column=3, value=v).font = FONT
    for r, w in enumerate(rep.warnings, start=10):
        ws.cell(row=r, column=2, value=w).font = FONT
    heads = ["구획", "표", "행", "열", "값", "결과", "방법 · 이유", "조서 위치"]
    widths = [12, 7, 30, 22, 18, 8, 44, 18]
    for title, items in (("못 채움", [x for x in rep.items if x.verdict == "못 채움"]), ("전체", rep.items)):
        ws = wb.create_sheet(title)
        for k, w in enumerate(widths, start=1):
            ws.column_dimensions[chr(64 + k)].width = w
        _write_header(ws, 1, heads)
        for rr, x in enumerate(items, start=2):
            for cc, v in enumerate([x.sheet, x.table, x.label, x.col, x.value, x.verdict, x.how, x.where], start=1):
                cell = ws.cell(row=rr, column=cc, value=v)
                cell.font = FONT
                cell.border = _BORDER
                if cc == 5 and isinstance(v, float):
                    cell.number_format = _NUMFMT if v == int(v) else "#,##0.00"
                if x.verdict == "못 채움":
                    cell.fill = _FILL_BAD
    wb.properties.creator = "이재현"
    wb.save(out_path)
