"""정산표 → DSD 다시 만들기(build) - 합성 뼈대 · 조서로 레시피 단계 · 빌드 밖 수정 감지."""
from __future__ import annotations

import openpyxl

from dsdxlsx import build as B
from dsdxlsx import recon as R
from dsdxlsx.dsd_io import read_dsd
from dsdxlsx.export_xlsx import export_xlsx
from dsdxlsx.parse import parse_contents

from fixtures.factory import write_dsd
from test_fill import _skeleton, _wp
from test_rollinterim import _row


def _setup(tmp_path):
    skel = tmp_path / "skel.dsd"
    write_dsd(skel, _skeleton().data)
    wp = _wp(tmp_path)
    wb = openpyxl.load_workbook(wp)
    ws = wb.create_sheet("추가")
    ws["B3"], ws["C3"] = "보증금", 77
    ws["D5"] = 0.05143
    wb.save(wp)
    return str(skel), wp


def _locate(skel: str, tmp_path, label: str):
    """합성 문서에서 label 행이 있는 (시트, 표 ID, 행 ID)."""
    out = tmp_path / "look.xlsx"
    export_xlsx(parse_contents(read_dsd(skel).contents), open(skel, "rb").read(), skel, str(out))
    for ws in openpyxl.load_workbook(out).worksheets:
        for r in ws.iter_rows():
            v = r[0].value
            if isinstance(v, str) and B.ROW_ID.match(v) and label in str(r[1].value):
                tid, rid = v.split(".")
                return ws.title, tid, rid
    raise AssertionError(label)


def _build(tmp_path, skel, wp, recipe, name="out.dsd"):
    out = str(tmp_path / name)
    rep = B.build(skel, wp, out, recipe=recipe, fill_conf={}, side_sheets=R.load_workpaper(wp))
    return out, rep


def test_build_without_recipe_equals_fill(tmp_path):
    from dsdxlsx import fill as F
    from dsdxlsx import splice
    skel, wp = _setup(tmp_path)
    out, rep = _build(tmp_path, skel, wp, {})
    doc = parse_contents(read_dsd(skel).contents)
    ops, _ = F.fill_ops(doc, R.load_workpaper(wp))
    assert read_dsd(out).contents == splice.apply(doc.data, ops)
    assert rep.filled >= 10


def test_recipe_order_table_cell_template_and_text(tmp_path):
    skel, wp = _setup(tmp_path)
    sheet, tid, rid = _locate(skel, tmp_path, "보증금")
    recipe = {
        "tables": [{"sheet": sheet, "table": tid, "match": "order", "first": rid, "src": "추가",
                    "rows": "3", "cols": {"C": "C"}}],
        "cells": [{"sheet": sheet, "table": tid, "row_label": "보증금", "col": "D", "src": "추가",
                   "template": "{D5*100:.2f}"}],
        "text": [{"start": "보증금", "end": "</TABLE>", "replace": [["보증금", "임차보증금"]]}],
    }
    out, rep = _build(tmp_path, skel, wp, recipe)
    doc = parse_contents(read_dsd(out).contents)
    assert _row(doc, "임차보증금")[1:3] == ["77", "5.14"]
    assert any("글 바꾸기" in s for s in rep.steps)


def test_recipe_label_match_and_missing_row_is_reported(tmp_path):
    skel, wp = _setup(tmp_path)
    sheet, tid, _ = _locate(skel, tmp_path, "보증금")
    recipe = {"tables": [{"sheet": sheet, "table": tid, "src": "추가", "rows": "1-5", "skip_header": 1,
                          "cols": {"C": "C"}}]}
    out, rep = _build(tmp_path, skel, wp, recipe)
    assert _row(parse_contents(read_dsd(out).contents), "보증금")[1] == "77"
    assert all("보증금" not in w for w in rep.warnings)


def test_outside_edit_detection(tmp_path):
    skel, wp = _setup(tmp_path)
    out, _ = _build(tmp_path, skel, wp, {})
    assert B.outside_edits(out) is None             # 기준 없음
    B.save_last_build(out)
    assert B.outside_edits(out) == []
    pkg = read_dsd(out)
    edited = pkg.contents.replace("보증금".encode(), "보증금(수정)".encode(), 1)
    from dsdxlsx.dsd_io import write_dsd as wd
    wd(out, pkg, edited)
    lines = B.outside_edits(out)
    assert lines and any("보증금(수정)" in ln for ln in lines)


def test_rows_spec_and_template():
    assert B._rows_spec("7-9,54") == [7, 8, 9, 54]
    assert B._rows_spec([157, "159-160"]) == [157, 159, 160]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["D1"], ws["D2"], ws["D3"] = 1234567, 0.05143, "-"
    assert B._template("{D1:,.0f}\nUSD {D1:,.0f}", ws) == "1,234,567\nUSD 1,234,567"
    assert B._template("{D2*100:.2f}", ws) == "5.14"
    assert B._template("{D3:.2f}", ws) is None      # 숫자 아님 - 칸을 건드리지 않음
